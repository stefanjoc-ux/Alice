"""Digital teams: drawings read visually, Assume and flag, and files added to a job that has started (Stefan, 9 Oct 2026).

Vision calls are stand-ins (no real model): drawing pages go as images only when needed (a Drawing's pages, an image file, a PDF page
with little text), through the Purview label and check_outbound, with their cost recorded against the member and on the job; with
Assume and flag the team carries on and lists its assumptions in the outputs, while Ask me still asks; a file added mid-job re-runs only
the work that depends on it, as a new version, and a file added from the question panel answers that question."""
import _util  # first: throwaway data folder, dummy keys, no real model calls
from _util import t
import base64, io, json, re
from types import SimpleNamespace
import app, substrate_store as s, assistants, org_research, usage_meter, documents, doc_library, permissions
import teams, team_qs, team_costs, team_files
from fastapi.testclient import TestClient
from pypdf import PdfReader, PdfWriter
cl = TestClient(app.app); H = {'x-admin-token': app.ADMIN_TOKEN}
teams.BACKGROUND = False

USAGE = SimpleNamespace(usage={'input_tokens': 10000, 'output_tokens': 1000})        # Sonnet: $0.03 a call
VISION, CALLS = [], []
STATE = {'plan': None, 'measure_q': [], 'reading': None}
SPEC, DRAW = 'FICTIONAL hall spec.md', 'FICTIONAL hall drawings.pdf'


def fake_call(provider, system, messages, max_tokens=1500, timeout=60, workload='', meta=None):
    content = messages[0]['content']
    usage_meter.log(USAGE, 'claude' if provider.startswith('claude') else 'openai', 'claude-sonnet-5-5' if provider.startswith('claude') else 'gpt-6-luna', workload)
    if isinstance(content, list):                                  # a page sent as an image
        att = [b for b in content if b.get('type') in ('document', 'image', 'input_file', 'input_image')]
        text = next(b['text'] for b in content if b.get('type') in ('text', 'input_text'))
        VISION.append({'provider': provider, 'workload': workload, 'types': [b['type'] for b in att], 'text': text, 'system': system,
                       'pages': [len(PdfReader(io.BytesIO(base64.b64decode(b['source']['data']))).pages) for b in att if b.get('type') == 'document']})
        page = re.search(r'PAGE: (\d+)', text).group(1)
        if STATE['reading']: return STATE['reading']
        return json.dumps({'drawing_number': f'A-10{page}', 'title': 'Ground floor plan', 'scale': '1:100 at A1',
                           'dimensions': [f'Hall internal width 12.0 m (page {page})'], 'levels': ['FFL 10.150'], 'notes': ['Walls: 140 mm blockwork'], 'unreadable': ''})
    CALLS.append({'workload': workload, 'payload': content, 'system': system})
    if 'Lead QS' in workload and 'COST PLAN FIGURES' in content:
        return json.dumps({'accept': True, 'summary': 'A hall.', 'assumptions': [], 'exclusions': ['VAT'], 'risks': [], 'note': 'On.'})
    if 'Lead QS' in workload and 'SUGGESTED MARKET ADJUSTMENT' in content:
        return json.dumps({'accept': False, 'reason': 'No.', 'summary': 'Rejected.'})
    if 'Lead QS' in workload: return json.dumps(STATE['plan'])
    if 'Measurement' in workload:
        if STATE['measure_q']: return json.dumps(STATE['measure_q'].pop(0))
        part = content.split('THIS PART: take off ONLY these elements:')[1].split('from the DOCUMENTS')[0]
        items = [{'element': e, 'description': f'{e} item', 'quantity': 10, 'unit': 'm2', 'source': {'document': SPEC, 'page': '1'}}
                 for e in ('Walls', 'Roof') if e in part]
        return json.dumps({'accept': True, 'items': items, 'summary': '', 'note': 'Price these.'})
    if 'Market Trends' in workload: return json.dumps({'matches': [], 'commentary': 'Nothing to compare.', 'summary': ''})
    if 'Checker' in workload or 'Writer' in workload: return json.dumps({'accept': True, 'output': 'Done: ' + workload, 'summary': 'Done.', 'note': 'On.'})
    raise AssertionError('unexpected call ' + workload)


def fake_ask(prompt, query, provider, workload='', **kw):
    usage_meter.log(USAGE, 'claude', 'claude-sonnet-5-5', workload)
    if 'MEASURED ITEMS' not in query: return json.dumps({'findings': [], 'commentary': 'No market evidence.'}), {}
    CALLS.append({'workload': workload, 'payload': query, 'system': prompt})
    refs = re.findall(r'^(Q\d+) \|', query.split('MEASURED ITEMS IN THIS PART')[1].split('RATE LIBRARY')[0], re.M)
    return json.dumps({'rates': [{'ref': r, 'source': 'published', 'rate': 50, 'unit': 'm2', 'source_url': 'https://fictional-prices.example/q1',
                                  'source_date': '2026-06'} for r in refs], 'searches': ['rates']}), {'https://fictional-prices.example/q1': 'Prices'}


assistants._call = fake_call
org_research._ask = fake_ask


def pdf(*pages):
    """A fictional PDF: each page either text (a string, from Alice's own PDF writer) or None (a blank page, as a drawing scan is)."""
    w = PdfWriter()
    for p in pages:
        if p is None: w.add_blank_page(595, 842)
        else: w.append(PdfReader(io.BytesIO(documents.to_pdf('FICTIONAL', p))), pages=(0, 1))
    b = io.BytesIO(); w.write(b); return b.getvalue()


PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')
LONG = 'FICTIONAL specification for a community hall. ' * 12
b64 = lambda raw: base64.b64encode(raw).decode()


def plan(*els, **kw):
    return {'plan': 'Measure.', 'elements': [{'name': e, 'documents': []} for e in els], 'documents': [{'name': SPEC, 'use': 'all'}],   # no documents named: every one
            'location': '', 'summary': 'Planned.', 'note': 'Go.', 'questions': [], **kw}


def start(tid, title, uploads, mode=None):
    r = cl.post(f'/admin/api/teams/{tid}/jobs', headers=H, json={'job_type': 'cost-estimate', 'title': title, 'client': '',
                'brief': 'A fictional single-storey hall for a community trust, to be estimated.', 'uploads': uploads})
    assert r.status_code == 200, r.text
    return r.json()


team = teams.from_template('quantity-surveying', name='Files test team')
TID = team['id']
teams.set_autonomy(TID, 'signoff')
t('the QS template starts on Assume and flag; a team saved before the setting asks', teams.missing_info(teams.get(TID)) == 'assume'
  and teams.missing_info({'name': 'old'}) == 'ask' and teams.MISSING_INFO['assume'] == 'Assume and flag')

# ---------------- drawings read visually, only when needed ----------------
STATE['plan'] = plan('Walls', 'Roof')
up = [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. [Page 1] walls and roof.'},
      {'name': DRAW, 'kind': 'drawing', 'data': b64(pdf(None, 'FICTIONAL elevation notes ' * 3))},
      {'name': 'FICTIONAL scanned schedule.pdf', 'kind': 'schedule', 'data': b64(pdf(LONG, None))},
      {'name': 'FICTIONAL site photo plan.png', 'kind': 'drawing', 'data': b64(PNG)}]
j = start(TID, 'FICTIONAL hall with drawings', up)
t('the job runs to sign-off', j['status'] == 'waiting' and j['pending'][0]['kind'] == 'signoff')
R = team_files.reads(j['id'])
got = sorted((r['doc_name'], r['page']) for r in R)
t('pages read as images: every page of the drawing, the image, and only the scanned page of the schedule (never the text spec or a page with text)',
  got == sorted([(DRAW, 1), (DRAW, 2), ('FICTIONAL scanned schedule.pdf', 2), ('FICTIONAL site photo plan.png', 1)]) and all(r['status'] == 'read' for r in R))
t('each page is sent on its own: a one-page PDF cut from the document, or the image itself', len(VISION) == 4
  and sum(1 for v in VISION if v['types'] == ['document'] and v['pages'] == [1]) == 3 and sum(1 for v in VISION if v['types'] == ['image']) == 1)
t('…to Sonnet 5.5 (the Lead QS\'s model), asking for dimensions, levels, notes, scale and drawing number, citing the page',
  all(v['provider'] == 'claude_sonnet' for v in VISION) and 'drawing number' in VISION[0]['system'] and 'scale' in VISION[0]['system']
  and all(re.search(r'DOCUMENT: .+\nPAGE: \d', v['text']) for v in VISION))
t('read once per job: the Measurement Surveyor reuses the readings (no second round of image calls)', len({v['workload'] for v in VISION}) == 1 and 'Lead QS' in VISION[0]['workload'])
meas = next(c for c in CALLS if 'Measurement' in c['workload'])
t('members see each reading under its page, with the drawing number, scale and page cited', '[Page 1, read from the image by Claude Sonnet 5.5; drawing A-101' in meas['payload']
  and 'scale 1:100 at A1' in meas['payload'] and f'Dimension: Hall internal width 12.0 m (page 1) ({DRAW}, page 1, drawing A-101)' in meas['payload'])
with s.db() as c:
    tc = [dict(r) for r in c.execute("SELECT * FROM team_costs WHERE job_id=? AND kind='drawing'", (j['id'],))]
t('image costs are recorded against the member who needed them (a team cost of kind drawing)', len(tc) == 4 and all(r['member_id'] == 'lead-qs' and abs(r['usd'] - 0.03) < 1e-9 for r in tc))
t('…on each page\'s reading, and in the job\'s AI cost', all(abs(r['cost_usd'] - 0.03) < 1e-9 for r in R)
  and teams._row(j['id'])['ai_cost'] >= 0.12 - 1e-9)
pg = cl.get(f'/admin/api/teams/jobs/{j["id"]}/page').json()
t('the job page lists the pages read, the model, the member and the cost', pg['drawings']['read'] == 4 and abs(pg['drawings']['total']['usd'] - 0.12) < 1e-9
  and pg['drawings']['pages'][0]['role'] == 'Lead QS' and pg['drawings']['pages'][0]['model'] == 'claude-sonnet-5-5')
t('…and the job\'s history says which pages were read', any(x['kind'] == 'drawings' and '4 pages read as images' in x['text'] for x in pg['timeline']))

# ---------------- the checks on every page ----------------
VISION.clear()
real_label = doc_library._label_action
doc_library._label_action = lambda p, raw, fam: 'Purview label "FICTIONAL Secret" is blocked' if 'labelled' in p.name else real_label(p, raw, fam)
STATE['reading'] = json.dumps({'drawing_number': 'X-1', 'scale': '', 'dimensions': [], 'levels': [], 'notes': ['OFFICIAL-SENSITIVE'], 'unreadable': ''})
up2 = [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls.'},
       {'name': 'FICTIONAL labelled drawing.pdf', 'kind': 'drawing', 'data': b64(pdf(None))},
       {'name': 'FICTIONAL marked drawing.pdf', 'kind': 'drawing', 'data': b64(pdf(None))}]
j2 = start(TID, 'FICTIONAL hall, checks', up2)
doc_library._label_action = real_label
STATE['reading'] = None
R2 = {r['doc_name']: r for r in team_files.reads(j2['id'])}
t('a page whose Purview label may not go to the model is never sent, and says why', R2['FICTIONAL labelled drawing.pdf']['status'] == 'refused'
  and 'Purview' in R2['FICTIONAL labelled drawing.pdf']['reason'] and len(VISION) == 1)
t('a reading that comes back with a protective marking is not kept or passed on', R2['FICTIONAL marked drawing.pdf']['status'] == 'refused'
  and 'OFFICIAL-SENSITIVE' not in R2['FICTIONAL marked drawing.pdf']['text'])
meas2 = [c for c in CALLS if 'Measurement' in c['workload']][-1]
t('…and members are told the page was not read (never the refusal\'s own words, which may name a marking or label)', j2['status'] == 'waiting'
  and '[Page 1 was not read as an image: held back by Alice\'s rules' in meas2['payload'] and 'OFFICIAL-SENSITIVE' not in meas2['payload'] and 'FICTIONAL Secret' not in meas2['payload'])
import rules_engine
blocked = {'n': 0}
real_check = rules_engine.check_outbound
def fussy(text, target='chat message', provider=None, packs=True):
    if '(drawing page)' in target:
        blocked['n'] += 1
        if 'refuse me' in text: raise rules_engine.RuleViolation('FICTIONAL rule: this page may not be sent.')
    return real_check(text, target, provider, packs)
rules_engine.check_outbound = fussy
VISION.clear()
j3 = start(TID, 'FICTIONAL hall, outbound', [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls.'},
                                             {'name': 'FICTIONAL noted drawing.pdf', 'kind': 'drawing', 'data': b64(pdf('refuse me'))}])
rules_engine.check_outbound = real_check
r3 = team_files.reads(j3['id'])
t('check_outbound runs on every page sent (its own text included), and a refusal keeps the page back', blocked['n'] == 1 and r3[0]['status'] == 'refused'
  and 'may not be sent' in r3[0]['reason'] and not VISION)

# ---------------- Assume and flag ----------------
CALLS.clear()
STATE['plan'] = plan('Walls', 'Roof', assumed=[{'assumption': 'The hall is single storey', 'why': 'No sections were provided', 'affects': 'Frame and roof'}],
                     questions=['Is there a mezzanine?'])
ja = start(TID, 'FICTIONAL hall, assume', [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls.'}])
pa = cl.get(f'/admin/api/teams/jobs/{ja["id"]}/page').json()
t('Assume and flag: the member carries on instead of asking (no reason given that it matters)', pa['status'] == 'waiting' and pa['pending'][0]['kind'] == 'signoff'
  and not any(x['kind'] == 'question' for x in pa['timeline']))
asm = [a['assumption'] for a in pa['assumed']]
t('…every assumption is listed, with what was missing and who made it', 'The hall is single storey' in asm and any(a.startswith('Not asked: Is there a mezzanine?') for a in asm)
  and pa['assumed'][0]['role'] == 'Lead QS' and pa['assumed'][0]['why'] == 'No sections were provided')
t('…the members are told to assume, list and ask only when it matters materially', 'make a reasonable assumption' in CALLS[0]['system'] and 'why_ask' in CALLS[0]['system'])
md, _cp = team_qs._md(teams._row(ja['id']), teams._row(ja['id'])['outputs'])
t('…in the cost plan (Word)', '# Assumed where information was missing' in md and 'The hall is single storey (missing: No sections were provided) Affects: Frame and roof.' in md)
docs = teams._row(ja['id'])['outputs']['documents']
xl = next(d for d in docs if d['kind'] == 'Excel')
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(documents.get(xl['id'])['data']))
t('…and in the workbook\'s Assumptions sheet', any(r[0] == 'Assumed (information missing)' and 'single storey' in (r[1] or '') for r in wb['Assumptions'].iter_rows(values_only=True)))
import pricing_templates
ws = openpyxl.Workbook(); pricing_templates._alice_sheet(ws, {'ref': 'J-1', 'title': 'x', 'outputs': teams._row(ja['id'])['outputs']}, _cp, [])
t('…and on Alice\'s figures sheet of a filled template', any('single storey' in str(r[0]) for r in ws["Alice's figures"].iter_rows(values_only=True)))
STATE['plan'] = plan('Walls', questions=['Is the hall two storeys?'], why_ask='Two storeys would roughly double the frame and floor costs.')
jw = start(TID, 'FICTIONAL hall, material question', [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls.'}])
t('…but a question that would change the result materially still goes to Stefan, saying why', jw['status'] == 'waiting' and jw['pending'][0]['kind'] == 'question'
  and jw['pending'][0]['content']['why'] == 'Two storeys would roughly double the frame and floor costs.')
card = cl.get(f'/admin/api/cards/review/h-{jw["pending"][0]["id"]}').json()
t('…and the question card says why it matters', 'Why it matters: Two storeys' in json.dumps(card))
v0 = teams.get(TID)['version']
r = cl.put(f'/admin/api/teams/{TID}/missing-info', json={'mode': 'ask'}, headers=H)
t('the setting is saved as a new team version', r.status_code == 200 and teams.get(TID)['version'] == v0 + 1 and teams.missing_info(teams.get(TID)) == 'ask'
  and teams.versions(TID)[0]['what'] == 'When information is missing: Ask me')
STATE['plan'] = plan('Walls', questions=['Is there a mezzanine?'])
jk = start(TID, 'FICTIONAL hall, ask me', [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls.'}])
t('Ask me still asks (no reason needed)', jk['status'] == 'waiting' and jk['pending'][0]['kind'] == 'question' and 'mezzanine' in jk['pending'][0]['note']
  and 'ask Stefan in "questions" rather than guessing' in [c for c in CALLS if 'Lead QS' in c['workload']][-1]['system'])
t('an unknown setting is refused', cl.put(f'/admin/api/teams/{TID}/missing-info', json={'mode': 'guess'}, headers=H).status_code == 422)
teams.set_missing_info(TID, 'assume')

# ---------------- a file added mid-job: only what depends on it, as a new version ----------------
STATE['plan'] = plan('Walls', 'Roof')
jf = start(TID, 'FICTIONAL hall, add a file', [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls and roof.'}])
t('the job reaches sign-off at v1', jf['status'] == 'waiting' and jf['version'] == 1)
CALLS.clear(); VISION.clear()
r = cl.post(f'/admin/api/teams/jobs/{jf["id"]}/files', headers=H, json={'uploads': [{'name': 'FICTIONAL roof drawing rev B.pdf', 'kind': 'drawing', 'data': b64(pdf(None))}],
                                                                       'elements': ['Roof'], 'note': 'Revision B of the roof.'})
x = r.json()
J = teams._row(jf['id'])
t('adding a file to a measured job starts a new version', r.status_code == 200 and x['rerun'] == 'started' and J['version'] == 2
  and teams.job_versions(jf['id'])['versions'][-1]['kind'] == 'files')
wl = [c['workload'] for c in CALLS]
t('…which re-runs only what depends on it: no new plan, the Roof measured again, only its items priced again',
  not any('Lead QS' in w and 'plan' in c['system'].lower() and 'COST PLAN FIGURES' not in c['payload'] and 'SUGGESTED' not in c['payload'] for w, c in zip(wl, CALLS))
  and [c for c in CALLS if 'Measurement' in c['workload']] and all('ONLY these elements: Roof' in c['payload'] for c in CALLS if 'Measurement' in c['workload'])
  and all('Walls item' not in c['payload'] for c in CALLS if 'MEASURED ITEMS' in c['payload']))
t('…the new drawing was read as an image for the take-off', len(VISION) == 1 and 'Measurement' in VISION[0]['workload'])
items = {i['element']: i for i in J['outputs']['price']['items']}
t('…the Walls item and its price are kept exactly', items['Walls']['ref'] == 'Q1' and items['Walls']['rate'] == 50 and J['status'] == 'waiting')
steps = teams._steps(jf['id'])
t('…and the job\'s history records the file, who added it and the note', any(s_['kind'] == 'files' and 'FICTIONAL roof drawing rev B.pdf (Drawing)' in s_['note']
  and 'Revision B' in s_['note'] for s_ in steps))
t('the new file joins the documents of the elements it concerns', 'FICTIONAL roof drawing rev B.pdf' in J['outputs']['plan']['element_documents']['Roof']
  and 'FICTIONAL roof drawing rev B.pdf' not in J['outputs']['plan']['element_documents'].get('Walls', []))
t('a file with the same name as one the job has is refused', cl.post(f'/admin/api/teams/jobs/{jf["id"]}/files', headers=H, json={'uploads': [{'name': SPEC, 'kind': 'spec', 'text': 'x y z'}]}).status_code == 400)
t('secrets in an added file are refused, and nothing is added', cl.post(f'/admin/api/teams/jobs/{jf["id"]}/files', headers=H,
  json={'uploads': [{'name': 'FICTIONAL keys.md', 'kind': 'spec', 'text': 'key sk-ant-api03-' + 'A' * 40}]}).status_code == 400
  and not any(d['name'] == 'FICTIONAL keys.md' for d in teams._docs_in(jf['id'])))
# while the team works, the re-run waits for it to stop
teams._set(jf['id'], status='running')
r = cl.post(f'/admin/api/teams/jobs/{jf["id"]}/files', headers=H, json={'uploads': [{'name': 'FICTIONAL wall schedule.md', 'kind': 'schedule', 'text': 'FICTIONAL wall schedule lines.'}],
                                                                       'elements': ['Walls']})
t('while the team is working, the re-run waits (nothing is re-run under it)', r.json()['rerun'] == 'pending' and teams._row(jf['id'])['version'] == 2
  and teams._row(jf['id'])['outputs']['_files_pending']['elements'] == ['Walls'])
teams._set(jf['id'], status='waiting')
team_files.apply_pending(jf['id'])
J = teams._row(jf['id'])
t('…and starts as soon as the team stops for you: v3, only the Walls measured again', J['version'] == 3 and '_files_pending' not in J['outputs']
  and [c for c in CALLS if 'Measurement' in c['workload']][-1]['payload'].count('ONLY these elements: Walls') == 1)
so = next(s_ for s_ in teams._steps(jf['id']) if s_['status'] == 'pending' and s_['kind'] == 'signoff')
cl.post(f'/admin/api/teams/steps/{so["id"]}', json={'action': 'approve'}, headers=H)
r = cl.post(f'/admin/api/teams/jobs/{jf["id"]}/files', headers=H, json={'uploads': [{'name': 'FICTIONAL late.md', 'kind': 'spec', 'text': 'Late.'}]})
t('a signed-off job takes no files (copy it as a new job instead)', teams._row(jf['id'])['status'] == 'done' and r.status_code == 400 and 'copy it' in r.json()['detail'])

# a job not measured yet just uses the file
STATE['plan'] = plan('Walls', questions=['Is there a mezzanine?'], why_ask='It changes the floor area.')
jn = start(TID, 'FICTIONAL hall, early file', [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls.'}])
r = cl.post(f'/admin/api/teams/jobs/{jn["id"]}/files', headers=H, json={'uploads': [{'name': 'FICTIONAL brief note.md', 'kind': 'brief', 'text': 'A note on the brief.'}]})
t('a file added before anything is planned or measured re-runs nothing: the team uses it from here on', r.json()['rerun'] == 'none'
  and teams._row(jn['id'])['version'] == 1)

# ---------------- a file added from the question panel answers it ----------------
STATE['plan'] = plan('Walls', 'Roof')
STATE['measure_q'] = [{'accept': True, 'items': [], 'questions': ['Please provide the roof drawing: the roof is not in the documents.'],
                       'why_ask': 'Without it the roof cannot be measured.', 'summary': '', 'note': ''}]
jq = start(TID, 'FICTIONAL hall, asks for a drawing', [{'name': SPEC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls.'}])
q = jq['pending'][0]
t('a member asking for a document: the question panel offers Add a file', q['kind'] == 'question' and q['content']['asks_file'] is True
  and any(a['label'] == 'Add a file to answer it' for a in cl.get(f'/admin/api/cards/review/h-{q["id"]}').json()['actions']))
VISION.clear()
r = cl.post(f'/admin/api/teams/jobs/{jq["id"]}/files', headers=H, json={'uploads': [{'name': 'FICTIONAL roof drawing.png', 'kind': 'drawing', 'data': b64(PNG)}],
                                                                       'answers': q['id']})
J = teams._row(jq['id'])
qs_ = next(s_ for s_ in teams._steps(jq['id']) if s_['id'] == q['id'])
t('the file answers the question: answered, saying what was added, and the member carries on (no new version)', r.status_code == 200 and qs_['status'] == 'answered'
  and 'FICTIONAL roof drawing.png (Drawing)' in qs_['decision_note'] and J['version'] == 1 and J['status'] == 'waiting' and J['outputs'].get('measure'))
t('…the member saw the new drawing, read as an image', len(VISION) == 1 and 'FICTIONAL roof drawing.png' in [c for c in CALLS if 'Measurement' in c['workload']][-1]['payload'])
t('a question already answered cannot be answered by a file', cl.post(f'/admin/api/teams/jobs/{jq["id"]}/files', headers=H,
  json={'uploads': [{'name': 'FICTIONAL other.md', 'kind': 'spec', 'text': 'Other.'}], 'answers': q['id']}).status_code == 400)

# ---------------- a team without its own rule: from the first stage ----------------
gen = teams.create('FICTIONAL writing team', members=[{'id': 'w', 'role': 'Writer'}, {'id': 'c', 'role': 'Checker'}],
                   job_types=[{'id': 'jt', 'name': 'Note', 'stages': [{'key': 's1', 'title': 'Write', 'member': 'w'}, {'key': 's2', 'title': 'Check', 'member': 'c'}]}],
                   autonomy='signoff')
r = cl.post(f'/admin/api/teams/{gen["id"]}/jobs', headers=H, json={'job_type': 'jt', 'title': 'FICTIONAL note', 'brief': 'Write a fictional note about the hall project please.',
            'uploads': [{'name': 'FICTIONAL notes.md', 'kind': 'brief', 'text': 'Notes.'}]})
jg = r.json()
CALLS.clear()
r = cl.post(f'/admin/api/teams/jobs/{jg["id"]}/files', headers=H, json={'uploads': [{'name': 'FICTIONAL more notes.md', 'kind': 'brief', 'text': 'More notes.'}]})
t('a team without its own rule runs again from its first stage, as a new version', r.json()['rerun'] == 'started' and teams._row(jg['id'])['version'] == 2
  and [c['workload'] for c in CALLS] == ['Digital team: Writer', 'Digital team: Checker'] and 'FICTIONAL more notes.md' in CALLS[0]['payload'])

# ---------------- routes and pages ----------------
t('adding files needs Use (run jobs) on the job; the setting needs Manage on the team', permissions.ROUTES.get('POST /admin/api/teams/jobs/{jid}/files') == 'job:use:run'
  and permissions.ROUTES.get('PUT /admin/api/teams/{tid}/missing-info') == 'team:manage')
html = cl.get(f'/admin/teams/{TID}/jobs/{jf["id"]}').text
t('the job page: Brief and files › Add, Add a file beside the answer, assumptions and pages read as images', all(x in html for x in (
  'function filesPanel(', "'Add files'", "'Add a file'", 'function assumedPanel(', 'function drawingsPanel(', "'/files'")))
t('the Rules and autonomy tab: When information is missing', "'When information is missing'" in html and "'/missing-info'" in html)
t('the Start screen accepts images as drawings', ".png,.jpg,.jpeg,.webp" in html)
ins = cl.post(f'/admin/api/teams/{TID}/inspect', headers=H, json={'name': 'FICTIONAL site plan.png', 'data': b64(PNG)}).json()
t('…and guesses an image is a drawing, read for its scale visually', ins['guess'] == 'drawing' and ins['type'] == 'Image' and ins['pages'] == 1 and ins['scale'] is None)
