"""Provisional sums and clearer decisions for digital teams (Stefan, 9 Oct 2026). Fictional data only; no real model is called.

"Provisional sum" is a source in the rate-source rule (allowed by default; a rule saved before it existed allows it until changed), switchable
per job on the Start screen and the decision panel. For work the documents do not let the team measure, the Cost Surveyor proposes a lump
sum with the range found, its reasoning and every page cited with its date; Alice checks the pages, keeps PS in their own section with a
subtotal worked out in code, and includes them in the total, in Word, Excel and filled pricing templates. Items can be excluded as not in
scope (listed under exclusions, never assumptions). A note asking for provisional sums is acted on by the Lead QS, or the rule that stops it
is named with a link. Nothing is left unpriced without Stefan's say, and the estimate and PS buttons are shown, or why not is explained."""
import _util  # first: throwaway data folder, dummy keys, no real model calls
from _util import t
import io, json, re, shutil, subprocess, tempfile
from decimal import Decimal
from pathlib import Path
import openpyxl
from openpyxl.styles import Font
import substrate_store as s
s.init()
import doc_library as L
root = Path(tempfile.mkdtemp(prefix='alice-docs-')) / 'Documents'
L.ROOT = root.resolve()
L.add_source('SharePoint', 'sharepoint', 'SharePoint: QS site')
(root / 'SharePoint' / 'Pricing templates').mkdir()
import app, assistants, org_research, rules_engine, documents
import teams, team_qs, pricing_templates as PT
from fastapi.testclient import TestClient
cl = TestClient(app.app); H = {'x-admin-token': app.ADMIN_TOKEN}
teams.BACKGROUND = False
TPL = root / 'SharePoint' / 'Pricing templates'

# ---------------- the rule ----------------
rs = rules_engine.rate_sources()
t('"Provisional sum" is a source in the rate-source rule, allowed by default and last in the order', rules_engine.RATE_SOURCES['provisional'] == 'Provisional sum'
  and rs['order'][-1] == 'provisional' and 'provisional' in rs['allowed'] and 'provisional sum' in rules_engine.rule('rate_sources')['description'])
with s.db() as c:                                          # a rule saved before provisional sums existed
    c.execute('UPDATE rules SET params=? WHERE id=?', (json.dumps({'order': ['published', 'library', 'built_up', 'estimate'], 'allowed': ['published', 'library']}), 'rate_sources'))
rules_engine._rules_changed() if hasattr(rules_engine, '_rules_changed') else None
rs = rules_engine.rate_sources()
t('a rule saved before provisional sums existed allows them until it is changed; what it said about the others stands', rs['allowed'] == ['published', 'library', 'provisional'])
rp = next(x for x in cl.get('/admin/api/rules').json()['rules'] if x['id'] == 'rate_sources')['params']
t('…and the Rules page shows them in its editor, ticked, so they can be switched off there', rp['order'][-1] == 'provisional' and 'provisional' in rp['allowed'])
rules_engine.update_rule('rate_sources', new_params={'order': ['published', 'library', 'built_up', 'estimate', 'provisional'], 'allowed': ['published', 'library', 'built_up']})
t('…once saved without them, they are not allowed', 'provisional' not in rules_engine.rate_sources()['allowed'])
DEFAULT = {'order': ['published', 'library', 'built_up', 'estimate', 'provisional'], 'allowed': ['published', 'library', 'built_up', 'provisional']}
rules_engine.update_rule('rate_sources', new_params=DEFAULT)
t('where prices come from (team page) lists provisional sums from the rule', [x['key'] for x in team_qs.price_rules()][-2:] == ['provisional', 'unpriced'])

# ---------------- the team and stand-in models ----------------
team = teams.from_template('quantity-surveying', name='Provisional sums test team')
TID = team['id']
teams.set_autonomy(TID, 'signoff')
DOC = 'FICTIONAL hall spec.md'
PLAN = {'plan': 'Measure.', 'elements': [{'name': 'Walls', 'documents': [DOC]}, {'name': 'Drainage', 'documents': [DOC]}],
        'documents': [{'name': DOC, 'use': 'all'}], 'location': 'Perth', 'summary': 'Planned.', 'note': 'Go.', 'questions': []}
ITEMS = [{'element': 'Walls', 'description': 'Wall item 1', 'quantity': 10, 'unit': 'm2', 'source': {'document': DOC, 'page': '1'}},
         {'element': 'Walls', 'description': 'Wall item 2', 'quantity': 20, 'unit': 'm2', 'source': {'document': DOC, 'page': '1'}},
         {'element': 'Drainage', 'description': 'Connection to the existing sewer', 'quantity': 0, 'unit': '', 'unmeasured': True, 'source': {'document': DOC, 'page': '2'}},
         {'element': 'Drainage', 'description': 'Existing manhole', 'quantity': 1, 'unit': 'nr', 'source': {'document': DOC, 'page': '2'}}]
CALLS, SEARCHES = [], []
STATE = {'requests': []}
PAGES = {'https://fictional-prices.example/walls': 'Wall prices', 'https://fictional-costs.example/sewer': 'Typical sewer connection costs, Scotland',
         'https://fictional-costs.example/drainage': 'Drainage costs guide'}
PS = {'Connection to the existing sewer': {'sum': 2500, 'low': 1800, 'high': 3200, 'reasoning': 'Typical connections to an existing public sewer in Perth and Kinross cost £1,800 to £3,200.',
             'sources': [{'source_url': 'https://fictional-costs.example/sewer', 'source_title': 'Typical sewer connection costs', 'source_date': '2026-07', 'cost': 2400},
                         {'source_url': 'https://made-up.example/nowhere', 'source_date': '2026-01'}]}}


def fake_call(provider, system, messages, max_tokens=1500, timeout=60, workload='', meta=None):
    payload = messages[0]['content']
    CALLS.append({'workload': workload, 'payload': payload, 'system': system})
    if 'Lead QS' in workload and "THE USER'S NOTE TO ACT ON" in payload:
        return json.dumps(STATE['requests'].pop(0) if STATE['requests'] else {'provisional_sums': [], 'estimates': []})
    if 'Lead QS' in workload and 'COST PLAN FIGURES' in payload:
        return json.dumps({'accept': True, 'summary': 'A hall.', 'assumptions': [], 'exclusions': ['VAT'], 'risks': [], 'note': 'On.'})
    if 'Lead QS' in workload: return json.dumps(PLAN)
    if 'Measurement' in workload:
        part = payload.split('THIS PART: take off ONLY these elements:')[1].split('from the DOCUMENTS')[0]
        return json.dumps({'accept': True, 'items': [x for x in ITEMS if x['element'] in part], 'summary': '', 'note': 'Price these.'})
    if 'Market Trends' in workload: return json.dumps({'matches': [], 'commentary': 'Nothing to compare.', 'summary': ''})
    raise AssertionError('unexpected call ' + workload)


def fake_ask(prompt, query, provider, workload='', **kw):
    if 'MEASURED ITEMS' not in query: return json.dumps({'findings': [], 'commentary': 'No market evidence.'}), {}
    rows = re.findall(r'^(Q\d+) \| [^|]+ \| ([^|]+?) \|', query.split('MEASURED ITEMS IN THIS PART')[1].split('RATE LIBRARY')[0], re.M)
    SEARCHES.append({'refs': [r for r, _ in rows], 'query': query, 'prompt': prompt})
    out = []
    for r, d in rows:
        if d == 'Wall item 1': out.append({'ref': r, 'source': 'published', 'rate': 50, 'unit': 'm2', 'source_url': 'https://fictional-prices.example/walls', 'source_date': '2026-06'})
        if d == 'Wall item 2': out.append({'ref': r, 'source': 'published', 'rate': 40, 'unit': 'm2', 'source_url': 'https://fictional-prices.example/walls', 'source_date': '2026-06'})
        if d in PS: out.append({'ref': r, 'source': 'provisional', 'provisional': PS[d]})
    return json.dumps({'rates': out, 'searches': ['typical costs']}), dict(PAGES)


assistants._call = fake_call
org_research._ask = fake_ask


def start(title='FICTIONAL hall', **kw):
    body = {'job_type': 'cost-estimate', 'title': title, 'brief': 'A fictional single-storey hall with a new drainage connection, to be estimated.',
            'uploads': [{'name': DOC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1: walls. Page 2: connect to the existing sewer.'}], **kw}
    r = cl.post(f'/admin/api/teams/{TID}/jobs', json=body, headers=H)
    assert r.status_code == 200, r.text
    return r.json()


def items_of(jid): return {i['ref']: i for i in teams._row(jid)['outputs']['price']['items']}
def ref_of(jid, desc): return next(i['ref'] for i in teams._row(jid)['outputs']['price']['items'] if i['description'] == desc)
SEWER, MANHOLE = 'Connection to the existing sewer', 'Existing manhole'
PS_MANHOLE = {'sum': 900, 'low': 600, 'high': 1200, 'reasoning': 'Adjusting an existing manhole typically costs £600 to £1,200.',
              'sources': [{'source_url': 'https://fictional-costs.example/drainage', 'source_date': '2026-05'}]}
def page(jid): return cl.get(f'/admin/api/teams/jobs/{jid}/page').json()


# ---------------- unmeasured work, priced as a provisional sum ----------------
j = start()
meas = {i['ref']: i for i in teams._row(j['id'])['outputs']['measure']['items']}
t('the Measurement Surveyor is told unmeasurable work can be listed for a provisional sum', any(team_qs.UNMEASURED_NOTE in c_['payload'] for c_ in CALLS if 'Measurement' in c_['workload']))
t('…unmeasured work becomes one item, quantity 1, unit item, flagged', meas['Q3']['quantity'] == 1 and meas['Q3']['unit'] == 'item' and meas['Q3'].get('unmeasured') is True)
q = SEARCHES[-1]['query']
t('the Cost Surveyor is given the location, the date and that provisional sums are allowed, with the unmeasured item marked',
  'LOCATION: Perth' in q and re.search(r'DATE: \d{4}-\d{2}-\d{2}', q) and '5. provisional: allowed' in q and 'Connection to the existing sewer | 1.0 | item | UNMEASURED' in q
  and 'provisional:' in SEARCHES[-1]['prompt'])
it = items_of(j['id'])
pv = it['Q3'].get('provisional') or {}
t('a provisional sum carries its range, reasoning and the pages cited with their dates (badge PS)', it['Q3']['rate_source'] == 'provisional' and it['Q3']['rate'] == 2500.0
  and pv['low'] == 1800.0 and pv['high'] == 3200.0 and pv['reasoning'].startswith('Typical connections') and it['Q3']['quantity'] == 1 and it['Q3']['unit'] == 'item'
  and [x['source_url'] for x in pv['sources']] == ['https://fictional-costs.example/sewer'] and pv['sources'][0]['source_date'] == '2026-07' and pv['sources'][0]['cost'] == 2400.0)
t('…a page the search did not return is dropped', pv['sources_dropped'] == 1)
t('an item no allowed source priced stays unpriced (never a provisional sum by itself)', it['Q4']['rate_source'] == 'unpriced')

# ---------------- the arithmetic, in code ----------------
cp = team_qs.final_plan(teams._row(j['id'])['outputs'], team_qs.SETTINGS)
works = Decimal('500') + Decimal('800')
cons = works + Decimal('2500')
pre = (cons * Decimal('12') / 100).quantize(Decimal('0.01'))
con = ((cons + pre) * Decimal('10') / 100).quantize(Decimal('0.01'))
fee = ((cons + pre + con) * Decimal('10') / 100).quantize(Decimal('0.01'))
t('provisional sums have their own subtotal, worked out in code, and are included in the total', cp['works'] == 1300.0 and cp['provisional_total'] == 2500.0
  and cp['construction'] == 3800.0 and cp['total'] == float(cons + pre + con + fee) and [p['ref'] for p in cp['provisional']] == ['Q3']
  and cp['elements'] == [{'element': 'Walls', 'subtotal': 1300.0}])
cf = team_qs.compute([{'ref': 'A', 'element': 'Walls', 'description': 'x', 'quantity': 10, 'unit': 'm2', 'rate': 10, 'rate_source': 'web'},
                      {'ref': 'B', 'element': 'Drainage', 'description': 'y', 'quantity': 1, 'unit': 'item', 'rate': 1000, 'rate_source': 'provisional',
                       'provisional': {'low': 900, 'high': 1100}}], 1.1, team_qs.SETTINGS, 10)
t('…a regional factor and a market adjustment apply to the measured works, never to a provisional sum', cf['works'] == 110.0 and cf['provisional_total'] == 1000.0
  and cf['market_adjustment'] == 11.0 and cf['construction'] == 1110.0)
pg = page(j['id'])
pl = pg['view']['plan']
t('the job page shows the PS badge, the Provisional sums section with its subtotal, and the construction split', pl['labels']['provisional'] == 'PS'
  and pl['counts']['provisional'] == 1 and pl['provisional'][0]['ref'] == 'Q3' and pl['provisional_total'] == 2500.0 and 'provisional sum' in pl['provisional_line'])

# ---------------- clearer decisions: nothing left unpriced without your say ----------------
dec = pg['view']['decision']
t('at sign-off the panel offers both asks and the PS switch', dec['state'] == 'signoff' and dec['can_estimate'] and dec['ps_on'] and dec['ps_rule_on'] and dec['ps_setting'] is None
  and [i['ref'] for i in dec['items']] == ['Q4'])
r = cl.post(f'/admin/api/teams/jobs/{j["id"]}/rates', json={'entries': [{'ref': 'Q4'}]}, headers=H)
t('"Use these" with no rate and nothing ticked asks what to do, naming the item and the choices', r.status_code == 400 and 'Q4' in r.json()['detail']
  and 'provisional sums' in r.json()['detail'] and items_of(j['id'])['Q4']['rate_source'] == 'unpriced' and not items_of(j['id'])['Q4'].get('decision'))
r = cl.post(f'/admin/api/teams/jobs/{j["id"]}/rates', json={'entries': [{'ref': 'Q4', 'exclude': ''}], 'go_on': True}, headers=H)
t('…an exclusion needs a short reason', r.status_code == 400)
with s.acting('Stefan'):
    r = cl.post(f'/admin/api/teams/jobs/{j["id"]}/rates', json={'entries': [{'ref': 'Q4', 'exclude': 'existing manhole, not new work'}]}, headers=H)
it = items_of(j['id'])
asm = teams._row(j['id'])['outputs']['assemble']
t('Exclude: not in scope records the item with the reason and who decided', r.status_code == 200 and it['Q4']['rate_source'] == 'excluded' and it['Q4']['rate'] is None
  and it['Q4']['exclusion']['reason'] == 'existing manhole, not new work' and it['Q4']['exclusion']['by'] == 'Stefan')
t('…it is listed under exclusions, not assumptions', any(x.startswith('Q4 Existing manhole: not in scope (existing manhole, not new work)') for x in asm['exclusions'])
  and not any(a.startswith('Q4 ') for a in asm['assumptions']))
pg = page(j['id'])
t('…and the total appears, without it', pg['view']['plan']['totals'] and pg['view']['plan']['undecided'] == 0 and pg['view']['plan']['totals']['total'] == cp['total']
  and pg['view']['plan']['exclusions'][0].startswith('Q4 '))
row = teams._row(j['id'])
md, _ = team_qs._md(row, row['outputs'])
t('the Word cost plan has a Provisional sums section with its subtotal, and the exclusion under Exclusions', '# Provisional sums' in md
  and '| **Provisional sums subtotal** | | | | **£2,500.00** |' in md and 'https://fictional-costs.example/sewer' in md and '£1,800.00 to £3,200.00' in md
  and md.index('Q4 Existing manhole: not in scope') > md.index('# Exclusions') and '| Provisional sums | £2,500.00 |' in md)
fin = team_qs.finish(row, teams.get(TID), teams._job_team(row)[1])
xl = next(d_ for d_ in fin['documents'] if d_['kind'] == 'Excel')
wb = openpyxl.load_workbook(io.BytesIO(documents.get(xl['id'])['data']))
psr = [list(r_) for r_ in wb['Provisional sums'].iter_rows(values_only=True)]
summ = {r_[0]: r_[1] for r_ in wb['Summary'].iter_rows(values_only=True)}
asr = [r_ for r_ in wb['Assumptions'].iter_rows(values_only=True)]
t('the Excel workbook has a Provisional sums sheet with the range, the sources and a subtotal', psr[1][0] == 'Q3' and psr[1][3] == 1800 and psr[1][4] == 3200
  and psr[1][5] == 2500 and 'fictional-costs.example/sewer' in psr[1][7] and psr[-1][0] == 'Provisional sums subtotal' and psr[-1][5] == 2500)
t('…and its summary splits measured works and provisional sums, both in the total', summ['Measured works subtotal'] == 1300 and summ['Provisional sums'] == 2500
  and summ['Construction subtotal'] == 3800 and summ['Total excluding VAT'] == cp['total'])
t('…the exclusion is listed as an exclusion', any(r_[0] == 'Exclusion' and str(r_[1]).startswith('Q4 ') for r_ in asr) and not any(r_[0] == 'Assumption' and str(r_[1]).startswith('Q4 ') for r_ in asr)
  and not any(r_[0] == 'Q3' or r_[0] == 'Q4' for r_ in wb['Cost plan'].iter_rows(values_only=True)))
note_title, note = team_qs.summary_note(row, 'Provisional sums test team')
t('the filed summary names the provisional sums and the exclusion; a lump sum is never filed as a past rate', 'Provisional sums: Q3' in note and 'Excluded (not in scope): Q4' in note
  and not re.search(r'^RATE \| Q3 ', note, re.M))
RATES_IN = [x for x in team_qs.history({'id': 'x', 'client': '', 'team_id': TID, 'job_type': 'cost-estimate', 'team_version': 1}, {'provider': 'claude_sonnet', 'role': 'Market Trends QS'})]
t('…a provisional sum is never counted as a past rate', not any(h_['unit'] == 'item' and h_['rate'] == 2500 for h_ in RATES_IN))

# re-pricing keeps the exclusion
r = cl.post(f'/admin/api/teams/jobs/{j["id"]}/reprice', json={'refs': ['Q1']}, headers=H)
t('a later re-price keeps the exclusion and never prices it again', r.status_code == 200 and items_of(j['id'])['Q4']['rate_source'] == 'excluded'
  and 'Q4' not in SEARCHES[-1]['refs'])

# ---------------- switched off per job ----------------
j2 = start('FICTIONAL hall, no PS', provisional=False)
it = items_of(j2['id'])
mh = ref_of(j2['id'], MANHOLE)
t('switched off on the Start screen: unmeasurable work is not listed for a provisional sum, and none is used on that job',
  SEWER not in [i['description'] for i in it.values()] and it[mh]['rate_source'] == 'unpriced' and not team_qs.ps_allowed(teams._row(j2['id'])['outputs']))
t('…the Cost Surveyor is told so', '5. provisional: not allowed on this job' in SEARCHES[-1]['query'])
PS[MANHOLE] = PS_MANHOLE
r = cl.post(f'/admin/api/teams/jobs/{j2["id"]}/reprice', json={'refs': [mh]}, headers=H)
t('…a provisional sum it offers is refused, and the refusal says why', items_of(j2['id'])[mh]['rate_source'] == 'unpriced'
  and any(x['ref'] == mh and x['source'] == 'provisional' and 'switched off on this job' in x['reason'] for x in teams._row(j2['id'])['outputs']['price']['refused_rates']))
dec = page(j2['id'])['view']['decision']
t('the decision panel shows the job\'s own setting', dec['ps_on'] is False and dec['ps_setting'] is False and dec['ps_rule_on'] is True)
r = cl.post(f'/admin/api/teams/jobs/{j2["id"]}/provisional', json={'refs': [mh], 'note': 'Use a provisional sum for the manhole.'}, headers=H)
pj = page(j2['id'])
it = items_of(j2['id'])
t('"Ask the team for provisional sums" works even with the job switched off, for the items chosen; only those are priced again', r.status_code == 200
  and SEARCHES[-1]['refs'] == [mh] and it[mh]['rate_source'] == 'provisional' and it[mh]['rate'] == 900.0 and f'{mh} | Drainage | Existing manhole | 1.0 | nr | PROVISIONAL SUM ALLOWED' in SEARCHES[-1]['query'])
with s.db() as c: logged = c.execute("SELECT count(*) FROM activity WHERE action='team_provisional_asked' AND target=?", (j2['id'],)).fetchone()[0]
t('…recorded with who asked', logged == 1 and any(x['kind'] == 'provisional' for x in pj['steps']))
r = cl.put(f'/admin/api/teams/jobs/{j2["id"]}/provisional', json={'on': True}, headers=H)
with s.db() as c: logged = c.execute("SELECT count(*) FROM activity WHERE action='team_provisional_switched' AND target=?", (j2['id'],)).fetchone()[0]
t('the switch on the decision panel turns them on for this job (logged)', r.status_code == 200 and team_qs.ps_allowed(teams._row(j2['id'])['outputs']) and logged == 1)
del PS[MANHOLE]

# Re-price can allow provisional sums for its re-run only
j3 = start('FICTIONAL hall, re-price', provisional=False)
PS[MANHOLE] = PS_MANHOLE
mh = ref_of(j3['id'], MANHOLE)
r = cl.post(f'/admin/api/teams/jobs/{j3["id"]}/reprice', json={'refs': [mh], 'provisional': True}, headers=H)
del PS[MANHOLE]
t('Re-price with "allow provisional sums" prices those items as provisional sums on that re-run', r.status_code == 200 and items_of(j3['id'])[mh]['rate_source'] == 'provisional'
  and 'provisional sums allowed for them' in teams.job_versions(j3['id'])['versions'][-1]['what'])
t('…without changing the job\'s own switch', team_qs.ps_setting(teams._row(j3['id'])['outputs']) is False)

# the rule switched off for provisional sums
rules_engine.update_rule('rate_sources', new_params={**DEFAULT, 'allowed': ['published', 'library', 'built_up']})
j4 = start('FICTIONAL hall, rule says no')
t('with the rule not allowing them, a job without its own setting uses none', items_of(j4['id'])['Q3']['rate_source'] == 'unpriced')
j5 = start('FICTIONAL hall, allowed here', provisional=True)
t('…and the Start screen can allow them for one job (as estimates)', items_of(j5['id'])['Q3']['rate_source'] == 'provisional')
rules_engine.update_rule('rate_sources', new_params=DEFAULT)

# a provisional sum that does not hold up is refused
PS[SEWER] = {**PS[SEWER], 'sum': 5000}
j6 = start('FICTIONAL hall, bad PS')
t('a sum outside the range found is refused, with the reason', items_of(j6['id'])['Q3']['rate_source'] == 'unpriced'
  and any(x['ref'] == 'Q3' and x['reason'] == 'the sum is outside the range found' for x in teams._row(j6['id'])['outputs']['price']['refused_rates']))
PS[SEWER] = {**PS[SEWER], 'sum': 2500, 'sources': [{'source_url': 'https://made-up.example/nowhere', 'source_date': '2026-01'}]}
j7 = start('FICTIONAL hall, uncited PS')
t('…as is one with no page the search returned', items_of(j7['id'])['Q3']['rate_source'] == 'unpriced'
  and any(x['ref'] == 'Q3' and 'no page the search returned' in x['reason'] for x in teams._row(j7['id'])['outputs']['price']['refused_rates']))
PS[SEWER] = {**PS[SEWER], 'sources': [{'source_url': 'https://fictional-costs.example/sewer', 'source_date': '2026-07'}]}

# ---------------- a note asking for provisional sums is acted on ----------------
j8 = start('FICTIONAL hall, note')
p = page(j8['id'])['pending'][0]
STATE['requests'].append({'provisional_sums': ['Q4'], 'estimates': [], 'reply': 'Sending the manhole to the Cost Surveyor.'})
PS[MANHOLE] = PS_MANHOLE
n = len(SEARCHES)
r = cl.post(f'/admin/api/teams/steps/{p["id"]}', json={'action': 'send_back', 'note': 'Please put a provisional sum in for the existing manhole works.'}, headers=H)
pj = page(j8['id'])
it = items_of(j8['id'])
t('a sign-off note asking for a provisional sum: the Lead QS sends the item to the Cost Surveyor, which prices only it', r.status_code == 200
  and len(SEARCHES) == n + 1 and SEARCHES[-1]['refs'] == ['Q4'] and it['Q4']['rate_source'] == 'provisional')
t('…the Cost Surveyor sees Stefan\'s note as feedback', 'Please put a provisional sum in for the existing manhole' in SEARCHES[-1]['query']
  and any(x['kind'] == 'request' for x in pj['steps']))
t('…and the job comes back for sign-off with it, not with the same question', pj['pending'][0]['kind'] == 'signoff' and not pj['view']['decision'])
del PS[MANHOLE]

j9 = start('FICTIONAL hall, note blocked', provisional=False)
p = page(j9['id'])['pending'][0]
STATE['requests'].append({'provisional_sums': ['all'], 'estimates': []})
n = len(SEARCHES)
cl.post(f'/admin/api/teams/steps/{p["id"]}', json={'action': 'send_back', 'note': 'Use provisional sums for anything you could not price.'}, headers=H)
pj = page(j9['id'])
cf = pj['view']['conflicts']
t('a note asking for what this job does not allow is explained plainly: what, which setting stops it, and where to change it', len(SEARCHES) == n and cf
  and 'switched off on this job' in cf[-1]['text'] and ref_of(j9['id'], MANHOLE) in cf[-1]['text'] and 'asked for in your note' in cf[-1]['text']
  and any(l_['label'] == 'Ask the team for provisional sums' and l_['href'].endswith('#provisional') for l_ in cf[-1]['links']))
t('…the job comes back for your decision with the explanation (no repeated question)', pj['pending'][0]['kind'] == 'signoff' and not any(x['kind'] == 'question' for x in pj['steps']))
rules_engine.update_rule('rate_sources', new_params={**DEFAULT, 'allowed': ['published', 'library', 'built_up']})
j10 = start('FICTIONAL hall, note and rule')
p = page(j10['id'])['pending'][0]
STATE['requests'].append({'provisional_sums': [], 'estimates': [ref_of(j10['id'], MANHOLE)]})
cl.post(f'/admin/api/teams/steps/{p["id"]}', json={'action': 'send_back', 'note': 'Estimate the manhole please.'}, headers=H)
cf = page(j10['id'])['view']['conflicts']
t('…when the rule stops it, the rule is named with a link to it', cf and 'the rule “Where digital teams\' rates come from” does not allow team estimates' in cf[-1]['text']
  and cf[-1]['links'][0]['href'] == '/admin/rules?rule=rate_sources#rules')
rules_engine.update_rule('rate_sources', new_params=DEFAULT)

# ---------------- the buttons are shown, or why not is explained ----------------
teams.set_autonomy(TID, 'approve')
j11 = start('FICTIONAL hall, every hand-off')
for _ in range(3):                                        # plan → measure → price → assemble, approving each hand-off
    p = page(j11['id'])['pending'][0]
    if p['kind'] == 'handoff' and p['stage'] == 'assemble': break
    cl.post(f'/admin/api/teams/steps/{p["id"]}', json={'action': 'approve'}, headers=H)
pj = page(j11['id'])
dec = pj['view']['decision']
t('at a hand-off after pricing (as on the copy of J-26A165) the estimate and PS buttons are offered', pj['pending'][0]['stage'] == 'assemble' and dec['state'] == 'later' and dec['can_estimate'])
r = cl.post(f'/admin/api/teams/jobs/{j11["id"]}/estimate', json={'refs': [ref_of(j11['id'], MANHOLE)]}, headers=H)
t('…and asking from there works', r.status_code == 200 and SEARCHES[-1]['refs'] == [ref_of(j11['id'], MANHOLE)])
cp_ = cl.post(f'/admin/api/teams/jobs/{j["id"]}/copy', json={'title': 'FICTIONAL hall copy'}, headers=H).json()
dec = page(cp_['id'])['view']['decision']
t('a copy of a job offers them too', dec and dec['can_estimate'])
with s.db() as c:                                        # a member's question is waiting: the panel says why the asks are not offered yet
    c.execute("UPDATE team_steps SET kind='question', content=? WHERE id=?", (json.dumps({'questions': ['Which drainage?'], 'role': 'Lead QS'}), page(j11['id'])['pending'][0]['id']))
dec = page(j11['id'])['view']['decision']
t('…when they cannot be offered, the panel says why and links to the rule', dec and not dec['can_estimate'] and 'question first' in dec['cannot_ask_why']
  and dec['rule_href'] == '/admin/rules?rule=rate_sources#rules')
html = cl.get(f'/admin/teams/{TID}/jobs/{j["id"]}').text
js = '\n;\n'.join(re.findall(r'<script>(.*?)</script>', html, re.S))
t('the page offers both buttons, the PS switch, Exclude: not in scope and asks what to do about undecided items',
  "'Ask the team to estimate these'" in js and "'Ask the team for provisional sums'" in js and "psb.id='provisional'" in js and "'Provisional sums on this job'" in js
  and "'Exclude: not in scope'" in js and 'function askWhat' in js and "'Provisional sums'" in js)
start_html = cl.get(f'/admin/teams/{TID}/start').text
t('the Start screen has the per-job tick', "'Allow provisional sums on this job'" in start_html)
node = shutil.which('node')
if node:
    f = Path(tempfile.mkdtemp()) / 'page.js'; f.write_text(js, encoding='utf-8')
    res = subprocess.run([node, '--check', str(f)], capture_output=True, text=True, timeout=60)
    t('the page script parses (node --check)', res.returncode == 0 or print(res.stderr[:800]))
teams.set_autonomy(TID, 'signoff')

# ---------------- filled pricing templates ----------------
def tpl(ps_section):
    """FICTIONAL: headings for Walls (and a Provisional sums section when asked), amount formulas, totals."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Cost plan'
    ws.append(['Ref', 'Description', 'Qty', 'Unit', 'Rate £', 'Amount £', 'Notes'])
    ws['B2'] = 'Walls'; ws['B2'].font = Font(bold=True)
    for r in (3, 4): ws[f'F{r}'] = f'=C{r}*E{r}'
    ws['B5'] = 'Walls total'; ws['F5'] = '=SUM(F3:F4)'
    if ps_section:
        ws['B6'] = 'Provisional sums'; ws['B6'].font = Font(bold=True)
        ws['F7'] = '=C7*E7'
        ws['B8'] = 'Provisional sums total'; ws['F8'] = '=SUM(F7:F7)'
    b = io.BytesIO(); wb.save(b); return b.getvalue()


(TPL / 'FICTIONAL with PS.xlsx').write_bytes(tpl(True))
(TPL / 'FICTIONAL without PS.xlsx').write_bytes(tpl(False))
WITH, WITHOUT = 'SharePoint/Pricing templates/FICTIONAL with PS.xlsx', 'SharePoint/Pricing templates/FICTIONAL without PS.xlsx'
cl.put(f'/admin/api/teams/{TID}/pricing-templates', json={'folder': 'SharePoint/Pricing templates'}, headers=H)
for p_ in (WITH, WITHOUT): PT.confirm(p_, PT.mapping(p_, TID)['mapping'])
t('a template\'s own Provisional sums heading is found as a section', 'Provisional sums' in [e['name'] for e in PT.confirmed(WITH)['sheets'][0]['elements']])
r = cl.put(f'/admin/api/teams/jobs/{j["id"]}/template', json={'path': WITH}, headers=H)
fl = PT.fills(j['id'])[-1]
ws = openpyxl.load_workbook(io.BytesIO(documents.get(fl['doc_id'])['data']))['Cost plan']
cells = {r_[1]: (n_, r_) for n_, r_ in enumerate(ws.iter_rows(values_only=True), 1) if r_[1]}
t('a provisional sum goes into the template\'s own Provisional sums section, as one item with its sum and PS source', r.status_code == 200
  and cells['Connection to the existing sewer'][0] == 7 and cells['Connection to the existing sewer'][1][2] == 1 and cells['Connection to the existing sewer'][1][4] == 2500
  and cells['Connection to the existing sewer'][1][6].startswith('PS: provisional sum (range £1,800.00 to £3,200.00)'))
t('…an excluded item is not put in the rows', 'Existing manhole' not in cells)
af = {r_[0]: r_[1] for r_ in openpyxl.load_workbook(io.BytesIO(documents.get(fl['doc_id'])['data']))["Alice's figures"].iter_rows(values_only=True) if r_ and r_[0]}
t('…Alice\'s figures list the provisional sums with their subtotal, and the exclusion', af['Provisional sums subtotal'] == 2500 and af['Measured works'] == 1300
  and str(af['Excluded: not in scope']).startswith('Q4 Existing manhole'))
t('…and the template\'s formulas agree with them', not fl['differences'])
r = cl.put(f'/admin/api/teams/jobs/{j["id"]}/template', json={'path': WITHOUT}, headers=H)
fl = PT.fills(j['id'])[-1]
ws = openpyxl.load_workbook(io.BytesIO(documents.get(fl['doc_id'])['data']))['Cost plan']
rows = [r_ for r_ in ws.iter_rows(values_only=True)]
heads = [n_ for n_, r_ in enumerate(rows) if r_[1] == 'Provisional sums']
t('a template with no PS section gets them after its rows under "Provisional sums", and the difference is listed', heads and rows[heads[0] + 1][1] == 'Connection to the existing sewer'
  and any(d_['where'] == 'Provisional sums' for d_ in fl['differences']))
