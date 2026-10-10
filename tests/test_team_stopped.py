"""Stopped jobs, the job page's pricing template card, removing templates from a team's list, and macro-enabled templates
(Stefan, 9 Oct 2026). Fictional templates only, made here. No real model is called.

A stopped job still allows changing and refilling its pricing template, Re-price, Re-measure, downloads and Copy as a new job; only
carrying on needs Resume, which makes a new version and picks up exactly where it stopped. A job from before versions is v1.
"Remove from this list" hides a template from one team's list and pickers without touching the file or its mapping; defaults fall
back. A macro-enabled template's filled copy keeps its macros."""
import _util  # first: throwaway data folder, dummy keys, no real model calls
from _util import t
import io, json, re, shutil, subprocess, tempfile, zipfile
from pathlib import Path
import openpyxl
from openpyxl.styles import Font
import substrate_store as s
s.init()
import doc_library as L
root = Path(tempfile.mkdtemp(prefix='alice-docs-')) / 'Documents'
L.ROOT = root.resolve()
L.add_source('SharePoint', 'sharepoint', 'SharePoint: QS site')
(root / 'SharePoint' / 'Pricing templates').mkdir()
(root / 'SharePoint' / 'Filled').mkdir()
import app, assistants, org_research, organisations, documents
import teams, team_qs, pricing_templates as PT
from fastapi.testclient import TestClient
cl = TestClient(app.app); H = {'x-admin-token': app.ADMIN_TOKEN}
teams.BACKGROUND = False
TPL = root / 'SharePoint' / 'Pricing templates'


def xlsx(wb):
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def template_headings():
    """FICTIONAL: one sheet, section headings, amount formulas, element and grand totals."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Cost plan'
    ws['A1'] = 'FICTIONAL pricing template'
    ws.append([]); ws.append(['Ref', 'Description', 'Qty', 'Unit', 'Rate £', 'Amount £', 'Notes'])
    ws['B4'] = 'Walls'; ws['B4'].font = Font(bold=True)
    for r in (5, 6): ws[f'F{r}'] = f'=C{r}*E{r}'
    ws['B7'] = 'Walls total'; ws['F7'] = '=SUM(F5:F6)'
    ws['B8'] = 'Roof'; ws['B8'].font = Font(bold=True)
    ws['F9'] = '=C9*E9'
    ws['B10'] = 'Roof total'; ws['F10'] = '=SUM(F9:F9)'
    ws['B12'] = 'Grand total'; ws['F12'] = '=F7+F10'
    return xlsx(wb)


VBA = b'FICTIONAL VBA project: Sub Recalc() ... End Sub ' * 20


def macro_enabled(raw):
    """FICTIONAL: the same template saved as a macro-enabled workbook (.xlsm) holding a VBA project."""
    src = zipfile.ZipFile(io.BytesIO(raw)); out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in src.namelist():
            d = src.read(n)
            if n == '[Content_Types].xml':
                d = d.decode().replace('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml', 'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
                d = d.replace('</Types>', '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>').encode()
            if n == 'xl/_rels/workbook.xml.rels':
                d = d.decode().replace('</Relationships>', '<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
                                                           'Target="vbaProject.bin"/></Relationships>').encode()
            z.writestr(n, d)
        z.writestr('xl/vbaProject.bin', VBA)
    return out.getvalue()


(TPL / 'FICTIONAL shared cost plan.xlsx').write_bytes(template_headings())
(TPL / 'FICTIONAL second plan.xlsx').write_bytes(template_headings())
(TPL / 'FICTIONAL macro plan.xlsm').write_bytes(macro_enabled(template_headings()))
SHARED = 'SharePoint/Pricing templates/FICTIONAL shared cost plan.xlsx'
SECOND = 'SharePoint/Pricing templates/FICTIONAL second plan.xlsx'
MACRO = 'SharePoint/Pricing templates/FICTIONAL macro plan.xlsm'

team = teams.from_template('quantity-surveying', name='Stopped jobs test team')
TID = team['id']
DOC = 'FICTIONAL hall spec.md'
PLAN = {'plan': 'Measure.', 'elements': [{'name': 'Walls', 'documents': [DOC]}, {'name': 'Roof', 'documents': [DOC]}],
        'documents': [{'name': DOC, 'use': 'all'}], 'location': '', 'summary': 'Planned.', 'note': 'Go.', 'questions': []}
ITEMS = [{'element': 'Walls', 'description': f'Wall item {n}', 'quantity': 10 * n, 'unit': 'm2', 'source': {'document': DOC, 'page': '1'}} for n in (1, 2)] + [
    {'element': 'Roof', 'description': 'Roof covering', 'quantity': 120, 'unit': 'm2', 'source': {'document': DOC, 'page': '2'}}]
CALLS = []
RATES = {'Wall item 1': 50, 'Wall item 2': 40, 'Roof covering': 30}


def fake_call(provider, system, messages, max_tokens=1500, timeout=60, workload='', meta=None):
    payload = messages[0]['content']
    CALLS.append(workload)
    if 'Lead QS' in workload and 'COST PLAN FIGURES' in payload:
        return json.dumps({'accept': True, 'summary': 'A hall.', 'assumptions': [], 'exclusions': ['VAT'], 'risks': [], 'note': 'On.'})
    if 'Lead QS' in workload: return json.dumps(PLAN)
    if 'Measurement' in workload:
        part = payload.split('THIS PART: take off ONLY these elements:')[1].split('from the DOCUMENTS')[0]
        return json.dumps({'accept': True, 'items': [x for x in ITEMS if x['element'] in part], 'summary': '', 'note': 'Price these.'})
    if 'Market Trends' in workload: return json.dumps({'matches': [], 'commentary': 'Nothing to compare.', 'summary': ''})
    raise AssertionError('unexpected call ' + workload)


def fake_ask(prompt, query, provider, workload='', **kw):
    CALLS.append('search')
    if 'MEASURED ITEMS' not in query: return json.dumps({'findings': [], 'commentary': 'No market evidence.'}), {}
    rows = re.findall(r'^(Q\d+) \| [^|]+ \| ([^|]+?) \|', query.split('MEASURED ITEMS IN THIS PART')[1].split('RATE LIBRARY')[0], re.M)
    out = [{'ref': r, 'source': 'published', 'rate': RATES[d], 'unit': 'm2', 'source_url': 'https://fictional-prices.example/p',
            'source_title': 'FICTIONAL price book', 'source_date': '2026-06'} for r, d in rows if d in RATES]
    return json.dumps({'rates': out, 'searches': ['rates']}), {'https://fictional-prices.example/p': 'FICTIONAL price book'}


assistants._call = fake_call
org_research._ask = fake_ask

cl.put(f'/admin/api/teams/{TID}/pricing-templates', json={'folder': 'SharePoint/Pricing templates', 'outputs': 'SharePoint/Filled', 'default': SHARED}, headers=H)
for p_ in (SHARED, SECOND, MACRO): PT.confirm(p_, PT.mapping(p_, TID)['mapping'])


def start(title, autonomy='signoff', **kw):
    r = cl.post(f'/admin/api/teams/{TID}/jobs', json={'job_type': 'cost-estimate', 'title': title, 'autonomy': autonomy,
                'brief': 'A fictional single-storey hall, to be estimated in full.',
                'uploads': [{'name': DOC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1 walls, page 2 roof.'}], **kw}, headers=H)
    assert r.status_code == 200, r.text
    return r.json()


def approve_pending(jid, upto=20):
    for _ in range(upto):
        p = [x for x in teams.job_detail(jid)['pending'] if x['kind'] == 'handoff']
        if not p: return
        cl.post(f'/admin/api/teams/steps/{p[0]["id"]}', json={'action': 'approve'}, headers=H)


def stop(jid):
    r = cl.post(f'/admin/api/teams/jobs/{jid}/stop', headers=H)
    assert r.status_code == 200, r.text
    return r.json()


# ---------------- a job stopped at sign-off (like J-26A165) ----------------
j = start('FICTIONAL hall A')
t('the job runs to sign-off', j['status'] == 'waiting' and j['pending'][0]['kind'] == 'signoff')
stop(j['id'])
pg = cl.get(f'/admin/api/teams/jobs/{j["id"]}/page').json()
t('a stopped job offers Re-price, Re-measure, Resume and Copy as a new job', pg['status'] == 'stopped'
  and pg['can'] == {'reprice': True, 'remeasure': True, 'resume': True, 'copy': True, 'add_files': True, 'decide_flags': True})
n = len(CALLS)
fills0 = len(PT.fills(j['id']))
r = cl.put(f'/admin/api/teams/jobs/{j["id"]}/template', json={'path': SECOND}, headers=H)
t('a stopped job may change its pricing template ("Use this template"), and it is filled again from the same items', r.status_code == 200
  and teams._row(j['id'])['pricing_template'] == SECOND and len(PT.fills(j['id'])) == fills0 + 1 and PT.fills(j['id'])[-1]['template'] == SECOND)
r = cl.post(f'/admin/api/teams/jobs/{j["id"]}/template/fill', headers=H)
t('…and Fill the template again works on it too', r.status_code == 200 and r.json()['pricing_template']['fill']['status'] == 'filled' and len(PT.fills(j['id'])) == fills0 + 2)
t('…without any member working again (no model call) and the job stays stopped', len(CALLS) == n and teams._row(j['id'])['status'] == 'stopped')
docs = teams._row(j['id'])['outputs']['documents']
ok = all(cl.get(f'/documents/{d_["id"]}/download').status_code == 200 for d_ in docs)
t('a stopped job\'s documents still download (Word, Excel and the filled template)', len(docs) >= 3 and ok)
cp = cl.post(f'/admin/api/teams/jobs/{j["id"]}/copy', json={'title': 'FICTIONAL hall A again'}, headers=H)
t('Copy as a new job works on a stopped job', cp.status_code == 200 and cp.json()['id'] != j['id'] and cp.json()['title'] == 'FICTIONAL hall A again')
RATES['Wall item 2'] = 44
r = cl.post(f'/admin/api/teams/jobs/{j["id"]}/reprice', json={'refs': ['Q2']}, headers=H)
row = teams._row(j['id'])
t('Re-price works on a stopped job, as a new version, and comes back for sign-off', r.status_code == 200 and row['version'] == 2 and row['status'] == 'waiting'
  and next(i for i in row['outputs']['price']['items'] if i['ref'] == 'Q2')['rate'] == 44)
t('…the stopped version stays readable', cl.get(f'/admin/api/teams/jobs/{j["id"]}/versions/1').status_code == 200)

j2 = start('FICTIONAL hall B')
stop(j2['id'])
r = cl.post(f'/admin/api/teams/jobs/{j2["id"]}/remeasure', json={'elements': ['Roof']}, headers=H)
row = teams._row(j2['id'])
t('Re-measure works on a stopped job, as a new version', r.status_code == 200 and row['version'] == 2 and row['status'] == 'waiting'
  and [v['kind'] for v in teams.job_versions(j2['id'])['versions']] == ['first', 'remeasure'])

# ---------------- Resume: a new version that carries on where it stopped ----------------
j3 = start('FICTIONAL hall C')
stop(j3['id'])
before = {d_['id'] for d_ in teams._row(j3['id'])['outputs']['documents']}
n = len(CALLS)
r = cl.post(f'/admin/api/teams/jobs/{j3["id"]}/resume', json={'note': 'Carry on, please.'}, headers=H)
row = teams._row(j3['id'])
d3 = r.json()
vs = teams.job_versions(j3['id'])['versions']
t('Resume on a stopped job makes a new version (Resumed) with your note', r.status_code == 200 and row['version'] == 2
  and [(v['version'], v['kind_label']) for v in vs] == [(1, 'First run'), (2, 'Resumed')] and vs[1]['note'] == 'Carry on, please.')
t('…it carries on where it stopped: your sign-off waits again, no member works again', d3['status'] == 'waiting'
  and [p['kind'] for p in d3['pending']] == ['signoff'] and len(CALLS) == n)
newdocs = row['outputs']['documents']
t('…and the documents are rebuilt for the new version', not ({d_['id'] for d_ in newdocs} & before) and any('v2' in d_['name'] for d_ in newdocs))
t('…the first version is kept readable', cl.get(f'/admin/api/teams/jobs/{j3["id"]}/versions/1').status_code == 200)
with s.db() as c:
    act = c.execute("SELECT actor, detail FROM activity WHERE action='team_job_resumed' AND target=?", (j3['id'],)).fetchone()
t('…and it is logged, with who', act and 'v2' in act['detail'] and act['actor'])
r = cl.post(f'/admin/api/teams/jobs/{j3["id"]}/resume', headers=H)
t('a job that is not stopped is not made a new version by Resume', r.status_code == 200 and teams._row(j3['id'])['version'] == 2)

j4 = start('FICTIONAL hall D', autonomy='approve')
t('a job that holds every hand-off waits after the plan', j4['pending'][0]['kind'] == 'handoff')
stop(j4['id'])
n = len(CALLS)
d4 = cl.post(f'/admin/api/teams/jobs/{j4["id"]}/resume', json={}, headers=H).json()
t('resumed at a hand-off: the same hand-off waits for you again, nothing re-run', d4['status'] == 'waiting' and [p['kind'] for p in d4['pending']] == ['handoff']
  and d4['pending'][0]['stage'] == j4['pending'][0]['stage'] and len(CALLS) == n and d4['version'] == 2)
approve_pending(j4['id'])
t('…and the job goes on from there to sign-off', teams.job_detail(j4['id'])['pending'][0]['kind'] == 'signoff')

# a job stopped after measuring, before pricing: the template is filled from the measured items
j5 = start('FICTIONAL hall E', autonomy='approve', template='')
cl.post(f'/admin/api/teams/steps/{j5["pending"][0]["id"]}', json={'action': 'approve'}, headers=H)
t('a job measured but not priced yet waits at the measuring hand-off', bool((teams._row(j5['id'])['outputs'].get('measure') or {}).get('items'))
  and not (teams._row(j5['id'])['outputs'].get('price') or {}).get('items'))
stop(j5['id'])
pg = cl.get(f'/admin/api/teams/jobs/{j5["id"]}/page').json()
t('…stopped there, it offers Re-measure and Resume, not Re-price (nothing priced)', pg['can']['remeasure'] and pg['can']['resume'] and not pg['can']['reprice'])
r = cl.put(f'/admin/api/teams/jobs/{j5["id"]}/template', json={'path': SHARED}, headers=H)
r2 = cl.post(f'/admin/api/teams/jobs/{j5["id"]}/template/fill', headers=H)
fl = PT.fills(j5['id'])
ws = openpyxl.load_workbook(io.BytesIO(documents.get(fl[-1]['doc_id'])['data']))['Cost plan']
cells = {r_[1]: r_ for r_ in ws.iter_rows(values_only=True) if r_[1]}
t('…its template can be changed and filled from the measured items: quantities in, rates left blank', r.status_code == 200 and r2.json()['pricing_template']['fill']['status'] == 'filled'
  and cells['Wall item 1'][2] == 10 and cells['Wall item 1'][4] is None and 'not priced' in cells['Wall item 1'][6])

# ---------------- a job from before versions existed is v1 ----------------
j6 = start('FICTIONAL hall F')
stop(j6['id'])
with s.db() as c:
    c.execute('DELETE FROM team_job_versions WHERE job_id=?', (j6['id'],))
pg = cl.get(f'/admin/api/teams/jobs/{j6["id"]}/page').json()
t('a job from before versions is shown as v1', pg['version'] == 1 and [v['version'] for v in pg['versions']['versions']] == [1] and pg['can']['resume'])
r = cl.post(f'/admin/api/teams/jobs/{j6["id"]}/resume', json={}, headers=H)
vs = teams.job_versions(j6['id'])['versions']
t('…and resumes as v2, keeping v1 readable', r.status_code == 200 and [(v['version'], v['kind']) for v in vs] == [(1, 'first'), (2, 'resume')]
  and cl.get(f'/admin/api/teams/jobs/{j6["id"]}/versions/1').status_code == 200)
j7 = start('FICTIONAL hall G')
stop(j7['id'])
with s.db() as c:
    c.execute('DELETE FROM team_job_versions WHERE job_id=?', (j7['id'],))
for k_, body in (('template', None), ('fill', None), ('reprice', {'refs': ['Q1']})):
    if k_ == 'template': rr = cl.put(f'/admin/api/teams/jobs/{j7["id"]}/template', json={'path': SECOND}, headers=H)
    elif k_ == 'fill': rr = cl.post(f'/admin/api/teams/jobs/{j7["id"]}/template/fill', headers=H)
    else: rr = cl.post(f'/admin/api/teams/jobs/{j7["id"]}/reprice', json=body, headers=H)
    t(f'a stopped job from before versions: {k_} works', rr.status_code == 200)
t('…the re-price is v2 after a v1 made for it', [v['version'] for v in teams.job_versions(j7['id'])['versions']] == [1, 2])
t('…its v1 documents still download', all(cl.get(f'/documents/{d_["id"]}/download').status_code == 200
                                        for d_ in teams.version_view(j7['id'], 1)['documents']) and teams.version_view(j7['id'], 1)['documents'])
j8 = start('FICTIONAL hall I')
stop(j8['id'])
with s.db() as c:
    c.execute('DELETE FROM team_job_versions WHERE job_id=?', (j8['id'],))
rr = cl.post(f'/admin/api/teams/jobs/{j8["id"]}/remeasure', json={'elements': ['Walls']}, headers=H)
t('a stopped job from before versions: remeasure works, as v2', rr.status_code == 200 and teams._row(j8['id'])['version'] == 2
  and [v['kind'] for v in teams.job_versions(j8['id'])['versions']] == ['first', 'remeasure'])
t('…and Copy as a new job works on it', cl.post(f'/admin/api/teams/jobs/{j7["id"]}/copy', json={}, headers=H).status_code == 200)

# ---------------- the job page: Re-price, Re-measure and Resume in the header; no "null" ----------------
html = cl.get(f'/admin/teams/{TID}/jobs/{j["id"]}').text
js = '\n;\n'.join(re.findall(r'<script>(.*?)</script>', html, re.S))
t('the header offers Resume next to Re-price, Re-measure and Copy as a new job', "d.can.resume" in js and "toggle('copy','Copy as a new job')" in js
  and js.index("d.can.resume") < js.index("toggle('copy','Copy as a new job')"))


def bare_nulls(src):
    """Native .append(...) calls given a null argument: the browser prints it as the text "null"."""
    out = []
    for m in re.finditer(r'\.append\(', src):
        i, d, cur, args, q = m.end(), 1, '', [], None
        while d and i < len(src):
            ch = src[i]
            if q:
                cur += ch
                if ch == '\\': cur += src[i + 1]; i += 1
                elif ch == q: q = None
            elif ch in '\'"`': q = ch; cur += ch
            elif ch in '([{': d += 1; cur += ch
            elif ch in ')]}':
                d -= 1
                if d: cur += ch
            elif ch == ',' and d == 1: args.append(cur); cur = ''
            else: cur += ch
            i += 1
        args.append(cur)
        out += [a.strip()[:80] for a in args if re.search(r':\s*null\s*$', a.strip()) or a.strip() == 'null']
    return out


import teams_ui
t('no "null" on the page: nothing hands null to the browser\'s own append (the Pricing template card under Alice\'s own layout)',
  bare_nulls(teams_ui.SCRIPT) == [] and bare_nulls(js) == [])
t('…Alice\'s own layout gets a plain description instead', 'Alice’s Word cost plan and Excel workbook, without a template of yours.' in js)
t('a job with Alice\'s own layout says so, with nothing from a template', PT.for_job({'pricing_template': ''}) == {'path': '', 'name': "Alice's own layout", 'own': True, 'from': ''})
node = shutil.which('node')
if node:
    f = Path(tempfile.mkdtemp()) / 'page.js'; f.write_text(js, encoding='utf-8')
    res = subprocess.run([node, '--check', str(f)], capture_output=True, text=True, timeout=60)
    t('the page script parses (node --check)', res.returncode == 0 or print(res.stderr[:800]))

# ---------------- Remove from this list ----------------
organisations.create('FICTIONAL Otter Council', client=True)
cl.put(f'/admin/api/teams/{TID}/pricing-templates/client-default', json={'client': 'FICTIONAL Otter Council', 'path': SECOND}, headers=H)
used = start('FICTIONAL hall H', template=SECOND)
used_fills = PT.fills(used['id'])
raw_before, mtime = (TPL / 'FICTIONAL second plan.xlsx').read_bytes(), (TPL / 'FICTIONAL second plan.xlsx').stat().st_mtime_ns
mapping_before = PT.confirmed(SECOND)
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/hidden', json={'path': SECOND, 'hidden': True}, headers=H)
P = r.json()
t('Remove from this list hides the template from the team\'s list', r.status_code == 200 and SECOND not in [x['path'] for x in P['templates']]
  and [x['path'] for x in P['hidden']] == [SECOND] and P['hidden'][0]['hidden_by'])
t('…the file in the document source is never touched', (TPL / 'FICTIONAL second plan.xlsx').read_bytes() == raw_before
  and (TPL / 'FICTIONAL second plan.xlsx').stat().st_mtime_ns == mtime)
t('…its saved mapping is kept', PT.confirmed(SECOND) == mapping_before and PT.describe(SECOND)['mapping'] == 'confirmed')
t('…an organisation\'s default falls back to the team\'s, and the page says so', PT.default_for(teams.get(TID), 'FICTIONAL Otter Council') == (SHARED, 'team')
  and any('FICTIONAL Otter Council' in f_ and 'FICTIONAL shared cost plan.xlsx' in f_ for f_ in P['fallbacks']) and 'FICTIONAL Otter Council' in P['message'])
t('…the client\'s default in the team\'s settings is left as it is (Add back restores it)', PT.client_default(teams.get(TID), 'FICTIONAL Otter Council')[0] == SECOND)
sp = cl.get(f'/admin/api/teams/{TID}/start').json()
t('…the Start a job screen no longer offers it, nor the organisation default pointing at it', SECOND not in [x['path'] for x in sp['templates']['templates']]
  and 'FICTIONAL Otter Council' not in sp['org_defaults'])
pg = cl.get(f'/admin/api/teams/jobs/{used["id"]}/page').json()
t('…a job that used it keeps it and its filled copies', pg['pricing_template']['path'] == SECOND and PT.fills(used['id']) == used_fills
  and cl.get(f'/documents/{used_fills[-1]["doc_id"]}/download').status_code == 200)
t('…but it is not offered in another job\'s picker', SECOND not in [x['path'] for x in cl.get(f'/admin/api/teams/jobs/{j3["id"]}/page').json()['pricing_template']['choices']])
t('…choosing it for a job is refused until it is added back', cl.put(f'/admin/api/teams/jobs/{j3["id"]}/template', json={'path': SECOND}, headers=H).status_code == 400
  and cl.post(f'/admin/api/teams/{TID}/jobs', json={'job_type': 'cost-estimate', 'title': 'x', 'brief': 'A fictional hall to estimate in full.',
                                                      'uploads': [{'name': DOC, 'kind': 'spec', 'text': 'FICTIONAL.'}], 'template': SECOND}, headers=H).status_code == 400)
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/hidden', json={'path': SHARED, 'hidden': True}, headers=H)
P = r.json()
t('removing the team\'s default: new jobs fall back to Alice\'s own layout, and the page says so', r.status_code == 200 and P['default'] is None
  and PT.default_for(teams.get(TID), '') == ('', '') and PT.default_for(teams.get(TID), 'FICTIONAL Otter Council') == ('', '')
  and any("team's default" in f_ and "Alice's own layout" in f_ for f_ in P['fallbacks']))
t('…the team\'s setting itself is kept, so adding it back restores it', P['settings']['default'] == SHARED)
t('a hidden template cannot be made the team\'s default', cl.put(f'/admin/api/teams/{TID}/pricing-templates', json={'default': SECOND}, headers=H).status_code == 400)
t('something not in the team\'s list cannot be removed', cl.put(f'/admin/api/teams/{TID}/pricing-templates/hidden', json={'path': 'SharePoint/nope.xlsx', 'hidden': True}, headers=H).status_code == 400)
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/hidden', json={'path': SHARED, 'hidden': False}, headers=H)
t('Add back: it is in the list and the team\'s default again', r.status_code == 200 and SHARED in [x['path'] for x in r.json()['templates']]
  and r.json()['default']['path'] == SHARED and "team's default again" in r.json()['message'] and PT.default_for(teams.get(TID), '') == (SHARED, 'team'))
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/hidden', json={'path': SECOND, 'hidden': False}, headers=H)
t('…and an organisation\'s default is used again', r.status_code == 200 and PT.default_for(teams.get(TID), 'FICTIONAL Otter Council') == (SECOND, 'client'))
with s.db() as c:
    log = [dict(x) for x in c.execute("SELECT action, actor, target FROM activity WHERE action IN ('pricing_template_hidden','pricing_template_restored') ORDER BY id")]
t('who removed or added back what is logged', [(x['action'], x['target']) for x in log] == [('pricing_template_hidden', SECOND), ('pricing_template_hidden', SHARED),
  ('pricing_template_restored', SHARED), ('pricing_template_restored', SECOND)] and all(x['actor'] for x in log))
other = teams.from_template('quantity-surveying', name='Another team')
cl.put(f'/admin/api/teams/{other["id"]}/pricing-templates', json={'folder': 'SharePoint/Pricing templates'}, headers=H)
cl.put(f'/admin/api/teams/{TID}/pricing-templates/hidden', json={'path': SECOND, 'hidden': True}, headers=H)
t('the list is per team: another team still lists it', SECOND in [x['path'] for x in cl.get(f'/admin/api/teams/{other["id"]}/pricing-templates').json()['templates']])
html = cl.get(f'/admin/teams/{TID}').text
t('the team page offers Remove from this list, Show hidden and Add back', "'Remove from this list'" in html and "'Show hidden'" in html and "'Add back'" in html)

# ---------------- macro-enabled templates keep their macros ----------------
jm = start('FICTIONAL macro hall', template=MACRO)
fl = PT.fills(jm['id'])[-1]
doc = documents.get(fl['doc_id'])
z = zipfile.ZipFile(io.BytesIO(doc['data']))
t('a macro-enabled template is filled as a macro-enabled workbook (.xlsm)', doc['name'].endswith('.xlsm') and doc['media_type'] == 'application/vnd.ms-excel.sheet.macroEnabled.12')
t('…its macros are kept exactly', 'xl/vbaProject.bin' in z.namelist() and z.read('xl/vbaProject.bin') == VBA
  and 'macroEnabled' in z.read('[Content_Types].xml').decode() and 'vbaProject' in z.read('xl/_rels/workbook.xml.rels').decode())
wsm = openpyxl.load_workbook(io.BytesIO(doc['data']), keep_vba=True)['Cost plan']
t('…with the items filled in', any(r_[1] == 'Roof covering' and r_[4] == 30 for r_ in wsm.iter_rows(values_only=True)))
t('…saved to the outputs folder as .xlsm', fl['library_path'].endswith('.xlsm') and (root / fl['library_path']).is_file())
t('…the template file itself is unchanged', zipfile.ZipFile(TPL / 'FICTIONAL macro plan.xlsm').read('xl/vbaProject.bin') == VBA)
dl = cl.get(f'/documents/{fl["doc_id"]}/download')
t('…and it downloads as a macro-enabled workbook', dl.status_code == 200 and 'macroEnabled' in dl.headers['content-type'])
t('the chat\'s document tool still offers only Word, Excel and PDF', documents.TOOL_SCHEMA['properties']['format']['enum'] == ['docx', 'xlsx', 'pdf'])
