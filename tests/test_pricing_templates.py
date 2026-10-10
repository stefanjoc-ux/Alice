"""Pricing templates and the Start a job screen (Stefan, 8 Oct 2026). Fictional sample templates only, made here.

The file role and pickers; the layout found in code on different layouts (headings, one sheet per element, CSV), a model only when
code cannot tell (checked in code); the mapping saved, editable and asked again when the file changes; items and rates landing in the
right cells with their sources, rows added in the template's formatting, totals worked out in code and the template's own formulas
checked against them; the template never modified; shared templates across clients, a client's template refused elsewhere; filled
copies tagged to the job's client; the Ready to start checklist; typical cost hidden with no history. No real model is called."""
import _util  # first: throwaway data folder, dummy keys, no real model calls
from _util import t
import base64, io, json, os, re, tempfile, time
from pathlib import Path
import openpyxl
from openpyxl.styles import Font
import substrate_store as s
s.init()
import doc_library as L
root = Path(tempfile.mkdtemp(prefix='alice-docs-')) / 'Documents'
L.ROOT = root.resolve()
L.add_source('SharePoint', 'sharepoint', 'SharePoint: QS site')
(root / 'SharePoint' / 'Pricing templates').mkdir()
(root / 'SharePoint' / 'Filled').mkdir()
(root / 'SharePoint' / 'Drawings').mkdir()
import app, assistants, org_research, rules_engine, usage_meter, clients, organisations
import teams, team_qs, team_start, pricing_templates as PT
from fastapi.testclient import TestClient
cl = TestClient(app.app); H = {'x-admin-token': app.ADMIN_TOKEN}
teams.BACKGROUND = False
TPL = root / 'SharePoint' / 'Pricing templates'


def xlsx(wb):
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def template_headings(amount_formula='=C{r}*E{r}', with_source=True):
    """FICTIONAL: one sheet, section headings, two empty rows per element with amount formulas, element and grand totals."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Cost plan'
    ws['A1'] = 'FICTIONAL pricing template'
    ws.append([]); ws.append(['Ref', 'Description', 'Qty', 'Unit', 'Rate £', 'Amount £'] + (['Notes'] if with_source else []))
    ws['B4'] = 'Walls'; ws['B4'].font = Font(bold=True)
    for r in (5, 6): ws[f'F{r}'] = amount_formula.format(r=r)
    ws['B7'] = 'Walls total'; ws['F7'] = '=SUM(F5:F6)'
    ws['B8'] = 'Roofs'; ws['B8'].font = Font(bold=True)
    ws['F9'] = amount_formula.format(r=9)
    ws['B10'] = 'Roofs total'; ws['F10'] = '=SUM(F9:F9)'
    ws['B12'] = 'Grand total'; ws['F12'] = '=F7+F10'
    for r in (5, 6, 9): ws[f'B{r}'].font = Font(italic=True)
    return xlsx(wb)


def template_sheets():
    """FICTIONAL: one sheet per element, headers in row 1, a total at the bottom of each."""
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for name in ('Walls', 'Roof'):
        ws = wb.create_sheet(name)
        ws.append(['Item', 'Specification', 'Quantity', 'Unit', 'Unit rate', 'Total', 'Source'])
        for r in range(2, 5): ws[f'F{r}'] = f'=C{r}*E{r}'
        ws['B6'] = 'Total'; ws['F6'] = '=SUM(F2:F5)'
    rm = wb.create_sheet('Read me'); rm['A1'] = 'FICTIONAL'
    return xlsx(wb)


def template_odd():
    """FICTIONAL: headers code cannot recognise."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Pricing'
    ws.append(['Line', 'What', 'How many', 'Measure', 'Each', 'Line value'])
    return xlsx(wb)


(TPL / 'FICTIONAL shared cost plan.xlsx').write_bytes(template_headings())
(TPL / 'FICTIONAL by sheet.xlsx').write_bytes(template_sheets())
(TPL / 'FICTIONAL odd headers.xlsx').write_bytes(template_odd())
(TPL / 'FICTIONAL simple.csv').write_text('Description,Qty,Unit,Rate,Amount,Notes\n', encoding='utf-8')
(TPL / 'FICTIONAL old.xls').write_bytes(b'\xd0\xcf\x11\xe0 not really')
SHARED = 'SharePoint/Pricing templates/FICTIONAL shared cost plan.xlsx'
BYSHEET = 'SharePoint/Pricing templates/FICTIONAL by sheet.xlsx'
ODD = 'SharePoint/Pricing templates/FICTIONAL odd headers.xlsx'
CSVT = 'SharePoint/Pricing templates/FICTIONAL simple.csv'

# ---------------- the team and stand-in models ----------------
team = teams.from_template('quantity-surveying', name='Templates test team')
TID = team['id']
teams.set_autonomy(TID, 'signoff')
DOC = 'FICTIONAL hall spec.md'
PLAN = {'plan': 'Measure.', 'elements': [{'name': 'Walls', 'documents': [DOC]}, {'name': 'Roof', 'documents': [DOC]}, {'name': 'Drainage', 'documents': [DOC]}],
        'documents': [{'name': DOC, 'use': 'all'}], 'location': '', 'summary': 'Planned.', 'note': 'Go.', 'questions': []}
ITEMS = [{'element': 'Walls', 'description': f'Wall item {n}', 'quantity': 10 * n, 'unit': 'm2', 'source': {'document': DOC, 'page': '1'}} for n in range(1, 4)] + [
    {'element': 'Roof', 'description': 'Roof covering', 'quantity': 120, 'unit': 'm2', 'source': {'document': DOC, 'page': '2'}},
    {'element': 'Drainage', 'description': 'Gullies', 'quantity': 4, 'unit': 'nr', 'source': {'document': DOC, 'page': '3'}}]
CALLS, SEARCHES, MAPPER = [], [], {'reply': None}


def fake_call(provider, system, messages, max_tokens=1500, timeout=60, workload='', meta=None):
    payload = messages[0]['content']
    CALLS.append({'workload': workload, 'payload': payload, 'system': system})
    if 'template layout' in workload: return MAPPER['reply']
    if 'Lead QS' in workload and 'COST PLAN FIGURES' in payload:
        return json.dumps({'accept': True, 'summary': 'A hall.', 'assumptions': [], 'exclusions': ['VAT'], 'risks': [], 'note': 'On.'})
    if 'Lead QS' in workload: return json.dumps(PLAN)
    if 'Measurement' in workload:
        part = payload.split('THIS PART: take off ONLY these elements:')[1].split('from the DOCUMENTS')[0]
        return json.dumps({'accept': True, 'items': [x for x in ITEMS if x['element'] in part], 'summary': '', 'note': 'Price these.'})
    if 'Market Trends' in workload: return json.dumps({'matches': [], 'commentary': 'Nothing to compare.', 'summary': ''})
    raise AssertionError('unexpected call ' + workload)


RATES = {'Wall item 1': 50, 'Wall item 2': 40, 'Roof covering': 30, 'Gullies': 120}


def fake_ask(prompt, query, provider, workload='', **kw):
    if 'MEASURED ITEMS' not in query: return json.dumps({'findings': [], 'commentary': 'No market evidence.'}), {}
    rows = re.findall(r'^(Q\d+) \| [^|]+ \| ([^|]+?) \|', query.split('MEASURED ITEMS IN THIS PART')[1].split('RATE LIBRARY')[0], re.M)
    SEARCHES.append({'refs': [r for r, _ in rows], 'query': query})
    out = [{'ref': r, 'source': 'published', 'rate': RATES[d], 'unit': 'nr' if d == 'Gullies' else 'm2', 'source_url': 'https://fictional-prices.example/p',
            'source_title': 'FICTIONAL price book', 'source_date': '2026-06'} for r, d in rows if d in RATES]
    return json.dumps({'rates': out, 'searches': ['rates']}), {'https://fictional-prices.example/p': 'FICTIONAL price book'}


assistants._call = fake_call
org_research._ask = fake_ask

# ---------------- the team's templates folder and the pickers ----------------
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates', json={'folder': 'SharePoint/Pricing templates', 'outputs': 'SharePoint/Filled'}, headers=H)
P = r.json()
names = {x['name']: x for x in P['templates']}
t('the team\'s templates folder is set like Parker\'s (a new team version) and its spreadsheets are listed', r.status_code == 200
  and set(names) == {'FICTIONAL shared cost plan.xlsx', 'FICTIONAL by sheet.xlsx', 'FICTIONAL odd headers.xlsx', 'FICTIONAL simple.csv', 'FICTIONAL old.xls'}
  and 'Pricing templates' in teams.versions(TID)[0]['what'])
t('…an old .xls is listed but marked unreadable, with what to do', not names['FICTIONAL old.xls']['readable'] and '.xlsx' in names['FICTIONAL old.xls']['mapping_label'])
t('…templates are shared until you tag one to a client', all(x['shared'] for x in P['templates']))
sp = cl.get(f'/admin/api/teams/{TID}/start').json()
t('the Start a job screen offers the file role Cost/pricing template', sp['roles']['template'] == 'Cost/pricing template'
  and list(sp['roles'].values()) == ['Drawing', 'Specification', 'Schedule', 'Cost/pricing template', 'Other'])
t('…and the team\'s templates to pick from', len(sp['templates']['templates']) == 5)

# ---------------- what a file is, before it is added ----------------
ins = lambda name, raw: cl.post(f'/admin/api/teams/{TID}/inspect', json={'name': name, 'data': base64.b64encode(raw).decode()}, headers=H).json()
x = ins('FICTIONAL contractor pricing.xlsx', template_headings())
t('an empty pricing spreadsheet is guessed to be a pricing template', x['guess'] == 'template' and x['type'] == 'Excel' and x['sheets'] == 1)
priced = openpyxl.Workbook(); ws = priced.active; ws.append(['Description', 'Qty', 'Unit', 'Rate', 'Amount']); ws.append(['Blocks', 10, 'm2', 45, 450])
t('a priced schedule is guessed to be a schedule', ins('FICTIONAL door schedule.xlsx', xlsx(priced))['guess'] == 'schedule')
x = ins('FICTIONAL ground floor plan.pdf', b'')
t('a file that cannot be read says so', not x['readable'] and x['problem'])
x = ins('FICTIONAL elevations drawing.txt', b'North elevation. Not to scale is not stated here.')
t('a drawing is guessed from its name, and its scale looked for', x['guess'] == 'drawing' and x['scale'] is True)
x = ins('FICTIONAL sections drawing.txt', b'Cross section through the hall.')
t('…a drawing with no scale is flagged', x['guess'] == 'drawing' and x['scale'] is False)
x = ins('FICTIONAL secret.txt', b'aws key AKIAIOSFODNN7EXAMPLE for the drawings')
with s.db() as c: blocks = c.execute("SELECT count(*) FROM activity WHERE action='rule_blocked'").fetchone()[0]
t('a file the rules refuse is flagged before it is added, without logging a block (a preview)', x['problem'] and blocks == 0)
x = ins('FICTIONAL old.xls', b'x')
t('an old .xls says it must be saved as .xlsx', not x['readable'] and '.xlsx' in x['problem'])

# ---------------- the layout found in code ----------------
m = cl.get('/admin/api/teams/pricing-templates/mapping', params={'path': SHARED, 'team': TID}).json()
mp = m['mapping']
t('headings layout: header row, every column and the element headings found in code', m['status'] == 'detected' and mp['detected_by'] == 'code'
  and mp['mode'] == 'headings' and mp['sheets'][0]['header_row'] == 3
  and mp['sheets'][0]['columns'] == {'ref': 'A', 'description': 'B', 'quantity': 'C', 'unit': 'D', 'rate': 'E', 'amount': 'F', 'source': 'G'}
  and [e['name'] for e in mp['sheets'][0]['elements']] == ['Walls', 'Roofs'] and [x['row'] for x in mp['sheets'][0]['totals']] == [7, 10, 12])
t('…with a preview of the sheet to check it against', m['preview'][0]['sheet'] == 'Cost plan' and m['preview'][0]['rows'][2][1] == 'Ref')
m2 = PT.mapping(BYSHEET, TID)['mapping']
t('one sheet per element: each sheet is an element; a Read me sheet is left out', m2['mode'] == 'sheets' and [x['sheet'] for x in m2['sheets']] == ['Walls', 'Roof']
  and m2['sheets'][0]['columns'] == {'ref': 'A', 'description': 'B', 'quantity': 'C', 'unit': 'D', 'rate': 'E', 'amount': 'F', 'source': 'G'})
m3 = PT.mapping(CSVT, TID)['mapping']
t('a CSV template is read too', m3['mode'] == 'single' and m3['sheets'][0]['columns'] == {'description': 'A', 'quantity': 'B', 'unit': 'C', 'rate': 'D', 'amount': 'E', 'source': 'F'})
n0 = len(CALLS)
g = cl.get('/admin/api/teams/pricing-templates/mapping', params={'path': ODD, 'team': TID}).json()
t('opening a template code cannot read never calls a model (a GET): it offers to ask', g['status'] == 'none' and g['can_ask_model'] and len(CALLS) == n0)
MAPPER['reply'] = json.dumps({'sheets': [{'sheet': 'Pricing', 'header_row': 1, 'columns': {'ref': 'A', 'description': 'B', 'quantity': 'C', 'unit': 'D', 'rate': 'E', 'amount': 'F'}}]})
m4 = cl.post('/admin/api/teams/pricing-templates/mapping/detect', json={'path': ODD, 'team': TID}, headers=H).json()
t('when code cannot tell, the team lead\'s model is asked, once, and its answer checked in code', m4['status'] == 'detected' and m4['mapping']['detected_by'] == 'model'
  and len(CALLS) == n0 + 1 and 'template layout' in CALLS[-1]['workload'] and m4['mapping']['sheets'][0]['columns']['rate'] == 'E')
with s.db() as c: run = c.execute("SELECT count(*) FROM agent_runs WHERE agent_id='team-template-mapper'").fetchone()[0]
t('…as a tracked agent run', run == 1)
MAPPER['reply'] = json.dumps({'sheets': [{'sheet': 'Nope', 'header_row': 1, 'columns': {'description': 'B'}}]})
m5 = PT.mapping(ODD, TID, redetect=True)
t('…a model answer that does not fit the workbook is not used: set it by hand', m5['status'] == 'none' and m5['mapping'] is None and 'by hand' in m5['problem'])

# ---------------- saved, editable, asked again when the file changes ----------------
mp['sheets'][0]['columns']['source'] = 'G'
r = cl.put('/admin/api/teams/pricing-templates/mapping', json={'path': SHARED, 'mapping': {'mode': 'headings', 'sheets': mp['sheets']}}, headers=H)
t('Stefan confirms the mapping; it is saved with the template', r.status_code == 200 and r.json()['status'] == 'confirmed' and PT.describe(SHARED)['mapping'] == 'confirmed')
bad = cl.put('/admin/api/teams/pricing-templates/mapping', json={'path': SHARED, 'mapping': {'sheets': [{'sheet': 'Cost plan', 'header_row': 3, 'columns': {'description': 'B', 'quantity': 'C', 'unit': 'C', 'rate': 'E'}}]}}, headers=H)
t('a mapping with two roles in one column, or a role missing, is refused', bad.status_code == 400)
r = cl.put('/admin/api/teams/pricing-templates/mapping', json={'path': CSVT, 'mapping': {'sheets': [{'sheet': 'FICTIONAL simple', 'header_row': 1, 'edited': True,
       'columns': {'description': 'A', 'quantity': 'B', 'unit': 'C', 'rate': 'D', 'amount': 'E'}}], 'edited': True}}, headers=H)
t('a corrected mapping is kept as yours (here: no notes column)', r.status_code == 200 and r.json()['mapping']['detected_by'] == 'you' and 'source' not in r.json()['mapping']['sheets'][0]['columns'])
PT.confirm(BYSHEET, m2)
PT.confirm(ODD, {'sheets': [{'sheet': 'Pricing', 'header_row': 1, 'columns': {'description': 'B', 'quantity': 'C', 'unit': 'D', 'rate': 'E', 'amount': 'F'}}]})
time.sleep(0.01)
p_odd = TPL / 'FICTIONAL odd headers.xlsx'
p_odd.write_bytes(template_odd() + b'')
os.utime(p_odd, ns=(p_odd.stat().st_atime_ns, p_odd.stat().st_mtime_ns + 5_000_000_000))
MAPPER['reply'] = json.dumps({'sheets': [{'sheet': 'Pricing', 'header_row': 1, 'columns': {'description': 'B', 'quantity': 'C', 'unit': 'D', 'rate': 'E'}}]})
ch = PT.mapping(ODD, TID)
t('when the file changes, the mapping is asked for again (the old one kept to compare)', PT.describe(ODD)['mapping'] == 'detected' and ch['status'] == 'changed'
  and ch['previous'] and PT.confirmed(ODD) is None)

# ---------------- filling: the right cells, sources, rows added, totals in code, formulas checked ----------------
before = (TPL / 'FICTIONAL shared cost plan.xlsx').read_bytes()
mtime = (TPL / 'FICTIONAL shared cost plan.xlsx').stat().st_mtime_ns
j = cl.post(f'/admin/api/teams/{TID}/jobs', json={'job_type': 'cost-estimate', 'title': 'FICTIONAL hall', 'brief': 'A fictional single-storey hall, to be estimated in full.',
            'uploads': [{'name': DOC, 'kind': 'spec', 'text': 'FICTIONAL specification. Page 1 walls, page 2 roof, page 3 drainage.'}], 'template': SHARED}, headers=H).json()
if j.get('status') != 'waiting': print('JOB', j.get('status'), j.get('error'), j.get('detail'))
t('the job runs with its pricing template to sign-off', j['status'] == 'waiting' and j['pending'][0]['kind'] == 'signoff')
row = teams._row(j['id'])
docs = row['outputs']['documents']
tdoc = next(d for d in docs if d.get('template'))
import documents
data = documents.get(tdoc['id'])['data']
wb = openpyxl.load_workbook(io.BytesIO(data))
ws = wb['Cost plan']
cells = {r_[1]: r_ for r_ in ws.iter_rows(min_row=1, values_only=True) if r_[1]}
t('the filled copy is named after the job and version', tdoc['name'].startswith('FICTIONAL-hall-v1') and tdoc['name'].endswith('.xlsx'))
t('items go under the matching element (Roof → the template\'s Roofs)', cells['Roof covering'][0] == 'Q4' and ws['B10'].value == 'Roof covering' and ws['B9'].value == 'Roofs')
t('…with quantity, unit and rate in the mapped columns', cells['Wall item 1'][2] == 10 and cells['Wall item 1'][3] == 'm2' and cells['Wall item 1'][4] == 50)
t('…and where the rate came from, with the job page\'s badge, in the notes column', cells['Wall item 1'][6].startswith('Published: FICTIONAL price book https://fictional-prices.example/p (2026-06)'))
t('an unpriced item has no rate and says so', cells['Wall item 3'][4] is None and cells['Wall item 3'][6].startswith('Unpriced'))
t('a row is added within the template\'s formatting when an element has more items than rows', ws['B7'].value == 'Wall item 3' and ws['B7'].font.i
  and ws['F7'].value == '=C7*E7' and ws['F8'].value == '=SUM(F5:F7)' and ws['F13'].value == '=F8+F11')
t('the template\'s own formulas are left in place', ws['F5'].value == '=C5*E5' and ws['F10'].value == '=C10*E10')
other = [r_ for r_ in ws.iter_rows(values_only=True) if r_[1] == 'Gullies']
t('an item whose element the template does not have is listed after its rows, with its amount as a value', other and other[0][5] == 480.0)
af = wb["Alice's figures"]
vals = {r_[0]: r_[1] for r_ in af.iter_rows(values_only=True) if r_ and r_[0]}
cp = team_qs.final_plan(row['outputs'])
t('Alice\'s own figures are written as values: worked out in code, the same as the job page', vals['Construction'] == cp['construction'] == 500 + 800 + 3600 + 480
  and vals['Total excluding VAT'] == cp['total'])
fl = PT.fills(j['id'])
t('the formulas agree with Alice except where an item sits outside the template\'s totals, which is listed', len(fl) == 1
  and [d_['where'] for d_ in fl[0]['differences']] == [PT.OTHER])
t('the template file itself is never changed', (TPL / 'FICTIONAL shared cost plan.xlsx').read_bytes() == before and (TPL / 'FICTIONAL shared cost plan.xlsx').stat().st_mtime_ns == mtime)
t('the filled copy is saved to the team\'s outputs folder when one is set', fl[0]['library_path'].startswith('SharePoint/Filled/') and (root / fl[0]['library_path']).is_file())

# a template whose formulas disagree with Alice's figures
(TPL / 'FICTIONAL marked-up.xlsx').write_bytes(template_headings(amount_formula='=C{r}*E{r}*1.1'))
MU = 'SharePoint/Pricing templates/FICTIONAL marked-up.xlsx'
PT.confirm(MU, PT.mapping(MU)['mapping'])
r = cl.put(f'/admin/api/teams/jobs/{j["id"]}/template', json={'path': MU}, headers=H)
n_calls = len(CALLS)
fl = PT.fills(j['id'])[-1]
t('changing the template after pricing fills it again from the same items, without re-running the team', r.status_code == 200 and len(CALLS) == n_calls and fl['template'] == MU)
t('…and every difference between the template\'s formulas and Alice\'s figures is listed', any(d_['where'].startswith('Q1') and d_['template'] == 550.0 and d_['alice'] == 500.0 for d_ in fl['differences'])
  and any('Walls total' in d_['where'] for d_ in fl['differences']))
pg = cl.get(f'/admin/api/teams/jobs/{j["id"]}/page').json()
t('the job page shows the template, where it came from, its mapping and the differences', pg['pricing_template']['path'] == MU and pg['pricing_template']['from'] == 'chosen'
  and pg['pricing_template']['mapping'] == 'confirmed' and pg['pricing_template']['fills'][-1]['differences'])

# a template with no notes column: the source goes in a cell comment
r = cl.put(f'/admin/api/teams/jobs/{j["id"]}/template', json={'path': CSVT}, headers=H)
fl = PT.fills(j['id'])[-1]
w2 = openpyxl.load_workbook(io.BytesIO(documents.get(fl['doc_id'])['data']))
s2 = w2.worksheets[0]
rate_cell = next(r_ for r_ in s2.iter_rows() if r_[0].value == 'Wall item 1')[3]
t('with no source column mapped, the source is a comment on the rate cell', rate_cell.value == 50 and rate_cell.comment and rate_cell.comment.text.startswith('Published:'))
t('…a CSV template gives a filled workbook, amounts as values', next(r_ for r_ in s2.iter_rows(values_only=True) if r_[0] == 'Roof covering')[4] == 3600.0)

# re-pricing a version refills the template, named for that version
r = cl.put(f'/admin/api/teams/jobs/{j["id"]}/template', json={'path': SHARED}, headers=H)
RATES['Wall item 3'] = 55
cl.post(f'/admin/api/teams/jobs/{j["id"]}/reprice', json={'refs': ['Q3']}, headers=H)
fl = PT.fills(j['id'])[-1]
t('each re-priced version fills the template again, named for its version', fl['version'] == 2 and 'v2' in documents.get(fl['doc_id'])['name'])

# ---------------- shared across clients; a client's own template refused elsewhere; copies tagged ----------------
for n in ('FICTIONAL Aardvark Council', 'FICTIONAL Badger Trust'):
    organisations.create(n, client=True)
v_before = teams.get(TID)['version']
for n in ('FICTIONAL Aardvark Council', 'FICTIONAL Badger Trust'):
    r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/client-default', json={'client': n, 'path': SHARED}, headers=H)
t('several clients can have the same shared template as their default in the team\'s settings (Add)', r.status_code == 200
  and PT.client_default(teams.get(TID), 'FICTIONAL Aardvark Council')[0] == PT.client_default(teams.get(TID), 'FICTIONAL Badger Trust')[0] == SHARED
  and {x['client'] for x in r.json()['client_defaults']} == {'FICTIONAL Aardvark Council', 'FICTIONAL Badger Trust'})
t('…each change is a new team version', teams.get(TID)['version'] == v_before + 2)
mk = lambda title, client, **kw: cl.post(f'/admin/api/teams/{TID}/jobs', json={'job_type': 'cost-estimate', 'title': title, 'client': client,
                                          'brief': 'A fictional single-storey hall, to be estimated in full.',
                                          'uploads': [{'name': DOC, 'kind': 'spec', 'text': 'FICTIONAL specification.'}], **kw}, headers=H)
ja, jb = mk('FICTIONAL Aardvark hall', 'FICTIONAL Aardvark Council').json(), mk('FICTIONAL Badger hall', 'FICTIONAL Badger Trust').json()
ra, rb = teams._row(ja['id']), teams._row(jb['id'])
t('a new job uses the client\'s default template, marked as such; its mapping is saved once', ra['pricing_template'] == rb['pricing_template'] == SHARED
  and ra['pricing_template_from'] == 'client')
fa, fb = PT.fills(ja['id'])[-1], PT.fills(jb['id'])[-1]
t('each filled copy is tagged to its own job\'s client', fa['client'] == 'FICTIONAL Aardvark Council' and fb['client'] == 'FICTIONAL Badger Trust')
r = cl.put('/admin/api/teams/pricing-templates/client', json={'path': BYSHEET, 'client': 'FICTIONAL Aardvark Council'}, headers=H)
t('Stefan can tag a template to a client', r.status_code == 200 and not r.json()['shared'] and r.json()['client'] == 'FICTIONAL Aardvark Council')
r = mk('FICTIONAL Badger extension', 'FICTIONAL Badger Trust', template=BYSHEET)
t('a client-tagged template is refused on another client\'s job, naming the rule', r.status_code == 400 and 'FICTIONAL Aardvark Council' in r.json()['detail']
  and 'Client-facing documents' in r.json()['detail'])
with s.db() as c: blk = c.execute("SELECT count(*) FROM activity WHERE action='rule_blocked' AND rule='client_documents' AND target LIKE 'Pricing template%'").fetchone()[0]
t('…and the refusal is logged as that rule\'s block', blk >= 1)
t('…but used on that client\'s own job', mk('FICTIONAL Aardvark extension', 'FICTIONAL Aardvark Council', template=BYSHEET).status_code == 200)
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/client-default', json={'client': 'FICTIONAL Badger Trust', 'path': BYSHEET}, headers=H)
t('a template tagged to client A cannot be client B\'s default', r.status_code == 400 and 'tagged to FICTIONAL Aardvark Council' in r.json()['detail']
  and PT.client_default(teams.get(TID), 'FICTIONAL Badger Trust')[0] == SHARED)
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/client-default', json={'client': 'FICTIONAL Aardvark Council', 'path': BYSHEET}, headers=H)
t('…but it can be client A\'s own default (Change)', r.status_code == 200 and PT.client_default(teams.get(TID), 'FICTIONAL Aardvark Council')[0] == BYSHEET)
t('Start a job preselects the client\'s default, else nothing (then the team\'s default)', cl.get(f'/admin/api/teams/{TID}/start').json()['org_defaults'].keys()
  == {'FICTIONAL Aardvark Council', 'FICTIONAL Badger Trust'})
ra2 = teams._row(mk('FICTIONAL Aardvark preselect', 'FICTIONAL Aardvark Council').json()['id'])
t('…a new job for the client uses the client\'s default', ra2['pricing_template'] == BYSHEET and ra2['pricing_template_from'] == 'client')
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/client-default', json={'client': 'FICTIONAL Aardvark Council', 'path': SHARED}, headers=H)
r = cl.put(f'/admin/api/teams/{TID}/pricing-templates/client-default', json={'client': 'FICTIONAL Badger Trust', 'path': ''}, headers=H)
t('Remove takes a client\'s default out (a new team version)', r.status_code == 200 and PT.client_default(teams.get(TID), 'FICTIONAL Badger Trust') == ('', '')
  and 'removed' in r.json()['message'])
cl.put(f'/admin/api/teams/{TID}/pricing-templates', json={'default': ODD}, headers=H)
rb2 = teams._row(mk('FICTIONAL Badger preselect', 'FICTIONAL Badger Trust').json()['id'])
t('…then the team\'s default is used', rb2['pricing_template'] == ODD and rb2['pricing_template_from'] == 'team')
cl.put(f'/admin/api/teams/{TID}/pricing-templates', json={'default': ''}, headers=H)
rb3 = teams._row(mk('FICTIONAL Badger own layout', 'FICTIONAL Badger Trust').json()['id'])
t('…and with no team default, Alice\'s own layout', rb3['pricing_template'] == '' and rb3['pricing_template_from'] == '')
rb4 = teams._row(mk('FICTIONAL Badger chosen', 'FICTIONAL Badger Trust', template=SHARED).json()['id'])
t('…the job can still choose another template', rb4['pricing_template'] == SHARED and rb4['pricing_template_from'] == 'chosen')
cl.put(f'/admin/api/teams/{TID}/pricing-templates/client-default', json={'client': 'FICTIONAL Badger Trust', 'path': SHARED}, headers=H)

# ---------------- D-0039: the old organisation-page choices move into the team's setting, after a preview ----------------
organisations.create('FICTIONAL Curlew Housing', client=True)
organisations.create('FICTIONAL Dunlin Estates', client=True)
with s.db() as c:        # choices made on the organisation pages before this release (the throwaway test database)
    for org, path in (('FICTIONAL Curlew Housing', SHARED), ('FICTIONAL Dunlin Estates', BYSHEET), ('FICTIONAL Badger Trust', ODD)):
        c.execute('INSERT INTO org_pricing_templates(org,path,set_by,set_at) VALUES (?,?,?,?)', (org, path, 'Owner', s.now()))
t('until the move is confirmed, an organisation\'s old choice still applies', PT.default_for(teams.get(TID), 'FICTIONAL Curlew Housing') == (SHARED, 'client'))
t('…the team\'s own setting wins over the old choice', PT.default_for(teams.get(TID), 'FICTIONAL Badger Trust') == (SHARED, 'client'))
v_before = teams.get(TID)['version']
pv = cl.get('/admin/api/teams/pricing-templates/org-move').json()
by = {it['org']: {x['team_id']: x for x in it['teams']} for it in pv['items']}
t('the preview lists each choice and what would happen in each team, changing nothing', pv['pending'] and pv['counts']['choices'] == 3
  and by['FICTIONAL Curlew Housing'][TID]['status'] == 'add' and by['FICTIONAL Badger Trust'][TID]['status'] == 'keep'
  and by['FICTIONAL Dunlin Estates'][TID]['status'] == 'skip' and 'tagged to FICTIONAL Aardvark Council' in by['FICTIONAL Dunlin Estates'][TID]['note']
  and teams.get(TID)['version'] == v_before)
r = cl.post('/admin/api/teams/pricing-templates/org-move', json={'counts': {**pv['counts'], 'add': pv['counts']['add'] + 1}}, headers=H)
t('…confirming with counts that differ from the preview is refused', r.status_code == 400 and teams.get(TID)['version'] == v_before)
r = cl.post('/admin/api/teams/pricing-templates/org-move', json={'counts': pv['counts']}, headers=H)
cd = PT.team_settings(teams.get(TID))['clients']
t('the move keeps existing choices: added to the team\'s setting, the team\'s own kept, a tag clash left out', r.status_code == 200
  and cd.get('FICTIONAL Curlew Housing') == SHARED and cd.get('FICTIONAL Badger Trust') == SHARED and 'FICTIONAL Dunlin Estates' not in cd
  and teams.get(TID)['version'] == v_before + 1)
with s.db() as c:
    kept = c.execute('SELECT count(*) FROM org_pricing_templates').fetchone()[0]
    c.execute('UPDATE org_pricing_templates SET path=? WHERE org=?', (ODD, 'FICTIONAL Curlew Housing'))
    moved_log = c.execute("SELECT count(*) FROM activity WHERE action='pricing_org_defaults_moved'").fetchone()[0]
t('…the old field is retired: its rows are kept (nothing deleted) but no longer read', kept == 3
  and PT.default_for(teams.get(TID), 'FICTIONAL Curlew Housing') == (SHARED, 'client') and PT.client_default(teams.get(TID), 'FICTIONAL Dunlin Estates') == ('', ''))
t('…logged, and the preview now says it is done', moved_log == 1 and cl.get('/admin/api/teams/pricing-templates/org-move').json()['pending'] is False
  and cl.post('/admin/api/teams/pricing-templates/org-move', json={'counts': pv['counts']}, headers=H).status_code == 400)
oh = cl.get('/admin/organisations').text
t('the organisation\'s Details page has no pricing template control any more', 'o-ptpl' not in oh and 'pricing-templates/org-default' not in oh
  and 'pricing-templates/all' not in oh)
t('…and the old routes are gone', cl.put('/admin/api/teams/pricing-templates/org-default', json={'org': 'x', 'path': ''}, headers=H).status_code in (404, 405))
rules_engine.update_rule('client_documents', enabled=False)
t('switched off on the Rules page, the client\'s template may be used anywhere', mk('FICTIONAL Badger annex', 'FICTIONAL Badger Trust', template=BYSHEET).status_code == 200)
rules_engine.update_rule('client_documents', enabled=True)
t('no template: Alice\'s own layout (an explicit choice)', teams._row(mk('FICTIONAL plain', '', template='').json()['id'])['pricing_template'] == '')

# a template whose Purview label is blocked is never used
from openpyxl.packaging.custom import StringProperty
lb = openpyxl.load_workbook(io.BytesIO(template_headings()))
G = '99999999-8888-7777-6666-555555555555'
for k, v in (('Enabled', 'true'), ('Name', 'OFFICIAL-SENSITIVE')): lb.custom_doc_props.append(StringProperty(name=f'MSIP_Label_{G}_{k}', value=v))
(TPL / 'FICTIONAL labelled.xlsx').write_bytes(xlsx(lb))
r = mk('FICTIONAL labelled job', '', template='SharePoint/Pricing templates/FICTIONAL labelled.xlsx')
t('a template whose Purview label is blocked is refused', r.status_code == 400 and 'Purview label' in r.json()['detail'])

# ---------------- add a template from the start screen ----------------
up = cl.post(f'/admin/api/teams/{TID}/pricing-templates', json={'name': 'FICTIONAL uploaded.xlsx', 'data': base64.b64encode(template_sheets()).decode()}, headers=H).json()
t('Add a template saves the upload into the team\'s templates folder and detects its layout', up['path'] == 'SharePoint/Pricing templates/FICTIONAL uploaded.xlsx'
  and (TPL / 'FICTIONAL uploaded.xlsx').is_file() and up['detected']['mapping']['mode'] == 'sheets')
up2 = cl.post(f'/admin/api/teams/{TID}/pricing-templates', json={'name': 'FICTIONAL uploaded.xlsx', 'data': base64.b64encode(template_sheets()).decode()}, headers=H).json()
t('…never overwriting a file that is there', up2['path'] == 'SharePoint/Pricing templates/FICTIONAL uploaded (2).xlsx')
t('…only spreadsheets', cl.post(f'/admin/api/teams/{TID}/pricing-templates', json={'name': 'x.docx', 'data': base64.b64encode(b'x').decode()}, headers=H).status_code == 400)

# ---------------- the Ready to start checklist ----------------
chk = lambda form: cl.post(f'/admin/api/teams/{TID}/start-check', json={'form': form}, headers=H).json()
c0 = chk({})
t('an empty form is not ready: name, what\'s needed, brief and documents are required', not c0['ready']
  and {i['key'] for i in c0['items'] if i['level'] == 'required' and not i['ok']} == {'title', 'job_type', 'brief', 'documents'})
form = {'job_type': 'cost-estimate', 'title': 'FICTIONAL hall', 'brief': 'A fictional single-storey hall, to be estimated.', 'location': '',
        'documents': [{'name': 'FICTIONAL plan drawing.pdf', 'role': 'drawing', 'scale': False}], 'template': {'path': ODD}}
c1 = chk(form)
keys = {i['key']: i for i in c1['items']}
t('filled in, it is ready, with warnings that may lead to estimates', c1['ready'] and not keys['location']['ok'] and keys['location']['level'] == 'warning'
  and not keys['scale:FICTIONAL plan drawing.pdf']['ok'] and 'approximate' in keys['scale:FICTIONAL plan drawing.pdf']['text'] and not keys['only_drawings']['ok'])
t('…a template whose mapping is not confirmed is a warning', keys['template']['level'] == 'warning' and 'confirm' in keys['template']['text'].lower())
form.update(template={'path': SHARED}, location='Perth', documents=[{'name': 'FICTIONAL spec.md', 'role': 'spec'}])
keys = {i['key']: i for i in chk(form)['items']}
t('…a confirmed shared template is shown as ready', keys['template']['ok'] and 'shared' in keys['template']['text'] and keys['location']['ok'])
form.update(template={'path': BYSHEET}, client_name='FICTIONAL Badger Trust')
with s.db() as c: nb = c.execute("SELECT count(*) FROM activity WHERE action='rule_blocked'").fetchone()[0]
c3 = chk(form)
with s.db() as c: nb2 = c.execute("SELECT count(*) FROM activity WHERE action='rule_blocked'").fetchone()[0]
t('…checked as you type without logging a block (that happens only if the job is started)', nb2 == nb)
t('…another client\'s template stops the job', not c3['ready'] and not {i['key']: i for i in c3['items']}['template']['ok'])
form.update(template=None)
t('…no template: Alice\'s own layout, as a note', {i['key']: i for i in chk(form)['items']}['template']['text'].startswith('No pricing template'))
form.update(documents=[{'name': 'x', 'role': 'spec', 'problem': 'Not saved: secrets'}])
t('…a file the rules refuse stops the job', not chk(form)['ready'])

# ---------------- what happens next: members, typical time and cost (hidden with no history) ----------------
other = teams.from_template('quantity-surveying', name='Fresh team')
fp = cl.get(f'/admin/api/teams/{other["id"]}/start').json()
t('typical run time and AI cost are hidden when the team has no finished jobs (never invented)', fp['job_types'][0]['typical'] is None)
t('each member is listed with what it will do on this job', [m_['role'] for m_ in fp['job_types'][0]['members']][:2] == ['Lead QS', 'Measurement Surveyor']
  and fp['job_types'][0]['members'][0]['lead'] and fp['job_types'][0]['members'][0]['does'])
cl.post(f'/admin/api/teams/steps/{teams.job_detail(ja["id"])["pending"][0]["id"]}', json={'action': 'approve'}, headers=H)
ty = cl.get(f'/admin/api/teams/{TID}/start').json()['job_types'][0]['typical']
t('with a finished job, the typical time to sign-off and AI cost come from this team\'s past jobs', ty and ty['jobs'] == 1 and ty['run_time'].startswith('about ')
  and 'text' in ty['ai_cost'])

# ---------------- approval and estimates chosen for this job ----------------
jp = mk('FICTIONAL approve each', '', template='', autonomy='approve').json()
t('approval set for this job overrides the team\'s (every hand-off waits)', jp['pending'] and jp['pending'][0]['kind'] == 'handoff' and jp['autonomy'] == 'approve')
SEARCHES.clear()
je = mk('FICTIONAL estimates job', '', template='', estimates=True).json()
t('team estimates allowed on this job reach the Cost Surveyor for every item', SEARCHES and all('ESTIMATE ALLOWED' in x['query'] for x in SEARCHES)
  and teams.job_page(je['id'])['estimates_all'])

# ---------------- the pages ----------------
html = cl.get(f'/admin/teams/{TID}/start').text
t('the Start a job screen is served and drawn by the Teams page', cl.get(f'/admin/teams/{TID}/start').status_code == 200
  and all(x in html for x in ("'Start a job for '", "'What happens next'", "'Ready to start'", "'Pick from a document library'", "'Drop the documents here'",
                              "'Allow team estimates on this job'", "'Cost/pricing template", "'Confirm mapping'", "'Pricing templates'")))
t('the team\'s Knowledge tab has Default template per client (Add, Change, Remove) and the move card', all(x in cl.get(f'/admin/teams/{TID}').text for x in (
  "'Default template per client'", "'/pricing-templates/client-default'", "btn('Change'", "btn('Remove'", "btn('Add'", "'Preview the move'", "'Confirm the move'")))
