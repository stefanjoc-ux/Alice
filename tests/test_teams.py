"""Digital teams: the quantity surveying team runs a cost estimate stage by stage. Hand-offs, send-backs (by a member and by
Stefan), questions on Actions, both autonomy modes, team edits versioned with Undo and jobs pinned to their version, Temple's
suggestions needing approval, arithmetic done in code, no rate without a source, the location factor only when sourced, Market
Trends with and without history, rules on every member call, and the fictional demo job end to end. No real model is called."""
import _util  # first: throwaway data folder, dummy keys, no real model calls
from _util import t
import io, json, zipfile
import app, substrate_store as s, agents, actions, action_cards, assistants, org_research, rules_engine, knowledge, temple_discuss, temple
import teams, team_qs
from fastapi.testclient import TestClient
cl = TestClient(app.app); H = {'x-admin-token': app.ADMIN_TOKEN}
teams.BACKGROUND = False
temple.save_settings(False, 'claude')
TID, JT = team_qs.TEAM_ID, 'cost-estimate'
teams.set_missing_info(TID, 'ask')        # these checks cover members asking Stefan (test_team_files covers Assume and flag)

# ---------------- stand-in models ----------------
SEEN = []          # (workload, system, payload) for every member call
PLAN = {'plan': 'Measure by element from the specification and schedules.', 'elements': ['Substructure', 'Roof', 'Finishes'],
        'documents': [{'name': 'FICTIONAL specification.md', 'use': 'quantities'}], 'location': '', 'summary': 'Planned the estimate.',
        'note': 'Measure substructure, roof and finishes.', 'questions': []}
GOOD_ITEMS = [{'element': 'Substructure', 'description': 'Concrete strip foundation 600 x 900 mm including excavation', 'quantity': 74, 'unit': 'm',
               'source': {'document': 'FICTIONAL specification.md', 'page': '1'}},
              {'element': 'Roof', 'description': 'Concrete interlocking roof tiles on battens and membrane', 'quantity': 381, 'unit': 'm2',
               'source': {'document': 'FICTIONAL specification.md', 'page': '2'}},
              {'element': 'Finishes', 'description': 'Sprung timber sports floor', 'quantity': 180, 'unit': 'm2',
               'source': {'document': 'FICTIONAL finishes schedule.csv', 'line': 'Line 3'}},
              {'element': 'Windows', 'description': 'Rooflight to main hall', 'quantity': 2, 'unit': 'nr',
               'source': {'document': 'FICTIONAL specification.md', 'page': '2'}, 'from_drawing': True}]
STATE = {'plan': [PLAN], 'measure': [], 'assemble': [], 'trends': []}


def reply(kind, default):
    q = STATE.get(kind) or []
    return q.pop(0) if q else default


def fake_call(provider, system, messages, max_tokens=1500, timeout=60, workload='', meta=None):
    payload = messages[0]['content']
    SEEN.append((workload, system, payload))
    if 'Lead QS' in workload and 'COST PLAN FIGURES' in payload:
        return json.dumps(reply('assemble', {'accept': True, 'summary': 'A modest single-storey hall.', 'assumptions': ['Level site, no abnormals.'],
                                             'exclusions': ['VAT', 'Loose furniture'], 'risks': [{'risk': 'Ground conditions', 'mitigation': 'Site investigation'}],
                                             'note': 'Over to Market Trends.', 'total': 999}))
    if 'Lead QS' in workload: return json.dumps(reply('plan', PLAN))
    if 'Measurement' in workload: return json.dumps(reply('measure', {'accept': True, 'items': GOOD_ITEMS, 'summary': 'Measured.', 'note': 'Please price these.'}))
    if 'Market Trends' in workload:
        return json.dumps(reply('trends', {'matches': [{'ref': 'Q1', 'history': ['H1']}], 'commentary': 'Foundations are a little dearer than last time.', 'summary': 'Compared.'}))
    raise AssertionError('unexpected call ' + workload)


assistants._call = fake_call
SEARCHED = []
RET = {'https://fictional-prices.example/foundations': 'Fictional price book: foundations', 'https://fictional-prices.example/location-factors': 'Fictional location factors'}
PRICE = {'rates': [{'ref': 'Q1', 'rate': 160, 'unit': 'm', 'source_url': 'https://fictional-prices.example/foundations', 'source_title': 'Foundations', 'source_date': '2026-05'},
                   {'ref': 'Q2', 'rate': 52, 'unit': 'm2', 'source_url': 'https://not-returned.example/tiles', 'source_date': '2026-04'},
                   {'ref': 'Q3', 'rate': 90, 'unit': 'm2', 'source_url': 'https://fictional-prices.example/foundations', 'source_date': ''},
                   {'ref': 'Q4', 'rate': 700, 'unit': 'nr', 'source_url': '', 'library_id': ''}],
         'location_factor': {'factor': 1.03, 'region': 'Scotland', 'source_url': 'https://fictional-prices.example/location-factors', 'source_date': '2026-02'},
         'searches': ['strip foundation rate per metre UK 2026'], 'summary': '', 'note': 'Priced.'}
PRICES = []


def fake_ask(prompt, query, provider, workload='', **kw):
    SEARCHED.append((prompt, query, provider))
    return json.dumps(PRICES.pop(0) if PRICES else PRICE), dict(RET)


org_research._ask = fake_ask

# Market Trends QS searches the market too (task 3): its queries get a plain answer with no findings here
_price_ask = org_research._ask
org_research._ask = lambda prompt, query, provider, workload='', **kw: ((json.dumps({'findings': [], 'commentary': 'No market evidence.'}), {})
                                                                         if 'MEASURED ITEMS' not in query else _price_ask(prompt, query, provider, workload, **kw))

# ---------------- the seeded team, the page, the agents ----------------
team = teams.get(TID)
t('the Quantity surveying team is seeded with its four members', [m['role'] for m in team['members']] == ['Lead QS', 'Measurement Surveyor', 'Cost Surveyor', 'Market Trends QS'])
jt = team['job_types'][0]
t('job type Cost estimate: plan → measure → price → assemble → trends, each with what is handed on and what the receiver checks',
  jt['name'] == 'Cost estimate' and [x['key'] for x in jt['stages']] == ['plan', 'measure', 'price', 'assemble', 'trends', 'adjust']
  and all(x['hands'] for x in jt['stages']) and all(x['checks'] for x in jt['stages'][1:]))
t('autonomy starts as Approve every hand-off', team['autonomy'] == 'approve')
import admin_ui
t('Teams page under Workspace, with its own icon', 'teams' in dict(admin_ui.NAV_GROUPS)['Workspace'] and 'teams' in admin_ui.NAV_ICONS and 'teams' in admin_ui.PAGES)
t('the page renders', cl.get('/admin/teams').status_code == 200)
t('member turns and Temple\'s refinements are agents in the Digital teams group',
  {'team-member', 'temple-team-coach'} <= {b[0] for b in agents.BUILTIN} and agents._AGENT_GROUP.get('team-member') == 'teams' and agents._AGENT_GROUP.get('temple-team-coach') == 'teams')

# ---------------- rate library ----------------
r = cl.post(f'/admin/api/teams/{TID}/rates/demo', headers=H)
t('the fictional demo rate library loads', r.status_code == 200 and r.json()['added'] == 10)
t('…and only once', cl.post(f'/admin/api/teams/{TID}/rates/demo', headers=H).status_code == 400)
import base64
bad = base64.b64encode(b'Description,Unit,Rate\nThing,m2,abc\n').decode()
t('a rate library without usable rows is refused', cl.post(f'/admin/api/teams/{TID}/rates', json={'name': 'x.csv', 'data': bad}, headers=H).status_code == 400)
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.append(['Code', 'Description', 'Unit', 'Rate', 'As of']); ws.append(['X1', 'Fictional rooflight', 'nr', 640, '2025-09']); buf = io.BytesIO(); wb.save(buf)
r = cl.post(f'/admin/api/teams/{TID}/rates', json={'name': 'rates.xlsx', 'data': base64.b64encode(buf.getvalue()).decode(), 'label': 'Fictional returns'}, headers=H)
t('an Excel rate library is read', r.status_code == 200 and r.json()['added'] == 1)
xbatch = r.json()['batch']
sec = base64.b64encode(b'Description,Unit,Rate\nThing sk-ant-api03-' + b'A' * 90 + b',m2,10\n').decode()
t('a rate library holding a secret is refused', cl.post(f'/admin/api/teams/{TID}/rates', json={'name': 's.csv', 'data': sec}, headers=H).status_code == 400)

# ---------------- the demo job, approving every hand-off ----------------
demo = cl.get('/admin/api/teams/demo-project').json()
t('the demo project is clearly fictional', 'FICTIONAL' in demo['title'] and all('FICTIONAL' in d['name'] and 'FICTIONAL' in d['text'] for d in demo['documents']))


def start(**kw):
    body = {'job_type': JT, 'title': demo['title'], 'brief': demo['brief'], 'location': demo['location'], 'uploads': demo['documents']}
    body.update(kw)
    r = cl.post(f'/admin/api/teams/{TID}/jobs', json=body, headers=H)
    assert r.status_code == 200, r.text
    return r.json()


def job(jid): return cl.get(f'/admin/api/teams/jobs/{jid}').json()
def step(sid, action, note=''): return cl.post(f'/admin/api/teams/steps/{sid}', json={'action': action, 'note': note}, headers=H)


j = start()
jid = j['id']
p = j['pending']
t('the job runs the Lead QS, then waits for Stefan to approve the first hand-off', j['status'] == 'waiting' and len(p) == 1 and p[0]['kind'] == 'handoff'
  and p[0]['role'] == 'Lead QS' and p[0]['to_role'] == 'Measurement Surveyor' and j['outputs']['plan']['location'] == 'Perth, Scotland')
t('the job records the team version it started on', j['team_version'] == teams.get(TID)['current_version'])
sec_ = next(x for x in actions.summary()['sections'] if x['key'] == 'teams')
t('the hand-off waits on Actions', sec_['count'] == 1 and sec_['items'][0]['type'] == 'team_handoff' and 'Lead QS → Measurement Surveyor' in sec_['items'][0]['title'])
card = cl.get('/admin/api/cards/review/h-' + p[0]['id']).json()
t('…as an information card with Discuss with Temple', card['kind_label'].startswith('Digital team') and card['discuss']['url'] == f'/admin/api/review-items/h-{p[0]["id"]}/discussion'
  and [x['key'] for x in card['sections']][:2] == ['what', 'what'] and any(x['key'] == 'where' for x in card['sections'])
  and card['sections'][1]['text'].startswith('Measure by element') and 'Elements: Substructure' in card['sections'][1]['text'])
SENT_T = []
def fake_complete(provider, system, messages):
    SENT_T.append((system, messages)); return 'Check that every element in the plan has a document to measure from.', 'claude-haiku-4-5-20251001'
temple_discuss._complete = fake_complete
r = cl.post(f'/admin/api/review-items/h-{p[0]["id"]}/discussion', json={'message': 'What should I check?'}, headers=H)
t('Discuss with Temple on a hand-off: Temple sees the job and the hand-off, and changes nothing', r.status_code == 200
  and 'Lead QS' in SENT_T[-1][1][0]['content'] and job(jid)['pending'][0]['status'] == 'pending')

t('a send-back needs a reason', step(p[0]['id'], 'send_back').status_code == 400)
r = step(p[0]['id'], 'send_back', 'Include external works in the elements.')
j = job(jid)
t('Stefan sends the hand-off back: the Lead QS redoes the plan with his note, then the hand-off waits again',
  r.status_code == 200 and 'Include external works in the elements.' in SEEN[-1][2] and j['pending'][0]['kind'] == 'handoff' and j['pending'][0]['role'] == 'Lead QS')
t('a decided step cannot be decided again', step(p[0]['id'], 'approve').status_code == 400)

# the receiver sends work back: the first take-off has no sources, so the Cost Surveyor's code check returns it
STATE['measure'] = [{'accept': True, 'items': [{'element': 'Roof', 'description': 'Tiles', 'quantity': 381, 'unit': 'm2', 'source': {'document': 'FICTIONAL specification.md'}}],
                     'summary': 'Measured.', 'note': 'Here you go.'}]
step(j['pending'][0]['id'], 'approve')
j = job(jid)
t('the Measurement Surveyor\'s first take-off: an item without a page or line is left out, with the reason',
  j['pending'][0]['role'] == 'Measurement Surveyor' and j['outputs']['measure']['items'] == [] and 'no page or schedule line' in j['outputs']['measure']['rejected'][0]['reason'])
step(j['pending'][0]['id'], 'approve')
j = job(jid)
sb = [x for x in j['steps'] if x['kind'] == 'sendback']
t('the Cost Surveyor checks before accepting and sends the work back with reasons (no web search made)', len(sb) == 1 and sb[0]['role'] == 'Cost Surveyor'
  and sb[0]['to_role'] == 'Measurement Surveyor' and 'No measured items' in sb[0]['note'] and not SEARCHED)
t('…the Measurement Surveyor redid it with those reasons, and the new hand-off waits', 'No measured items' in [w for w in SEEN if 'Measurement' in w[0]][-1][2]
  and j['pending'][0]['role'] == 'Measurement Surveyor' and len(j['outputs']['measure']['items']) == 4)
items = j['outputs']['measure']['items']
t('every quantity names its source; the item from a drawing is approximate', all(i['source_text'] for i in items) and items[3]['approximate'] and not items[0]['approximate'])

step(j['pending'][0]['id'], 'approve')
j = job(jid)
pr = j['outputs']['price']
byref = {i['ref']: i for i in pr['items']}
t('the Cost Surveyor searched the web with the provider search (same code as organisation research)', SEARCHED and 'Q1 |' in SEARCHED[-1][1])
t('a published rate is used only with a page the search returned and its date', byref['Q1']['rate_source'] == 'web' and byref['Q1']['source_url'] in RET and byref['Q1']['source_date'] == '2026-05')
t('a rate citing a page the search did not return is not used; the rate library fills in (rate from the library, not the model)',
  byref['Q2']['rate_source'] == 'library' and byref['Q2']['rate'] == 48.0 and any(x['ref'] == 'Q2' and 'not among the pages' in x['reason'] for x in pr['refused_rates']))
t('a rate without a date is not used', byref['Q3']['rate_source'] == 'library' and any(x['ref'] == 'Q3' and 'no date' in x['reason'] for x in pr['refused_rates']))
t('an item with no published rate and no library row is unpriced, never invented', byref['Q4']['rate_source'] == 'unpriced' and byref['Q4']['rate'] is None)
t('each item says which of the three its rate came from', all(i['rate_source'] in ('web', 'library', 'unpriced') for i in pr['items']))
t('the regional factor is applied only because a returned page supports it, and cited', pr['location']['applied'] and pr['location']['factor'] == 1.03
  and pr['location']['source_url'] == 'https://fictional-prices.example/location-factors')
tr = [x for x in j['steps'] if x['kind'] == 'turn' and x['stage'] == 'price'][-1]
t('each search and source is recorded with the member\'s run', tr['content']['searches'][0]['sources'] and tr['run_id'])
with s.db() as c:
    ev = [dict(r) for r in c.execute("SELECT * FROM agent_events WHERE run_id=? AND kind='read' AND target_type='web'", (tr['run_id'],))]
t('…and as web reads on the agent run (for Temple and Data touched)', {e['target_id'] for e in ev} == set(RET) and any('cited' in e['detail'] for e in ev))

step(j['pending'][0]['id'], 'approve')
j = job(jid)
cp = j['outputs']['assemble']['cost_plan']
exp = team_qs.compute(pr['items'], 1.03, team_qs.SETTINGS)
t('the cost plan is the code\'s arithmetic, not the model\'s total', cp == exp and cp['total'] != 999)
q1 = next(l for l in cp['lines'] if l['ref'] == 'Q1')
t('quantity × rate × factor in code, to the penny', q1['amount'] == 12195.20)          # 74 × 160 × 1.03
a = j['outputs']['assemble']
t('assumptions name unpriced and approximate items and the regional factor', any('Q4' in x and 'unpriced' in x for x in a['assumptions']) and any('approximate' in x for x in a['assumptions'])
  and any('Regional factor' in x for x in a['assumptions']) and a['exclusions'] and a['risks'])

step(j['pending'][0]['id'], 'approve')
j = job(jid)
trn = j['outputs']['trends']
t('Market Trends with no history says so plainly (no model call)', trn['history'] == 0 and 'no past cost plans or team jobs' in trn['report'].lower()
  and not any('Market Trends' in w[0] for w in SEEN))
step(j['pending'][0]['id'], 'approve')
j = job(jid)
t('no market adjustment suggested: the Lead QS has nothing to decide (no model call)', j['outputs']['adjust']['decision'] == 'none')
t('the final output waits for sign-off, with a Word and an Excel document', j['pending'][0]['kind'] == 'signoff' and {d['kind'] for d in j['outputs']['documents']} == {'Word', 'Excel'})
dl = {d['name'].rsplit('.', 1)[1]: cl.get('/documents/' + d['id'] + '/download') for d in j['outputs']['documents']}
wtext = zipfile.ZipFile(io.BytesIO(dl['docx'].content)).read('word/document.xml').decode()
from openpyxl import load_workbook
xwb = load_workbook(io.BytesIO(dl['xlsx'].content))
xtext = ' '.join(str(v) for ws_ in xwb.worksheets for row in ws_.iter_rows(values_only=True) for v in row if v is not None)
t('both documents are marked as a draft for a qualified QS to review', team_qs.DRAFT_MARK.split(':')[0] in wtext and 'qualified quantity surveyor' in wtext
  and 'qualified quantity surveyor' in xtext)
t('the Word cost plan cites each rate\'s source', 'fictional-prices.example/foundations' in wtext and 'Your rate library' in wtext)
t('the sign-off waits on Actions too', any(i['type'] == 'team_signoff' for i in next(x for x in actions.summary()['sections'] if x['key'] == 'teams')['items']))
r = step(j['pending'][0]['id'], 'approve')
j = r.json()
t('signed off: the job is done and saved to Knowledge with its rates', j['status'] == 'done' and j['knowledge_id'])
with s.db() as c: ktext = c.execute('SELECT text FROM files WHERE id=?', (j['knowledge_id'],)).fetchone()[0]
t('…with a RATES USED block for later comparisons', 'RATE | Q1 |' in ktext and team_qs.DRAFT_MARK.split(':')[0] in ktext)
t('AI cost is recorded on the job', j['ai_cost'] >= 0 and all('cost_usd' in x for x in j['steps']))

# ---------------- run on its own, sign off at the end; Market Trends with history ----------------
t('autonomy change is versioned', cl.put(f'/admin/api/teams/{TID}/autonomy', json={'autonomy': 'signoff'}, headers=H).status_code == 200
  and teams.versions(TID)[0]['what'].startswith('Autonomy'))
SEEN.clear()
j2 = start(title='FICTIONAL demo: second hall', location='')
pend = j2['pending']
t('“Run, I sign off at the end”: hand-offs go ahead and only the final output waits', j2['status'] == 'waiting' and len(pend) == 1 and pend[0]['kind'] == 'signoff'
  and all(x['status'] == 'auto' for x in j2['steps'] if x['kind'] == 'handoff'))
t('no location: national rates, and the factor is not applied', not j2['outputs']['price']['location']['applied'] and 'national rates' in j2['outputs']['price']['location']['note'].lower())
tr2 = j2['outputs']['trends']
t('Market Trends compares with the earlier job held in Alice, percentages in code', tr2['history'] > 0 and tr2['comparisons'] and tr2['comparisons'][0]['ref'] == 'Q1'
  and tr2['comparisons'][0]['past'][0]['difference_pct'] == team_qs.pct_change(tr2['comparisons'][0]['rate_now'], tr2['comparisons'][0]['past'][0]['rate']))
t('pct_change is exact', team_qs.pct_change(165, 150) == 10.0 and team_qs.pct_change(100, 160) == -37.5)

# history held in Knowledge (a cost plan saved elsewhere, with RATE lines), client separation on it
knowledge.create('note', 'Older cost plan (fictional)', 'Fictional earlier scheme.\nRATE | A1 | Strip foundations | m | 140 | library | Fictional book | 2024-11', 'Fictional', 'you')
knowledge.create('note', 'Other client plan (fictional)', 'Fictional client scheme.\nRATE | B1 | Strip foundations secret client | m | 999 | web | x | 2025-01', 'Fictional', 'you')
import clients
clients.create_client('Fictional Other Council')
kid = [i for i in knowledge.listing(limit=100)['items'] if i['title'] == 'Other client plan (fictional)'][0]['id']
clients.tag('file', [kid], 'Fictional Other Council')
hist = team_qs.history({'id': 'x', 'client': ''}, teams.get(TID)['members'][3])
t('history includes cost plans in Knowledge with RATE lines', any(h['source'] == 'Older cost plan (fictional)' and h['rate'] == 140 for h in hist))
t('…but never another client\'s', not any('secret client' in h['description'] for h in hist))

# ---------------- questions for Stefan ----------------
teams.set_autonomy(TID, 'approve')
STATE['plan'] = [dict(PLAN, questions=['Is the hall to include a stage?']), PLAN]
j3 = start(title='FICTIONAL demo: hall with a question')
qq = j3['pending'][0]
t('a member asks Stefan a question: it waits on Actions', qq['kind'] == 'question' and any(i['type'] == 'team_question' for i in next(x for x in actions.summary()['sections'] if x['key'] == 'teams')['items']))
t('an answer is required', step(qq['id'], 'answer').status_code == 400)
step(qq['id'], 'answer', 'No stage; a retractable platform only.')
j3 = job(j3['id'])
t('the answer reaches the member, which carries on', 'No stage; a retractable platform only.' in [w for w in SEEN if 'Lead QS' in w[0]][-1][2] and j3['pending'][0]['kind'] == 'handoff')

# ---------------- edits versioned with Undo; jobs pinned to their version ----------------
v_before = teams.get(TID)['current_version']
r = cl.put(f'/admin/api/teams/{TID}/members/measurement-surveyor', json={'instructions': 'NEW INSTRUCTIONS: measure in metric only.'}, headers=H)
t('editing a member makes a new version that says who and what', r.status_code == 200 and teams.get(TID)['current_version'] == v_before + 1
  and 'Measurement Surveyor: instructions' in teams.versions(TID)[0]['what'] and teams.versions(TID)[0]['changed_by'])
step(j3['pending'][0]['id'], 'approve')
t('a job keeps the team version it started on (old instructions)', 'NEW INSTRUCTIONS' not in [w for w in SEEN if 'Measurement' in w[0]][-1][1])
r = cl.post(f'/admin/api/teams/{TID}/restore', json={'undo': True}, headers=H)
t('Undo restores the previous instructions as a new version', r.status_code == 200 and 'NEW INSTRUCTIONS' not in teams._member(teams.get(TID), 'measurement-surveyor')['instructions']
  and teams.get(TID)['current_version'] == v_before + 2 and teams.versions(TID)[0]['what'].startswith('Undo'))
r = cl.post(f'/admin/api/teams/{TID}/restore', json={'version': v_before + 1}, headers=H)
t('an older version can be restored too', 'NEW INSTRUCTIONS' in teams._member(teams.get(TID), 'measurement-surveyor')['instructions'])
j4 = start(title='FICTIONAL demo: after the change')
t('a new job runs on the new version', j4['team_version'] == teams.get(TID)['current_version'])
t('removing a member who works on a stage is refused', cl.delete(f'/admin/api/teams/{TID}/members/cost-surveyor', headers=H).status_code == 400)
r = cl.post(f'/admin/api/teams/{TID}/members', json={'role': 'Services Engineer', 'provider': 'claude'}, headers=H)
mid = next(m['id'] for m in r.json()['members'] if m['role'] == 'Services Engineer')
st = [dict(x) for x in teams.get(TID)['job_types'][0]['stages']]
st.insert(3, {'key': '', 'title': 'Check services', 'member': mid, 'task': 'Check the services allowance.', 'hands': 'Notes', 'checks': 'Priced items.'})
r = cl.put(f'/admin/api/teams/{TID}/job-types/{JT}', json={'stages': st}, headers=H)
t('stages and hand-offs can be edited on the page (versioned)', r.status_code == 200 and len(teams.get(TID)['job_types'][0]['stages']) == 7
  and 'stages added' in teams.versions(TID)[0]['what'])
t('a member can be removed once no stage uses it', cl.put(f'/admin/api/teams/{TID}/job-types/{JT}', json={'stages': [x for x in st if x['member'] != mid]}, headers=H).status_code == 200
  and cl.delete(f'/admin/api/teams/{TID}/members/{mid}', headers=H).status_code == 200)
t('the built-in steps keep their checks after editing', [x['handler'] for x in teams.get(TID)['job_types'][0]['stages']] == ['qs_plan', 'qs_measure', 'qs_price', 'qs_assemble', 'qs_trends', 'qs_adjust'])

# ---------------- Temple's suggestions need approval ----------------
def coach_reply(provider, system, messages):
    SENT_T.append((system, messages))
    return ('The Measurement Surveyor keeps leaving out page numbers.\nSuggested instructions:\nAlways give the page or schedule line for every quantity.\n'
            'Why: Two take-offs were sent back for missing sources.'), 'claude-haiku-4-5-20251001'
temple_discuss._complete = coach_reply
before = teams._member(teams.get(TID), 'measurement-surveyor')['instructions']
r = cl.post(f'/admin/api/teams/{TID}/members/measurement-surveyor/discussion', json={'message': 'How could the instructions be better?'}, headers=H)
t('Temple sees how the member\'s jobs went (send-backs included)', r.status_code == 200 and 'No measured items' in SENT_T[-1][1][0]['content'])
sid = r.json()['suggestion']
t('Temple\'s suggestion waits; the member is unchanged until Stefan approves', sid and teams._member(teams.get(TID), 'measurement-surveyor')['instructions'] == before
  and any(i['type'] == 'team_suggestion' for i in next(x for x in actions.summary()['sections'] if x['key'] == 'teams')['items']))
r = cl.post(f'/admin/api/teams/suggestions/{sid}', json={'action': 'approve'}, headers=H)
t('approved: applied as a new version naming Temple\'s suggestion', r.status_code == 200 and teams._member(teams.get(TID), 'measurement-surveyor')['instructions'].startswith('Always give the page')
  and 'Temple' in teams.versions(TID)[0]['what'])
r = cl.post(f'/admin/api/teams/{TID}/members/measurement-surveyor/discussion', json={'message': 'Again?'}, headers=H)
cl.post(f'/admin/api/teams/suggestions/{r.json()["suggestion"]}', json={'action': 'reject'}, headers=H)
t('a rejected suggestion changes nothing', teams.get(TID)['current_version'] == teams.versions(TID)[0]['version'] and 'Temple\'s suggestion, approved' in teams.versions(TID)[0]['what'])
t('the discussion is kept with the member', len(cl.get(f'/admin/api/teams/{TID}/members/measurement-surveyor/discussion').json()['messages']) == 4)

# ---------------- rules on member calls ----------------
t('a brief holding a secret is refused before anything starts', cl.post(f'/admin/api/teams/{TID}/jobs', json={'job_type': JT, 'title': 'x', 'brief': 'Use key sk-ant-api03-' + 'B' * 90 + ' for the build please now'}, headers=H).status_code == 400)
t('a protectively marked document is refused', cl.post(f'/admin/api/teams/{TID}/jobs', json={'job_type': JT, 'title': 'x', 'brief': 'A fictional hall for the community trust here.',
                                                                                          'uploads': [{'name': 'a.txt', 'text': 'OFFICIAL-SENSITIVE\nfloor areas'}]}, headers=H).status_code == 400)
m = teams.get(TID)['members'][0]
try: teams.call_model(m, {}, 'sys', 'Payload with key sk-ant-api03-' + 'C' * 90); ok = False
except rules_engine.RuleViolation: ok = True
t('every member call goes through check_outbound', ok)
real_spend = rules_engine.check_spend
rules_engine.check_spend = lambda kind: (_ for _ in ()).throw(rules_engine.RuleViolation('Spending cap reached: chat is paused.'))
j5 = start(title='FICTIONAL demo: over the cap')
t('the spending cap stops a member turn; the job says why and waits for Resume', j5['status'] == 'blocked' and 'Spending cap' in j5['error'])
rules_engine.check_spend = real_spend
agents.set_status('team-member', 'paused', 'testing')
r = cl.post(f'/admin/api/teams/jobs/{j5["id"]}/resume', headers=H).json()
t('a paused agent does not run', r['status'] == 'blocked' and 'paused' in r['error'])
agents.set_status('team-member', 'active')
r = cl.post(f'/admin/api/teams/jobs/{j5["id"]}/resume', headers=H).json()
t('resumed once the agent is active again', r['status'] == 'waiting' and r['pending'])
t('a job can be stopped; its waiting items leave Actions', cl.post(f'/admin/api/teams/jobs/{j5["id"]}/stop', headers=H).json()['status'] == 'stopped'
  and not any(i.get('job_id') == j5['id'] for i in teams.waiting()))

# client separation on member knowledge
s.create_category('Cost data')
knowledge.create('note', 'General cost note (fictional)', 'Fictional general guidance on hall costs and finishes.', 'Fictional', 'you', category='Cost data')
k2 = knowledge.create('note', 'Other client cost note (fictional)', 'Fictional client-only guidance on hall costs and finishes.', 'Fictional', 'you', category='Cost data')['id']
clients.tag('file', [k2], 'Fictional Other Council')
lead = dict(teams.get(TID)['members'][0], categories=['Cost data'])
kn = teams._knowledge({'id': 'x', 'title': 'hall costs', 'brief': 'hall costs finishes', 'client': ''}, lead, 'claude')
t('members read knowledge in their categories, never another client\'s', 'General cost note' in kn and 'client-only' not in kn)
t('no categories: no knowledge', teams._knowledge({'id': 'x', 'title': 'hall costs', 'brief': 'hall costs finishes', 'client': ''}, dict(lead, categories=[]), 'claude') == '')

# ---------------- provider failures are explained ----------------
class FakeProviderError(Exception):
    status_code = 400
    message = 'Your credit balance is too low.'
FakeProviderError.__module__ = 'anthropic._exceptions'
def failing(*a, **k): raise FakeProviderError('Your credit balance is too low.')
assistants._call = failing
teams.set_autonomy(TID, 'signoff')
j6 = start(title='FICTIONAL demo: provider refuses')
t('a provider\'s refusal is shown as the provider\'s own reason, not the exception name', j6['status'] == 'blocked' and 'FakeProviderError' not in j6['error'] and j6['error'])
assistants._call = fake_call
cl.delete(f'/admin/api/teams/{TID}/rates/{xbatch}', headers=H)
t('a rate library upload can be removed', not any(r_['code'] == 'X1' for r_ in team_qs.library(TID)))

# ---------------- arithmetic ----------------
c_ = team_qs.compute([{'ref': 'A', 'element': 'E', 'quantity': 3, 'rate': 33.33, 'rate_source': 'web'}, {'ref': 'B', 'element': 'E', 'quantity': 1, 'rate': 0.005, 'rate_source': 'library'},
                      {'ref': 'C', 'element': 'F', 'quantity': 10, 'rate': None, 'rate_source': 'unpriced'}], 1.0, {'prelims_pct': 10, 'contingency_pct': 5, 'fees_pct': 0})
t('compute: line amounts, subtotals and percentages in code, unpriced items left out', c_['construction'] == 99.99 + 0.01 and c_['prelims'] == 10.0
  and c_['contingency'] == 5.5 and c_['total'] == 115.5 and len(c_['lines']) == 2)

# ---------------- the Agents page can show what member turns touched ----------------
r = cl.get('/admin/api/agents/team-member/runs')
t('the Agents page lists member runs', r.status_code == 200 and r.json())
r = cl.get('/admin/api/agents/team-member/touched')
rows = r.json()['items']
t('Data touched shows the web pages the Cost Surveyor read', r.status_code == 200 and all(u in json.dumps(rows) for u in RET))

# ---------------- a decision that arrives while a run is finishing is not lost ----------------
teams.set_autonomy(TID, 'approve')
j7 = start(title='FICTIONAL demo: race')
with teams._LOCK: teams._ACTIVE.add(j7['id'])
teams._set(j7['id'], status='running')
with s.db() as c: c.execute("UPDATE team_steps SET status='approved' WHERE job_id=? AND status='pending'", (j7['id'],))
t('a second run while one is active is queued, not dropped', teams._advance(j7['id']) is False and j7['id'] in teams._AGAIN)
with teams._LOCK: teams._ACTIVE.discard(j7['id'])
teams._set(j7['id'], stage=1)
teams._advance_safe(j7['id'])
t('…and the job goes on', job(j7['id'])['pending'] and j7['id'] not in teams._AGAIN)
