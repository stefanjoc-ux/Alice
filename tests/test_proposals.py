"""Proposal writer: template sections, context under the rules, writer and QA agents, one automatic revision,
pricing from the rate card (cost never sent or printed), the Word document built from the template."""
import _util  # first: throwaway data folder, dummy keys, no real model calls
from _util import t
import base64, io, json, shutil, tempfile, time, zipfile
from pathlib import Path
from types import SimpleNamespace as NS
import substrate_store as s
s.init()
import knowledge as K, doc_library as L, organisations as O, clients as C, agents as A, assistants as AS
import proposals as P, proposal_docx as PD
import anthropic
import app
from fastapi.testclient import TestClient
cl = TestClient(app.app); H = {'x-admin-token': app.ADMIN_TOKEN}

FIX = Path(__file__).parent / 'fixtures' / 'proposal_template.docx'
root = Path(tempfile.mkdtemp(prefix='alice-docs-')) / 'Documents'
L.ROOT = root.resolve()
L.add_source('Proposal Templates', 'sharepoint', 'SharePoint: Bids site, Templates library')
shutil.copy(FIX, root / 'Proposal Templates' / 'Proposal-template.docx')

# ---------------- the template ----------------
raw = FIX.read_bytes()
o = PD.outline(raw)
titles = [x['title'] for x in o['sections']]
t('template headings become the sections, in order', titles[:3] == ['Executive summary', 'Understanding of your requirements', 'Proposed approach'] and len(titles) == 9)
t('text under a heading is guidance; [keep] marks standard text', o['sections'][0]['guidance'].startswith('[Two or three') and
  next(x for x in o['sections'] if x['title'] == 'About us')['keep'])
t('placeholders are found in the body, header and footer', {'title', 'client', 'date', 'reference', 'author'} <= set(o['placeholders']))

# ---------------- settings on the assistant ----------------
a = AS.get('proposal-writer')
t('a Proposal writer assistant is ready to set up', a['kind'] == 'proposal' and a['provider'] == 'claude_sonnet')
card = [{'role': 'Solution architect', 'unit': 'day', 'cost': 650, 'sell': '1,200'}, {'role': 'Consultant', 'unit': 'day', 'cost': 450, 'sell': 850}]
body = {'name': 'Proposal writer', 'provider': 'claude_sonnet', 'kind': 'proposal', 'guidance': 'Plain UK English.',
        'settings': {'template': 'Proposal Templates/Proposal-template.docx', 'rate_card': card, 'min_margin': 50,
                     'sections': [{'title': 'Executive summary', 'guidance': 'Mention value for money.'}, {'title': 'Risks', 'guidance': 'Top five risks with mitigations.'}]}}
r = cl.put('/admin/api/assistants/proposal-writer', headers=H, json=body)
t('rate card and template are saved (and the old Format and flow sections stay stored)', r.status_code == 200 and r.json()['settings']['rate_card'][0]['sell'] == 1200
  and r.json()['settings']['template'].endswith('Proposal-template.docx') and [x['title'] for x in r.json()['settings']['sections']] == ['Executive summary', 'Risks'])
for bad, why in [({'rate_card': [{'role': 'X', 'cost': 'abc', 'sell': 1}]}, 'number'), ({'template': '../../etc/passwd'}, 'template'),
                 ({'sections': [{'title': 'A'}, {'title': 'a'}]}, 'twice'), ({'rate_card': [{'role': 'X', 'cost': -1, 'sell': 1}]}, 'between')]:
    rr = cl.put('/admin/api/assistants/proposal-writer', headers=H, json=dict(body, settings=dict(body['settings'], **bad)))
    t(f'refused: {why}', rr.status_code == 400 and why in rr.json()['detail'])
t('the type cannot change once created', cl.put('/admin/api/assistants/hr-policy', headers=H, json={'name': 'HR policy assistant', 'kind': 'proposal'}).json()['kind'] == 'qa')
t('templates are listed from the document sources', any(x['name'] == 'Proposal-template.docx' and x['source'] == 'Proposal Templates'
                                                      for x in cl.get('/admin/api/proposal-templates', headers=H).json()['templates']))
st = cl.get('/assistant/proposal-writer/setup').json()
secs = [x['title'] for x in st['sections']]
t('the page starts from the template\'s sections only: no Format and flow added', secs == titles and 'Risks' not in secs
  and 'Mention value for money.' not in st['sections'][0]['guidance'] and all(x['source'] == 'template' for x in st['sections']))

# ---------------- context material ----------------
O.create('Northshire Council', 'council', client=True, aliases=['NSC'])
O.create('Southvale Council', 'council', client=True)
f1 = O.propose_fact('Northshire Council', 'purpose', 'Northshire is moving its data platform to Microsoft Fabric in 2027.', 'Public web', 'https://example.gov.uk/fabric')
O.review_facts([f1['id']], 'approved')
m_ok = s.propose('Fabric baseline approach', 'Our Fabric data security baseline starts with Purview labelling and DLP before migration.', 'Stefan said')['id']; s.review(m_ok, 'approved')
m_nsc = s.propose('Northshire stakeholders', 'Northshire Council data security lead wants a six week discovery for Fabric.', 'Stefan said')['id']; s.review(m_nsc, 'approved')
m_svc = s.propose('Southvale Fabric pricing', 'Southvale Council Fabric discovery was priced at forty days of consultancy.', 'Stefan said')['id']; s.review(m_svc, 'approved')
C.tag('memory', [m_nsc], 'Northshire Council'); C.tag('memory', [m_svc], 'Southvale Council')

# ---------------- mock models ----------------
calls = []
def writer_json(sections, revised=False, client='Northshire Council'):
    return json.dumps({'sections': [{'title': x, 'body': ('Revised. ' if revised else '') + f'{client} needs a secure Fabric baseline. {x} text with enough words to count as a real section of a proposal for the client.'}
                                    for x in sections],
                       'resource_plan': [{'role': 'Solution architect', 'quantity': 10, 'purpose': 'Design'}, {'role': 'Consultant', 'quantity': 15.2, 'purpose': 'Delivery'},
                                         {'role': 'Astronaut', 'quantity': 3}],
                       'gaps': ['Start date not given']})
QA1 = {'verdict': 'needs_revision', 'score': 62, 'summary': 'Close.', 'requirements': [{'requirement': 'Six week discovery', 'status': 'partly', 'where': 'Proposed approach'}],
       'issues': [{'severity': 'high', 'section': 'Proposed approach', 'issue': 'Discovery length not stated', 'fix': 'Say six weeks'}], 'strengths': ['Clear']}
QA2 = {'verdict': 'client_ready', 'score': 88, 'summary': 'Ready.', 'requirements': [{'requirement': 'Six week discovery', 'status': 'met', 'where': 'Proposed approach'}],
       'issues': [{'severity': 'low', 'section': 'Next steps', 'issue': 'Add owners', 'fix': 'Name roles'}], 'strengths': ['Specific']}
MODE = {'qa': [QA1, QA2]}
class Msgs:
    def create(self, **kw):
        calls.append(kw)
        usage = NS(model_dump=lambda: {'input_tokens': 1000, 'output_tokens': 400})
        if kw['system'].startswith('You are the proposal writer'):
            want = [ln[2:].split(' [KEEP]')[0] for ln in kw['messages'][0]['content'].split('SECTIONS\n')[1].split('\n\nRATE CARD')[0].split('\n') if ln.startswith('- ')]
            client = kw['messages'][0]['content'].split('CLIENT: ')[1].split('\n')[0]
            text = writer_json(want, 'PREVIOUS DRAFT' in kw['messages'][0]['content'], client)
        elif kw['system'].startswith('You are Argus'):
            text = json.dumps(MODE['qa'].pop(0) if MODE['qa'] else QA2)
        else:
            text = '{}'
        return NS(content=[NS(type='text', text=text)], usage=usage)
class Anth:
    def __init__(s_, **k): s_.messages = Msgs()
    def __enter__(s_): return s_
    def __exit__(s_, *a): pass
anthropic.Anthropic = Anth

def wait(pid, secs=20):
    end = time.time() + secs
    while time.time() < end:
        p = cl.get(f'/assistant/proposal-writer/proposals/{pid}').json()
        if p['status'] != 'running':
            import threading    # the job also logs and adds its cost as it finishes: let its thread end first
            for th in threading.enumerate():
                if th.name in ('proposal-' + pid[:6], 'proposal-recheck-' + pid[:6]): th.join(10)
            return cl.get(f'/assistant/proposal-writer/proposals/{pid}').json()
        time.sleep(0.1)
    return p

brief = ('Northshire Council wants a partner to set a data security baseline before moving to Microsoft Fabric. They want a six week '
         'discovery, a Purview labelling baseline and a roadmap. Responses by the end of October.')
r = cl.post('/assistant/proposal-writer/proposals', json={'title': 'Fabric security baseline', 'organisation': 'NSC', 'brief': brief,
                                                          'sections': st['sections'], 'rate_card': card})
t('a proposal starts in the background', r.status_code == 200 and r.json()['id'])
p = wait(r.json()['id'])
t('it finishes with a Word document', p['status'] == 'done' and p['document_id'] and not print(p.get('error') or ''))
t('the organisation is recognised by its other name and marked as a client', p['organisation'] == 'Northshire Council' and p['client'] == 'Northshire Council')
w = [c for c in calls if c['system'].startswith('You are the proposal writer')]
q = [c for c in calls if c['system'].startswith('You are Argus')]
t('writer, QA, one automatic revision, QA again', len(w) == 2 and len(q) == 2 and 'PREVIOUS DRAFT' in w[1]['messages'][0]['content']
  and 'Discovery length not stated' in w[1]['messages'][0]['content'])
t('the QA reports are kept: needs revision, then client ready', [x['verdict'] for x in p['qa']] == ['needs_revision', 'client_ready'] and p['qa'][1]['score'] == 88)
t('the draft is the revision', p['draft']['sections'][0]['body'].startswith('Revised.'))
ctx = w[0]['messages'][0]['content']
t('the writer gets the client profile and relevant memories', 'moving its data platform to Microsoft Fabric' in ctx and 'Fabric baseline approach' in ctx)
t('this client\'s tagged memories are included', 'Northshire stakeholders' in ctx)
t('another client\'s memories never are', 'Southvale' not in ctx and not any('Southvale' in c['messages'][0]['content'] for c in calls))
t('cost rates never reach a model', not any('650' in c['messages'][0]['content'] for c in w)
  and not any('1200' in c['messages'][0]['content'] or '1,200' in c['messages'][0]['content'] for c in w))
t('roles not on the rate card are dropped; quantities rounded to half days', [x['role'] for x in p['draft']['resource_plan']] == ['Solution architect', 'Consultant']
  and p['draft']['resource_plan'][1]['quantity'] == 15.0 and p['draft']['dropped_roles'] == ['Astronaut'])
pr = p['pricing']
t('Alice prices the plan from the sell rates', pr['sell'] == 10 * 1200 + 15 * 850 and pr['rows'][-1] == ['Total', '', '', '£24,750'])
t('cost and margin are worked out for you', pr['cost'] == 10 * 650 + 15 * 450 and round(pr['margin'], 1) == round((24750 - 13250) / 24750 * 100, 1))
t('margin below your minimum is flagged', any('below your minimum of 50%' in x for x in pr['warnings']))
t('QA sees sell prices only', '£24,750' in q[0]['messages'][0]['content'] and '£13,250' not in q[0]['messages'][0]['content'] and '650' not in q[0]['messages'][0]['content'])
t('what Alice used is listed for you', p['context']['used']['organisation'] == 'Northshire Council' and 'Fabric baseline approach' in p['context']['used']['memories'])

import documents as D
doc = D.get(p['document_id'])['data']
text = K.docx_to_text(doc)
t('the document is the template, filled in', 'HEADER: Northshire Council | Fabric security baseline' in text and 'Prepared for Northshire Council' in text)
t('every template section is there in order, nothing added after them', text.index('# Executive summary') < text.index('# Proposed approach') and '# Risks' not in text)
t('standard [keep] text is kept word for word', 'We are an independent technology partner' in text and '[keep]' not in text)
t('the pricing table sits in the Commercials section, cost never appears', 'Row 4: Total |  |  | £24,750' in text and '£650' not in text and '13,250' not in text and 'All prices exclude VAT.' in text)
t('no guidance or placeholders are left', '[Two or three' not in text and '{{' not in text)
with zipfile.ZipFile(io.BytesIO(doc)) as z:
    t('the template look is untouched (styles, header, footer)', z.read('word/styles.xml') == zipfile.ZipFile(io.BytesIO(raw)).read('word/styles.xml')
      and any(n.startswith('word/header') for n in z.namelist()))
with s.db() as c:
    t('the document is not knowledge', not c.execute("SELECT 1 FROM files WHERE name LIKE 'Fabric-security%'").fetchone())
    t('the proposal is logged', c.execute("SELECT 1 FROM activity WHERE action='proposal_written' AND detail LIKE '%client_ready%'").fetchone() is not None)
runs = {x: A.runs(x)['runs'] for x in (P.WRITER, P.QA)}
t('both agents ran as tracked runs (2 writer, 2 QA)', len(runs[P.WRITER]) == 2 and len(runs[P.QA]) == 2)
touched = A.touched_items(P.WRITER)
t('the writer\'s Data touched shows the template, memories and the client profile, sent to Sonnet',
  any(i['target_type'] == 'document' for i in touched) and any(i['target_id'] == m_ok for i in touched)
  and all('Claude Sonnet 5.5 (Anthropic)' in i['sent_to'] for i in touched if i['read']))

# ---------------- QA passes first time: no revision ----------------
calls.clear(); MODE['qa'] = [QA2]
p2 = wait(cl.post('/assistant/proposal-writer/proposals', json={'title': 'Second proposal', 'organisation': 'Someone New Ltd', 'brief': brief,
                                                                  'use_memory': False, 'sections': [{'title': 'Summary'}, {'title': 'Pricing'}], 'rate_card': card}).json()['id'])
t('client ready first time: no revision', p2['status'] == 'done' and len(p2['qa']) == 1 and sum(1 for c in calls if c['system'].startswith('You are the proposal writer')) == 1)
t('an organisation not on the Organisations page is used as typed, not as a client', p2['organisation'] == 'Someone New Ltd' and p2['client'] == '')
t('with memory off, no memories are sent', 'APPROVED MEMORIES' not in calls[0]['messages'][0]['content'])

# ---------------- Alice's own checks ----------------
chk = P.alice_checks([{'title': 'A', 'body': 'We worked with Southvale Council on this. Price £5,000. [Insert case study]', 'keep': False},
                      {'title': 'B', 'body': '', 'keep': False}], 'Northshire Council')
kinds = ' '.join(c['issue'] for c in chk)
t('Alice flags other clients named, prices in text, placeholders and empty sections',
  'another client: Southvale Council' in kinds and 'Prices in the text' in kinds and 'Placeholder' in kinds and 'is empty' in kinds)

# ---------------- refusals ----------------
def start(**k):
    return cl.post('/assistant/proposal-writer/proposals', json={'title': 'Test', 'organisation': '', 'brief': brief, **k})
t('a brief that is too short is refused', start(brief='Fix it please').status_code == 400)
t('a secret in the brief is refused before anything is sent', start(brief=brief + ' key sk-proj-abcdefghijklmnopqrstuvwxyz0123456789ABCD').status_code == 400)
t('protectively marked briefs are refused', start(brief='OFFICIAL-SENSITIVE\n' + brief).status_code == 400)
t('cross-origin requests are refused', cl.post('/assistant/proposal-writer/proposals', headers={'origin': 'http://evil.example'},
                                               json={'title': 'T', 'brief': brief}).status_code == 403)
A.set_status(P.QA, 'paused', 'test')
t('a paused agent stops new proposals', start().status_code == 503)
A.set_status(P.QA, 'active')
t('proposal writers do not answer questions', cl.post('/assistant/proposal-writer/ask', json={'question': 'hello there'}).status_code == 400)

# failure is reported, not hidden
class Broken:
    def __init__(s_, **k): s_.messages = s_
    def create(s_, **k): return NS(content=[NS(type='text', text='not json at all')], usage=NS(model_dump=lambda: {'input_tokens': 1, 'output_tokens': 1}))
    def __enter__(s_): return s_
    def __exit__(s_, *a): pass
anthropic.Anthropic = Broken
pf = wait(start().json()['id'])
t('an unusable model answer fails the proposal with a clear message', pf['status'] == 'failed' and 'usable answer' in pf['error'])
anthropic.Anthropic = Anth

lst = cl.get('/assistant/proposal-writer/proposals').json()['proposals']
t('recent proposals are listed with their QA verdict', lst and any(x['verdict'] == 'client_ready' for x in lst))
t('another assistant cannot read a proposal', cl.get(f'/assistant/hr-policy/proposals/{p["id"]}').status_code == 404)
page = cl.get('/assistant/proposal-writer').text
t('the proposal page has the brief, notes and rate card, and no Format and flow', 'id="brief"' in page and 'id="notes"' in page and 'Rate card' in page and 'Format and flow' not in page)
_as = cl.get('/admin/assistants').text
t('the Assistants page offers the proposal settings, without default sections', 'Rate card and price book' in _as and 'Format and flow' not in _as
  and 'PE.sections' not in _as and 'sections:S.sections||[]' in _as)

# no template: Alice's own layout
cl.put('/admin/api/assistants/proposal-writer', headers=H, json=dict(body, settings=dict(body['settings'], template='')))
MODE['qa'] = [QA2]
p3 = wait(start(sections=[{'title': 'Summary'}, {'title': 'Investment'}], rate_card=card).json()['id'])
t3 = K.docx_to_text(D.get(p3['document_id'])['data']) if p3['status'] == 'done' else ''
t('without a template, Alice\'s own layout with the same sections and pricing', '# Summary' in t3 and '£24,750' in t3 and not print(p3.get('error') or ''))

# numbered lists restart in each section (Word would otherwise carry on counting)
import re as _re
out = PD.fill(raw, [{'title': 'A', 'body': '1. one\n2. two'}, {'title': 'B', 'body': '1. again'}], {})
with zipfile.ZipFile(io.BytesIO(out)) as z:
    ids = _re.findall(r'<w:numId w:val="(\d+)"', z.read('word/document.xml').decode())
    t('each numbered list restarts at 1', len(set(ids)) == 2 and z.read('word/numbering.xml').decode().count('<w:startOverride w:val="1"/>') == 2)
t('placeholders split across runs are still filled', 'Prepared for Acme' in K.docx_to_text(PD.fill(raw, [], {'client': 'Acme'})))

# ---------------- choosing a better model ----------------
t('premium models are offered for proposals, with rough costs', {'claude_opus', 'openai_astra'} <= {m['key'] for m in cl.get('/assistant/proposal-writer/setup').json()['models'] if m['premium']}
  and all(m['writer_cost'] is not None for m in cl.get('/assistant/proposal-writer/setup').json()['models']))
t('a question-answering assistant cannot use a premium model', cl.put('/admin/api/assistants/hr-policy', headers=H, json={'name': 'HR policy assistant', 'provider': 'claude_opus'}).status_code == 400)
t('provider families for the rules', AS.family('openai_astra') == 'openai' and AS.family('claude_opus') == 'claude')
cl.put('/admin/api/assistants/proposal-writer', headers=H, json=dict(body, settings=dict(body['settings'], template='Proposal Templates/Proposal-template.docx')))
calls.clear(); MODE['qa'] = [QA2]
p4 = wait(start(writer_model='claude_opus', qa_model='claude_sonnet', sections=[{'title': 'Summary'}], rate_card=card).json()['id'])
wm = [c['model'] for c in calls if c['system'].startswith('You are the proposal writer')]
qm = [c['model'] for c in calls if c['system'].startswith('You are Argus')]
t('each proposal can use another writer and QA model', p4['status'] == 'done' and wm == ['claude-opus-5-5'] and qm == ['claude-sonnet-5-5']
  and p4['inputs']['writer'] == 'claude_opus')
t('unknown models are refused', start(writer_model='gpt-99').status_code == 400)
t('Opus is recorded as where the data went', any('Claude Opus 5.5 (Anthropic)' in i['sent_to'] for i in A.touched_items(P.WRITER)))
t('the Agents page names the model in use', 'Claude Sonnet 5.5 (Proposal writer)' in next(a for a in cl.get('/admin/api/agents', headers=H).json()['agents'] if a['id'] == P.WRITER)['anatomy_live']['model'])
import openai
seen_o = []
class ORsp:
    def create(self, **kw):
        seen_o.append(kw)
        return NS(output_text='{}', usage=NS(model_dump=lambda: {'input_tokens': 10, 'output_tokens': 5}))
class OAI:
    def __init__(s_, **k): s_.responses = ORsp()
    def __enter__(s_): return s_
    def __exit__(s_, *a): pass
openai.OpenAI = OAI
AS._call('openai_astra', 'sys', [{'role': 'user', 'content': 'x'}])
t('GPT-6 Astra is called with reasoning on (it has no "none" setting)', seen_o[-1]['model'] == 'gpt-6-astra' and seen_o[-1]['reasoning']['effort'] == 'medium')

# ---------------- no pasted structure; older custom sections; re-checks and QA of your own document ----------------
anthropic.Anthropic = Anth
calls.clear(); MODE['qa'] = [QA2]
r = start(structure='1. Why now\n- budget pressure\n2. Approach', rate_card=card)
pn = wait(r.json()['id'])
wmsg = [c for c in calls if c['system'].startswith('You are the proposal writer')][0]['messages'][0]['content']
t('a new proposal takes the template\'s sections; a structure sent anyway is ignored', pn['status'] == 'done'
  and [x['title'] for x in pn['inputs']['sections']] == titles and [x['title'] for x in pn['draft']['sections']] == titles
  and 'STRUCTURE AND POINTS FROM THE AUTHOR' not in wmsg and 'Why now' not in wmsg and 'structure' not in pn['inputs'])
calls.clear(); MODE['qa'] = [QA2]
r = start(sections=[{'title': 'Summary', 'include': '- Fixed price\n- Starts in January'}, {'title': 'Pricing'}], rate_card=card)
p5 = wait(r.json()['id'])
with s.db() as c_:                      # an older proposal: written with its own sections and a pasted structure (before 7 Oct 2026)
    c_.execute('UPDATE proposals SET inputs=? WHERE id=?', (json.dumps(dict(p5['inputs'], structure='1. Why now\n- budget pressure')), p5['id']))
p5 = cl.get(f'/assistant/proposal-writer/proposals/{p5["id"]}').json()
wmsg = [c for c in calls if c['system'].startswith('You are the proposal writer')][0]['messages'][0]['content']
t('an older proposal with its own sections still opens with them', [x['title'] for x in p5['draft']['sections']] == ['Summary', 'Pricing']
  and p5['inputs']['structure'].startswith('1. Why now') and 'Include: - Fixed price' in wmsg)
t('and its Word document has those sections', all(x in K.docx_to_text(D.get(p5['document_id'])['data']) for x in ('Summary', 'Pricing')))

calls.clear(); MODE['qa'] = [QA2]
edited = [{'title': x['title'], 'body': ('My edited summary with the January start date and a fixed price for the council. ' * 3) if x['title'] == 'Summary' else x['body']}
          for x in p5['draft']['sections']]
r = cl.post(f'/assistant/proposal-writer/proposals/{p5["id"]}/recheck', json={'sections': edited})
p6 = wait(p5['id'])
q6 = [c for c in calls if c['system'].startswith('You are Argus')]
t('your edits go back to QA and the document is rebuilt', r.status_code == 200 and p6['status'] == 'done' and len(p6['qa']) == len(p5['qa']) + 1
  and p6['qa'][-1]['source'] == 'your edits' and len(q6) == 1 and 'My edited summary' in q6[0]['messages'][0]['content']
  and p6['document_id'] != p5['document_id'] and 'My edited summary' in K.docx_to_text(D.get(p6['document_id'])['data']))
t('no writer run for a re-check', not [c for c in calls if c['system'].startswith('You are the proposal writer')])
t('every section must come back, titles unchanged', cl.post(f'/assistant/proposal-writer/proposals/{p5["id"]}/recheck', json={'sections': edited[:1]}).status_code == 400)
t('another assistant cannot re-check it', cl.post(f'/assistant/hr-policy/proposals/{p5["id"]}/recheck', json={'sections': edited}).status_code == 404)

mine = D.get(D.create('docx', 'My version', '# Summary\nOur own words for the council: fixed price, January start, governance first. ' * 2
                      + '\n\n# Approach\nThree phases over six weeks with workshops and a roadmap at the end of it all.')['id'])['data']
calls.clear(); MODE['qa'] = [QA1]
r = cl.post(f'/assistant/proposal-writer/proposals/{p5["id"]}/qa-upload', json={'name': 'Fabric-v2.docx', 'data': base64.b64encode(mine).decode()})
p7 = wait(p5['id'])
q7 = [c for c in calls if c['system'].startswith('You are Argus')]
t('a revised version uploaded from Word is checked against the brief', r.status_code == 200 and p7['qa'][-1]['source'] == 'uploaded: Fabric-v2.docx'
  and 'Our own words for the council' in q7[0]['messages'][0]['content'] and [x['title'] for x in p7['qa'][-1]['sections']] == ['Summary', 'Approach'])
t('the uploaded document is not kept, and the draft is unchanged', p7['draft'] == p6['draft'] and p7['document_id'] == p6['document_id'])
t('a secret in an uploaded version is refused', cl.post(f'/assistant/proposal-writer/proposals/{p5["id"]}/qa-upload',
  json={'name': 'k.txt', 'data': base64.b64encode(('# A\nkey sk-proj-abcdefghijklmnopqrstuvwxyz0123456789ABCD ' + 'x ' * 200).encode()).decode()}).status_code == 400)

calls.clear(); MODE['qa'] = [QA2]
r = cl.post('/assistant/proposal-writer/proposals/qa-only', json={'title': 'Existing bid', 'organisation': 'NSC', 'brief': brief, 'name': 'bid.docx',
                                                                    'data': base64.b64encode(mine).decode(), 'qa_model': 'claude_sonnet'})
p8 = wait(r.json()['id'])
t('QA only: a proposal you already have is checked without writing anything', r.status_code == 200 and p8['status'] == 'done' and len(p8['qa']) == 1
  and not p8['document_id'] and p8['inputs']['qa_only'] and p8['organisation'] == 'Northshire Council'
  and not [c for c in calls if c['system'].startswith('You are the proposal writer')] and [x['title'] for x in p8['draft']['sections']] == ['Summary', 'Approach'])
t('QA only needs a brief', cl.post('/assistant/proposal-writer/proposals/qa-only', json={'title': 'X bid', 'brief': 'too short', 'name': 'a.docx',
                                                                                           'data': base64.b64encode(mine).decode()}).status_code == 400)
t('a QA-only proposal cannot be "edited and re-checked"', cl.post(f'/assistant/proposal-writer/proposals/{p8["id"]}/recheck',
  json={'sections': [{'title': 'Summary', 'body': 'x'}, {'title': 'Approach', 'body': 'y'}]}).status_code == 400)
secs = P.document_sections('t.docx', PD.fill(raw, [{'title': 'Executive summary', 'body': 'A ' * 150}, {'title': 'Proposed approach', 'body': 'B ' * 150}], {'title': 'T', 'client': 'C'}))
t('a document is split at its main headings', [x['title'] for x in secs][-2:] == ['Executive summary', 'Proposed approach'])
page = cl.get('/assistant/proposal-writer').text
t('the page offers QA-only mode and re-checks', all(x in page for x in ('Check one I already have', 'Edit the draft and check again', 'Upload a revised version for QA')))
t('no Format and flow: no section editor and no structure box', not any(x in page for x in ('Format and flow', 'id="structure"', 'Turn into sections', 'Paste a structure', 'PE.sections', 'id="secs"')))
t('the draft is shown as one document with editing in place', all(x in page for x in ('function drawDoc', "'doc-sec'", "'doc-ed'", 'Click into any section', 'From the template')))

# pasting a rate card table (the parser runs in the browser; checked with Node when it is installed)
import shutil, subprocess, proposal_ui
node = shutil.which('node')
if node:
    js = Path(tempfile.mkdtemp()) / 'pe.js'
    js.write_text('const document={};' + proposal_ui.PE_JS + 'module.exports=PE;', encoding='utf-8')
    out = subprocess.run([node, str(Path(__file__).parent / 'js' / 'rate_paste_test.js'), str(js)], capture_output=True, text=True, timeout=60).stdout
    t('a pasted rate card table is read: headings, order, units, £ and commas, markdown, no headings', out.strip().endswith('7 passed 0 failed'))
# the draft as one document, edited in place (the page's own functions, run with Node against a small stand-in for the browser)
SHIM = r"""
let FOCUS=null;
class Txt{constructor(t){this.nodeType=3;this.data=String(t)}get textContent(){return this.data}}
class El{constructor(tag){this.tagName=tag.toUpperCase();this.children=[];this.kids=[];this.attrs={};this.style={};this.dataset={};this.hidden=false;this._c=new Set();this.scrollHeight=40;
  const me=this;this.classList={add:(...k)=>k.forEach(x=>me._c.add(x)),remove:(...k)=>k.forEach(x=>me._c.delete(x)),contains:k=>me._c.has(k),toggle:(k,on)=>{on=on===undefined?!me._c.has(k):!!on;on?me._c.add(k):me._c.delete(k);return on}}}
 set className(v){this._c=new Set(String(v).split(/\s+/).filter(Boolean))}get className(){return [...this._c].join(' ')}
 get textContent(){return this.kids.map(k=>k.textContent).join('')}set textContent(v){this.kids=[new Txt(v)];this.children=[]}
 append(...n){for(let x of n){if(typeof x==='string')x=new Txt(x);this.kids.push(x);if(x instanceof El){x.parentNode=this;this.children.push(x)}}}
 prepend(...n){const k=this.kids;this.kids=[];this.children=this.children.filter(()=>false);this.append(...n);for(const x of k)this.append(x)}
 replaceChildren(...n){this.kids=[];this.children=[];this.append(...n)}
 setAttribute(k,v){this.attrs[k]=String(v)}getAttribute(k){return this.attrs[k]??null}removeAttribute(k){delete this.attrs[k]}
 focus(){FOCUS=this}scrollIntoView(){}}
const document={createElement:t=>new El(t),createTextNode:t=>new Txt(t),querySelectorAll:()=>[],querySelector:()=>({scrollTop:0})};
"""
page_js = '\n'.join(__import__('re').findall(r'<script>(.*?)</script>', cl.get('/assistant/proposal-writer').text, __import__('re').S))
def _between(src, a, b): i = src.index(a); return src[i:src.index(b, i)]
doc_js = (_between(page_js, '// ---------- the draft as one document', '// ---------- Parker: work on the form')
          + _between(page_js, 'function draftNow(){', 'function form(){')
          + _between(page_js, 'function editDraft(p){', '// ---------- proposals on the go'))
if node:
    js = Path(tempfile.mkdtemp()) / 'doc.js'
    js.write_text(SHIM + proposal_ui.PE_JS + """
const mk=PE.mk;let RO=false,DOC=null,EDS=null,CUR=null,DIRTY=false;const CALLS={api:[],follow:[]};
const $=id=>({scrollIntoView(){},value:''}),priceDiff=()=>[],updateBar=()=>{},when=t=>t;const api=async(path,method,body)=>{CALLS.api.push([path,method,body]);return {}};
const follow=id=>CALLS.follow.push(id),show=()=>{},alertBox=m=>console.log('ALERT',m);
""" + doc_js + """
module.exports={drawDoc:p=>{CUR=p;return drawDoc(p)},editDraft,calls:()=>CALLS,focused:()=>FOCUS,get:k=>({DOC,EDS,RO})[k],set:o=>{if('DOC' in o)DOC=o.DOC;if('EDS' in o)EDS=o.EDS;if('RO' in o)RO=o.RO}};
""", encoding='utf-8')
    res = subprocess.run([node, str(Path(__file__).parent / 'js' / 'draft_doc_test.js'), str(js)], capture_output=True, text=True, timeout=60)
    out = (res.stdout + res.stderr).strip()
    t('the draft renders as one document and edits in place save to the right sections (12 browser checks)', out.endswith('12 passed 0 failed') or print(out))
t('the rate card offers Paste a table', 'Paste a table' in proposal_ui.PE_JS and 'parseRates' in proposal_ui.PE_JS)

# ---------------- choosing the template on each proposal ----------------
cl.put('/admin/api/assistants/proposal-writer', headers=H, json=dict(body, settings=dict(body['settings'], template='Proposal Templates/Proposal-template.docx')))
st2 = cl.get('/assistant/proposal-writer/setup').json()
t('the page lists the templates in the document sources', any(x['path'].replace('\\', '/') == 'Proposal Templates/Proposal-template.docx' for x in st2['templates']))
o = cl.get('/assistant/proposal-writer/outline', params={'template': 'Proposal Templates/Proposal-template.docx'}).json()
t('picking a template gives its sections only', [x['title'] for x in o['sections']] == titles and 'client' in o['placeholders'])
t('no template: Alice\'s own layout', [x['title'] for x in cl.get('/assistant/proposal-writer/outline', params={'template': ''}).json()['sections']] == list(P.OWN_LAYOUT))
t('a path outside the document sources is refused', cl.get('/assistant/proposal-writer/outline', params={'template': '../../secret.docx'}).status_code == 400
  and start(template='../x.docx').status_code == 400)
calls.clear(); MODE['qa'] = [QA2]
p9 = wait(start(template='', rate_card=card).json()['id'])
t('a proposal can skip the template: Alice\'s own layout and sections', p9['status'] == 'done' and p9['inputs']['template'] == ''
  and [x['title'] for x in p9['draft']['sections']] == list(P.OWN_LAYOUT)
  and 'HEADER:' not in K.docx_to_text(D.get(p9['document_id'])['data']) and 'Our approach' in K.docx_to_text(D.get(p9['document_id'])['data']))
calls.clear(); MODE['qa'] = [QA2]
p10 = wait(start(sections=[{'title': 'Executive summary'}], rate_card=card).json()['id'])
t('without a choice, the assistant\'s default template is used', p10['inputs']['template'].endswith('Proposal-template.docx') and 'HEADER:' in K.docx_to_text(D.get(p10['document_id'])['data']))
t('the page has the template picker', 'id="tplsel"' in cl.get('/assistant/proposal-writer').text)

# ---------------- price book: ticked roles, fixed days, target margin, spreadsheets ----------------
book = [{'role': 'Solution architect', 'unit': 'day', 'cost': 650, 'sell': 1200, 'days': 4, 'use': True},
        {'role': 'Consultant', 'unit': 'day', 'cost': 450, 'sell': 850, 'use': False},
        {'role': 'Project manager', 'unit': 'day', 'cost': 400, 'sell': 700, 'days': '3', 'override': True}]
cb = P.clean_rate_card(book)
t('rate card rows keep fixed days, the tick and the override', cb[0]['days'] == 4 and cb[1]['days'] is None and not cb[1]['use']
  and cb[2]['days'] == 3 and cb[2]['use'] and cb[2]['override'] and not cb[0]['override'])
t('only ticked roles are used', [r['role'] for r in P.used(cb)] == ['Solution architect', 'Project manager'])
t('older rate cards without ticks use every role', len(P.used(P.clean_rate_card(card))) == 2)
t('sell rate from a target margin on sell', P.sell_for(700, 30) == 1000 and P.sell_for(650, 0) == 650)
calls.clear(); MODE['qa'] = [QA2]
pb = wait(start(sections=[{'title': 'Approach'}], rate_card=book).json()['id'])
wp = [c for c in calls if c['system'].startswith('You are the proposal writer')][0]['messages'][0]['content']
t('the writer sees only the ticked roles, with fixed days', '- Consultant' not in wp and 'Solution architect (per day) FIXED: 4 days' in wp and 'Project manager' in wp)
plan = {x['role']: x['quantity'] for x in pb['draft']['resource_plan']}
t('your fixed days beat the writer\'s plan; fixed roles it missed are added', plan.get('Solution architect') == 4 and plan.get('Project manager') == 3 and 'Consultant' not in plan)
t('pricing uses the ticked roles and fixed days', pb['pricing']['sell'] == 4 * 1200 + 3 * 700 and pb['pricing']['cost'] == 4 * 650 + 3 * 400)
sv = cl.put('/admin/api/assistants/proposal-writer', headers=H, json=dict(body, settings=dict(body['settings'], target_margin=35, auto_approve_references=False)))
t('target margin and the auto-approve switch are saved', sv.status_code == 200 and sv.json()['settings']['target_margin'] == 35
  and sv.json()['settings']['auto_approve_references'] is False)
t('the auto-approve switch is on unless you turn it off', P.clean_settings({})['auto_approve_references'] is True and P.clean_settings({})['target_margin'] == 30)
t('the page gets the target margin', cl.get('/assistant/proposal-writer/setup').json().get('target_margin') == 35)
cl.put('/admin/api/assistants/proposal-writer', headers=H, json=body)

import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Notes'; ws.append(['Read me'])
ws2 = wb.create_sheet('Rates'); ws2.append(['Pricing tool 2026']); ws2.append([])
ws2.append(['Role', 'Grade', 'Internal cost (day)', 'Day rate']); ws2.append(['Solution architect', 'SA', 650, 1200])
ws2.append(['Data engineer', 'SC', '£550', '£1,000']); ws2.append(['Data engineer', 'SC', 1, 2]); ws2.append(['Total', '', 1200, 2200])
bio = io.BytesIO(); wb.save(bio)
xr = P.rates_from_sheet('pricing.xlsx', bio.getvalue())
t('a pricing spreadsheet is read: right sheet, heading after a title, £ and commas, duplicates and totals skipped',
  xr['sheet'] == 'Rates' and xr['rows'] == [{'role': 'Solution architect', 'unit': 'day', 'cost': 650, 'sell': 1200},
                                            {'role': 'Data engineer', 'unit': 'day', 'cost': 550, 'sell': 1000}])
cr = P.rates_from_sheet('p.csv', b'Resource;Cost per hour;Sell per hour\nTester;40;75\nAnalyst;45;\n')
t('a CSV is read, with units from the headings and missing sell rates left blank',
  cr['rows'] == [{'role': 'Tester', 'unit': 'hour', 'cost': 40, 'sell': 75}, {'role': 'Analyst', 'unit': 'hour', 'cost': 45, 'sell': ''}])
pr_ = lambda name, raw, **h: cl.post('/assistant/proposal-writer/rates/parse', headers=h, json={'name': name, 'data': base64.b64encode(raw).decode()})
t('the page can load a pricing spreadsheet', pr_('pricing.xlsx', bio.getvalue()).json()['rows'][1]['role'] == 'Data engineer')
t('a sheet with no rates, a broken file and other types are refused', pr_('x.csv', b'a,b\n1,2\n').status_code == 400
  and pr_('x.xlsx', b'not a workbook').status_code == 400 and pr_('x.pdf', b'%PDF').status_code == 400)
t('cross-origin spreadsheet loads are refused', pr_('pricing.xlsx', bio.getvalue(), origin='http://evil.example').status_code == 403)
t('the HR assistant has no spreadsheet loading', cl.post('/assistant/hr-policy/rates/parse', json={'name': 'p.csv', 'data': base64.b64encode(b'Role,Cost\nA,1\n').decode()}).status_code == 404)
t('the rate card has the target margin, ticks, margin column and reference totals',
  all(x in proposal_ui.PE_JS for x in ('pe-target', 'pe-rtotal', 'Load a pricing spreadsheet', 'Apply to every role')))
pg_ = cl.get('/assistant/proposal-writer').text
t("Parker's logo is on the page and as the tab icon", 'class="parker"' in pg_ and 'image/svg+xml' in pg_ and 'id="go2"' in pg_)

# ---------------- a pricing tool with cost and sell sheets by country (hourly) ----------------
wb2 = openpyxl.Workbook(); p1 = wb2.active; p1.title = 'Personas'
p1.append(['IT Capability', 'Level', 'Persona Title']); [p1.append(['Advisory & Strategy', i, f'Persona {i}']) for i in range(1, 7)]
for title, rates in (('StdCosts', {'Strategic Investment Executive': 149.4667, 'Business Value Advisor': 80.2667, 'Offshore Analyst': None}),
                     ('SellPricing', {'Strategic Investment Executive': 233.7333, 'Business Value Advisor': 186.6667, 'Offshore Analyst': 43.3333})):
    w = wb2.create_sheet(title); w.append(['', 'EMEA Solutions Delivery'])
    w.append(['', '', '', '', '', '', '', 8, 8, 7.5]); w.append([])
    w.append(['', '', '', '', '', '', '', 'AT', 'DE', 'UK']); w.append(['', '', '', '', '', '', '', 'EUR', 'EUR', 'GBP'])
    w.append(['', 'Resource group', 'IT Capability', 'Persona', 'Role', 'Pricing Tool role?']); w.append(['', '(blank)', '(blank)', '(blank)'])
    for k, v in rates.items(): w.append(['', 'EMEA (SFIA)', 'Advisory & Strategy', k, 'EMEA-' + k, 'y', None, 99, 99, v])
bio2 = io.BytesIO(); wb2.save(bio2)
mx = P.rates_from_sheet('scheduler.xlsx', bio2.getvalue())
t('a pricing tool with cost and sell sheets: UK hourly rates become day rates (x 7.5 hours)', 'StdCosts and SellPricing' in mx['sheet']
  and mx['rows'][0] == {'role': 'Strategic Investment Executive', 'unit': 'day', 'cost': 1121.0, 'sell': 1753.0}
  and mx['rows'][2] == {'role': 'Offshore Analyst', 'unit': 'day', 'cost': '', 'sell': 325.0} and len(mx['rows']) == 3)
lv = openpyxl.Workbook(); lw = lv.active; lw.append(['IT Capability', 'Level', 'Persona Title', 'Pay band'])
for i in range(1, 7): lw.append(['Advisory & Strategy', i, f'Persona {i}', i])
bio3 = io.BytesIO(); lv.save(bio3)
try: P.rates_from_sheet('levels.xlsx', bio3.getvalue()); lvl_ok = False
except ValueError: lvl_ok = True
t('levels and grades (1 to 6) are not mistaken for rates', lvl_ok)
t('the rate card can be cleared, and a new spreadsheet replaces or adds', 'Remove every role from this rate card' in proposal_ui.PE_JS and 'OK replaces them' in proposal_ui.PE_JS)
t('the rate card has a tick-all / untick-all box', 'pe-all' in proposal_ui.PE_JS and 'Tick or untick every role shown' in proposal_ui.PE_JS)
cl_ = P.clean_rate_card([{'role': 'A', 'cost': 100, 'sell': 160, 'list': '200'}, {'role': 'B', 'cost': 100, 'sell': 160}])
t('the price book rate is kept beside the adjusted sell rate', cl_[0]['list'] == 200 and cl_[0]['sell'] == 160 and cl_[1]['list'] is None)
t('the rate card shows the price book rate and the change', all(x in proposal_ui.PE_JS for x in ("'Price book'", "'Change'", 'At price book', 'Adjusted sell price')))
pgx = cl.get('/assistant/proposal-writer').text
t('the rate card spans the page below Parker; the button still writes the proposal', 'id="full"' in pgx and 'form="f"' in pgx)

# ---------------- proposals in progress: saved as you work, a list to switch between ----------------
cl.put('/admin/api/assistants/proposal-writer', headers=H, json=body)
wsave = lambda **k: cl.post('/assistant/proposal-writer/work', json=k)
r1 = wsave(form={'title': 'Fabric baseline', 'organisation': 'NSC', 'brief': 'Six-week baseline.', 'rate_card': card + [{'role': 'Trainer', 'cost': 300, 'sell': 600, 'use': False}],
                 'sections': [{'title': 'Approach'}], 'references': [], 'template': ''}).json()
t('a proposal form saves itself: a new one gets an id', r1['created'] and len(r1['id']) == 32)
r2 = wsave(id=r1['id'], form={'title': 'Fabric data security baseline', 'brief': 'Six-week baseline, then a roadmap.', 'rate_card': card}).json()
pf = P.get(r1['id'])
t('later saves update the same proposal', r2['id'] == r1['id'] and not r2['created'] and pf['title'] == 'Fabric data security baseline' and pf['status'] == 'form'
  and pf['inputs']['form'] and len(pf['inputs']['rate_card']) == 2)
t('a half-typed rate keeps the last good rate card', wsave(id=r1['id'], form={'title': 'x', 'rate_card': [{'role': 'A', 'cost': 'abc', 'sell': 1}]}).status_code == 200
  and len(P.get(r1['id'])['inputs']['rate_card']) == 2)
wsave(id=r1['id'], form={'title': 'Fabric data security baseline', 'organisation': 'NSC', 'brief': 'Six-week baseline, then a roadmap.', 'rate_card': card})
t('secrets are not saved', wsave(id=r1['id'], form={'title': 'x', 'brief': 'key sk-proj-abcdefghijklmnopqrstuvwxyz0123456789ABCD'}).status_code == 400
  and P.get(r1['id'])['brief'] == 'Six-week baseline, then a roadmap.')
r3 = wsave(form={'title': 'Copilot readiness', 'brief': 'Review.'}).json()
lst = cl.get('/assistant/proposal-writer/proposals').json()['proposals']
t('the list shows proposals in progress, newest first', [x['id'] for x in lst][:2] == [r3['id'], r1['id']] and lst[0]['status'] == 'form')
t('a proposal in progress can be removed from the list', cl.post(f'/assistant/proposal-writer/work/{r3["id"]}/discard', json={}).status_code == 200
  and r3['id'] not in [x['id'] for x in cl.get('/assistant/proposal-writer/proposals').json()['proposals']] and P.get(r3['id'])['status'] == 'discarded')
t('written proposals cannot be discarded', cl.post(f'/assistant/proposal-writer/work/{p9["id"]}/discard', json={}).status_code == 404)
calls.clear(); MODE['qa'] = [QA2]
pw2 = cl.post('/assistant/proposal-writer/proposals', json={'title': 'Fabric data security baseline', 'organisation': 'NSC', 'brief': brief, 'sections': [{'title': 'Approach'}],
                                                            'rate_card': card, 'work_id': r1['id']}).json()['id']
done2 = wait(pw2)
t('writing it turns the form into the proposal (one item, not two)', pw2 == r1['id'] and done2['status'] == 'done')
r4 = wsave(id=r1['id'], form={'title': 'Fabric v2', 'brief': 'Changed.'}).json()
t('changes after writing start a new proposal in progress; the written one is kept', r4['created'] and r4['id'] != r1['id'] and P.get(r1['id'])['status'] == 'done')
t('another assistant\'s id cannot be used', cl.post('/assistant/hr-policy/work', json={'form': {'title': 'x'}}).status_code == 404)
t('cross-origin saves are refused', cl.post('/assistant/proposal-writer/work', headers={'origin': 'http://evil.example'}, json={'form': {'title': 'x'}}).status_code == 403)
t('a big price book can be sent with the proposal', cl.post('/assistant/proposal-writer/proposals', json={'title': 'T', 'brief': 'short'}).status_code == 400
  and len(P.clean_rate_card([{'role': f'R{i}', 'cost': 1, 'sell': 2} for i in range(60)])) == 60)
pg2 = cl.get('/assistant/proposal-writer').text
t('the page has the proposals list and saves as you work', all(x in pg2 for x in ('id="wb-list"', 'id="wb-new"', 'Saves itself as you work', "api('/work'")))
adm = cl.get('/admin').text
t('Alice opens assistants on a stage: the tile grows into it, a splash while it loads, back to Alice or a new window', all(x in adm for x in ('id="as-stage"', 'Back to Alice', 'Open in a new window', 'embed', 'as-splash', 'clipPath')))
t('assistant pages hide their own top bar on the stage', 'html.embed .topbar' in pg2 and 'html.embed .topbar' in cl.get('/assistant/hr-policy').text)

# ---------------- Argus; accepting and rejecting Argus's fixes; what each proposal cost ----------------
t('the QA agent is Argus', A.get(P.QA)['name'] == 'Argus: proposal QA' and P.QA_PROMPT.startswith('You are Argus'))
with s.db() as c:
    c.execute("UPDATE agents SET name='Proposal QA' WHERE id=?", (P.QA,)); c.execute("DELETE FROM activity WHERE action='agent_renamed'")
import importlib; importlib.reload(A)
t('an existing QA agent is renamed Argus once, logged', A.get(P.QA)['name'] == 'Argus: proposal QA')
done_ = P.get(pw2)
t('each written proposal carries what its AI calls cost', done_['context']['ai_cost']['writing'] > 0
  and any(x['ai_cost'] > 0 for x in cl.get('/assistant/proposal-writer/proposals').json()['proposals']))
calls.clear(); MODE['qa'] = [QA2]
fixes = [{'section': 'Approach', 'issue': 'No owners against next steps', 'fix': 'Add a role to each step', 'severity': 'low'}]
rv = cl.post(f'/assistant/proposal-writer/proposals/{pw2}/revise', json={'fixes': fixes})
d3 = wait(pw2)
wr = [c for c in calls if c['system'].startswith('You are the proposal writer')]
t('accepted fixes go to the writer, Argus checks again, the document is rebuilt', rv.status_code == 200 and d3['status'] == 'done'
  and 'Add a role to each step' in wr[0]['messages'][0]['content'] and 'Apply ONLY these fixes' in wr[0]['messages'][0]['content']
  and d3['qa'][-1]['source'] == 'your 1 accepted fix' and d3['document_id'])
t('the revision adds to the proposal\'s AI cost', P.get(pw2)['context']['ai_cost']['writing'] > done_['context']['ai_cost']['writing'])
t('no accepted fixes: refused', cl.post(f'/assistant/proposal-writer/proposals/{pw2}/revise', json={'fixes': [{'x': 1}]}).status_code == 400)
t('the page offers accept, reject and revise for each fix', all(x in cl.get('/assistant/proposal-writer').text for x in ("'Accept'", "'Reject'", 'accepted fix', '/revise')))

# ---------------- rejected fixes stay rejected; accepted ones are checked ----------------
QA3 = dict(QA2, issues=[{'severity': 'medium', 'section': 'Commercials', 'issue': 'Expenses terms are not stated in the commercials', 'fix': 'State the expenses policy'},
                        {'severity': 'low', 'section': 'Next steps', 'issue': 'Next steps have no named owners', 'fix': 'Name an owner for each step'}])
QA4 = dict(QA2, issues=[{'severity': 'medium', 'section': 'Commercials', 'issue': 'The commercials do not state expenses terms', 'fix': 'State how expenses are charged'},
                        {'severity': 'low', 'section': 'Next steps', 'issue': 'Next steps still have no named owners', 'fix': 'Name an owner for each step'},
                        {'severity': 'low', 'section': 'Approach', 'issue': 'Phase two has no exit criteria', 'fix': 'Add exit criteria'}])
calls.clear(); MODE['qa'] = [QA4]
rv2 = cl.post(f'/assistant/proposal-writer/proposals/{pw2}/revise', json={'fixes': [QA3['issues'][1]], 'rejected': [QA3['issues'][0]]})
d4 = wait(pw2)
qmsg = [c for c in calls if c['system'].startswith('You are Argus')][-1]['messages'][0]['content']
t('Argus is told what you accepted and what you rejected', rv2.status_code == 200 and "AUTHOR'S DECISIONS ON EARLIER FIXES" in qmsg
  and 'REJECTED (do not raise again)' in qmsg and 'Expenses terms are not stated' in qmsg and 'ACCEPTED (applied; check them)' in qmsg)
iss = [i['issue'] for i in d4['qa'][-1]['issues']]
t('a rejected point raised again in other words is left out', not any('expenses' in x.lower() for x in iss) and d4['qa'][-1]['rejected_not_raised'] == 1)
t('an accepted fix Argus says is still not done, and new points, are kept', 'Next steps still have no named owners' in iss and 'Phase two has no exit criteria' in iss)
t('your decisions are kept with the proposal for later checks', [d['decision'] for d in P.get(pw2)['context']['fix_decisions'] if 'xpenses' in d['issue']] == ['rejected'])
calls.clear(); MODE['qa'] = [QA4]
cl.post(f'/assistant/proposal-writer/proposals/{pw2}/recheck', json={'sections': [{'title': x['title'], 'body': x['body']} for x in P.get(pw2)['draft']['sections']]})
d5 = wait(pw2)
t('rejections still stand when you re-check your own edits', not any('expenses' in i['issue'].lower() for i in d5['qa'][-1]['issues']))
calls.clear(); MODE['qa'] = [QA2]
cl.post(f'/assistant/proposal-writer/proposals/{pw2}/revise', json={
    'fixes': [dict(QA4['issues'][2], note='Use the council\'s own gateway names')],
    'rejected': [dict(QA2['issues'][0], note='The client asked us not to name owners until contract')]})
d6 = wait(pw2)
wmsg = [c for c in calls if c['system'].startswith('You are the proposal writer')][-1]['messages'][0]['content']
amsg = [c for c in calls if c['system'].startswith('You are Argus')][-1]['messages'][0]['content']
t('your note on an accepted fix goes to the writer', "gateway names" in wmsg and 'follow the note' in wmsg)
t('your reason for rejecting goes to Argus with the decision', "author's note: The client asked us not to name owners until contract" in amsg)
t('notes are kept with the decisions', any(d.get('note') == 'The client asked us not to name owners until contract' for d in P.get(pw2)['context']['fix_decisions']))
t('a secret in a note is refused, and not kept', cl.post(f'/assistant/proposal-writer/proposals/{pw2}/revise', json={'fixes': [dict(QA2['issues'][0], note='sk-proj-abcdefghijklmnopqrstuvwxyz0123456789ABCD')]}).status_code == 400
  and not any('sk-proj' in (d.get('note') or '') for d in P.get(pw2)['context']['fix_decisions']))
t('the page has a why box for each decision', 'fixwhy' in cl.get('/assistant/proposal-writer').text)
calls.clear(); MODE['qa'] = [QA2]
before_rounds = len(P.get(pw2)['qa'])
ro = cl.post(f'/assistant/proposal-writer/proposals/{pw2}/revise', json={'fixes': [], 'rejected': [dict(QA4['issues'][2], note='Exit criteria are in the client\'s own gateway process')]})
d7 = wait(pw2)
t('only rejections: nothing is rewritten, Argus checks again with your reasons', ro.status_code == 200 and not [c for c in calls if c['system'].startswith('You are the proposal writer')]
  and len(d7['qa']) == before_rounds + 1 and d7['qa'][-1]['source'] == 'your 1 rejected fix' and "client's own gateway process" in calls[-1]['messages'][0]['content'])
t('nothing accepted or rejected: refused', cl.post(f'/assistant/proposal-writer/proposals/{pw2}/revise', json={'fixes': [{}]}).status_code == 400)

# ---------------- the rate card changes after writing: update the pricing ----------------
calls.clear(); MODE['qa'] = [QA2]
pp = P.get(pw2); was = pp['pricing']['sell']
role0 = pp['pricing']['lines'][0]['role']
newcard = [dict(r, use=(r['role'] == role0), days=(20 if r['role'] == role0 else '')) for r in card]
rp = cl.post(f'/assistant/proposal-writer/proposals/{pw2}/reprice', json={'rate_card': newcard})
d8 = wait(pw2)
t('updating the pricing reprices from the ticked roles, your days win, unticked roles come out', rp.status_code == 200
  and [(l['role'], l['quantity']) for l in d8['pricing']['lines']] == [(role0, 20)] and d8['pricing']['sell'] != was)
t('Argus checks the new price and the document is rebuilt', d8['qa'][-1]['source'] == 'new pricing' and d8['document_id'] != pp['document_id']
  and any(c['system'].startswith('You are Argus') for c in calls))
t('nothing ticked: refused', cl.post(f'/assistant/proposal-writer/proposals/{pw2}/reprice', json={'rate_card': [dict(r, use=False) for r in card]}).status_code == 400)
t('the page offers Update the pricing and shows what changed', all(x in cl.get('/assistant/proposal-writer').text for x in ('Update the pricing', 'has changed since this proposal was priced', '/reprice')))
