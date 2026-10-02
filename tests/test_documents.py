"""Documents from chat: Word, Excel and PDF built locally, checked by Alice's rules, kept for download (not knowledge)."""
import _util  # first: throwaway data folder, dummy keys, no real model calls
from _util import t
import io, json, uuid, zipfile
from types import SimpleNamespace as NS
import substrate_store as s
s.init()
import documents as D, knowledge as K
import app
from fastapi.testclient import TestClient
cl = TestClient(app.app)

MD = '''# Summary
Costs are in £ and the dates “quoted” are 2026–27. **Bold point.**

## Actions
- First action
- Second action

1. Step one
2. Step two

| Item | Owner | Cost |
|---|---|---|
| Container image | Stefan | 0 |
| Azure build | Stefan | 450.50 |
''' + '\n\n'.join(f'Paragraph {i}. ' + 'Lorem ipsum dolor sit amet, consectetur adipiscing elit. ' * 6 for i in range(40))

w = D.create('docx', 'Programme update', MD)
data = D.get(w['id'])['data']
with zipfile.ZipFile(io.BytesIO(data)) as z:
    doc = z.read('word/document.xml').decode()
t('a Word document is built with headings, bullets, numbering and a table', w['name'] == 'Programme-update.docx'
  and 'Heading2' in doc and 'w:numId w:val="1"' in doc and 'w:numId w:val="2"' in doc and '<w:tbl>' in doc and 'Azure build' in doc)
t('Word text reads back through Alice\'s own reader', 'Container image' in K.docx_to_text(data) and '£' in K.docx_to_text(data))

p = D.create('pdf', 'Programme update', MD)
from pypdf import PdfReader
rd = PdfReader(io.BytesIO(D.get(p['id'])['data']))
text = ''.join(pg.extract_text() for pg in rd.pages)
t('a PDF is built, paginated and readable', len(rd.pages) >= 2 and 'Programme update' in text and 'Azure build' in text and '£' in text and 'Page 1 of' in text)
t('the PDF title is set', (rd.metadata or {}).get('/Title') == 'Programme update')

x = D.create('xlsx', 'Budget 2027', sheets=[{'name': 'Budget', 'rows': [['Item', 'Q1', 'Q2', 'Total'], ['Licences', '1,200', '1300', '=B2+C2'], ['Azure', 450.5, 500, '=B3+C3']]},
                                            {'name': 'Notes', 'rows': [['Note'], ['Prices exclude VAT']]}])
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(D.get(x['id'])['data']))
ws = wb['Budget']
t('an Excel workbook keeps sheets, numbers and formulas', wb.sheetnames == ['Budget', 'Notes'] and ws['B2'].value == 1200 and ws['D2'].value == '=B2+C2'
  and ws['B3'].value == 450.5 and ws['A1'].font.bold and ws.freeze_panes == 'A2')
x2 = D.create('xlsx', 'From a table', '| A | B |\n|---|---|\n| 1 | 2 |')
t('a pipe table given as text also makes a spreadsheet', openpyxl.load_workbook(io.BytesIO(D.get(x2['id'])['data'])).active['B2'].value == 2)

for bad, msg in [(('txt', 'X', 'hello'), 'format'), (('docx', '', 'hello'), 'title'), (('pdf', 'X', '   '), 'content'), (('xlsx', 'X', ''), 'sheet')]:
    try: D.create(*bad); ok = False
    except ValueError as e: ok = msg in str(e).lower()
    t(f'refused: {msg}', ok)
try: D.create('docx', 'Keys', 'Our key is api_key=sk-proj-abcdefghijklmnopqrstuvwxyz0123456789ABCD'); ok = False
except ValueError: ok = True
t('secrets are never written to a document', ok)
try: D.create('pdf', 'Brief', 'OFFICIAL-SENSITIVE\nOperational details follow.'); ok = False
except ValueError: ok = True
t('protectively marked content is never written to a document', ok)

r = cl.get(f'/documents/{w["id"]}/download')
t('documents download with the right type and name', r.status_code == 200 and r.content == data
  and 'wordprocessingml' in r.headers['content-type'] and 'Programme-update.docx' in r.headers['content-disposition'])
t('unknown document: 404', cl.get('/documents/nope/download').status_code == 404)
with s.db() as c:
    t('documents are not knowledge (models cannot read them as files)', not c.execute("SELECT 1 FROM files WHERE name LIKE 'Programme-update%'").fetchone())
    t('creating a document is logged', c.execute("SELECT count(*) FROM activity WHERE action='document_created'").fetchone()[0] >= 4)

# ---- through chat: the model calls create_document, the answer carries the download ----
seen = []
class Blk(NS):
    def model_dump(self, exclude_none=True): return dict(vars(self))
class R:
    async def create(self, **kw):
        seen.append(kw)
        usage = NS(model_dump=lambda: {'input_tokens': 100, 'output_tokens': 20, 'input_tokens_details': {'cached_tokens': 0}})
        if len(seen) == 1:
            args = json.dumps({'format': 'pdf', 'title': 'Leave summary', 'content': '# Leave\n- 25 days plus bank holidays'})
            return NS(output=[Blk(type='function_call', name='create_document', arguments=args, call_id='c1')], output_text='', usage=usage)
        return NS(output=[Blk(type='message', content=[])], output_text='Here is your PDF summary of leave.', usage=usage)
class O:
    def __init__(s_, **k): s_.responses = R()
    async def __aenter__(s_): return s_
    async def __aexit__(s_, *a): pass
class M:
    def __init__(s_, *a, **k): pass
    async def __aenter__(s_): return s_
    async def __aexit__(s_, *a): pass
    async def list_tools(s_): return [NS(name=n, description='', inputSchema={'type': 'object'}) for n in app.ALLOWED_TOOLS]
app.AsyncOpenAI = O; app.Client = M
cid = cl.post('/chats').json()['id']
r = cl.post('/chat', json={'chat_id': cid, 'request_id': uuid.uuid4().hex, 'text': 'Make me a PDF summary of the leave policy', 'provider': 'openai'})
turn = cl.get('/chats/' + cid).json()['turns'][0]
t('chat offers create_document to the model', seen and any(x.get('name') == 'create_document' for x in seen[0]['tools']))
t('the answer carries the document for download', r.status_code == 200 and turn['reply'].startswith('Here is your PDF') and turn['documents']
  and turn['documents'][0]['name'] == 'Leave-summary.pdf')
t('the model is told not to paste links', 'do not add links' in json.dumps(seen[1]['input']))
t('the chat page shows document downloads', 'doc-chip' in cl.get('/').text and '/documents/' in cl.get('/').text)
