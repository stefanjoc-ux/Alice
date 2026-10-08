"""Documents Alice creates on request in chat: Word (.docx), Excel (.xlsx) and PDF.

Built with the standard library and openpyxl only (no new packages). Word and PDF take simple markdown-style text
(# headings, - bullets, 1. numbered items, | pipe | tables |, **bold**); Excel takes sheets of rows.
Before a document is saved, Alice's own rules run on its text (secrets and protective markings are never written).
Documents are kept for download from the chat. They are not knowledge: nothing a model writes becomes knowledge
without your approval.
"""
import io
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone

import substrate_store as store

FORMATS = {'docx': ('Word', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
           'xlsx': ('Excel', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
           'pdf': ('PDF', 'application/pdf')}
MAX_TEXT, MAX_ROWS, MAX_SHEETS, MAX_COLS = 200_000, 20_000, 20, 60

with store.db() as c:
    c.execute('''CREATE TABLE IF NOT EXISTS generated_documents (id TEXT PRIMARY KEY, chat_id TEXT NOT NULL DEFAULT '',
        name TEXT NOT NULL, format TEXT NOT NULL, size INTEGER NOT NULL, original BLOB NOT NULL, text TEXT NOT NULL DEFAULT '',
        created_at TEXT NOT NULL)''')

TOOL_DESCRIPTION = (
    'Create a Word (.docx), Excel (.xlsx) or PDF document for the user to download. Use it only when the user asks for a '
    'document, spreadsheet or PDF. For docx and pdf put the whole document in content as simple markdown: # and ## '
    'headings, - bullets, 1. numbered items, | pipe | tables | with a header row, **bold**. For xlsx give sheets: each '
    'with a name and rows (first row = column headings; numbers as numbers; formulas as strings starting with =). '
    'The user gets a download link automatically: do not paste links; say what you made in one line.')
TOOL_SCHEMA = {'type': 'object', 'properties': {
    'format': {'type': 'string', 'enum': ['docx', 'xlsx', 'pdf']},
    'title': {'type': 'string', 'description': 'Document title, also used for the file name'},
    'content': {'type': 'string', 'description': 'docx/pdf: the document as simple markdown'},
    'sheets': {'type': 'array', 'description': 'xlsx: sheets', 'items': {'type': 'object', 'properties': {
        'name': {'type': 'string'}, 'rows': {'type': 'array', 'items': {'type': 'array', 'items': {}}}}, 'required': ['rows']}}},
    'required': ['format', 'title']}


def _slug(title, ext):
    s = re.sub(r'[^A-Za-z0-9 _.-]+', '', title).strip().replace(' ', '-')[:80] or 'document'
    return s + ext


# ---------------- markdown-style text to blocks ----------------
def blocks(text):
    out, para = [], []
    lines = (text or '').replace('\r\n', '\n').split('\n')
    def flush():
        if para: out.append(('p', ' '.join(para))); para.clear()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip(); s = line.strip()
        if not s: flush(); i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)', s)
        if m: flush(); out.append(('h', min(len(m.group(1)), 3), m.group(2).strip())); i += 1; continue
        if s.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', x) for x in cells if x): rows.append(cells)
                i += 1
            if rows: out.append(('table', rows))
            continue
        m = re.match(r'^[-*•]\s+(.*)', s)
        if m:
            flush(); items = []
            while i < len(lines) and re.match(r'^\s*[-*•]\s+', lines[i]): items.append(re.sub(r'^\s*[-*•]\s+', '', lines[i]).strip()); i += 1
            out.append(('ul', items)); continue
        m = re.match(r'^\d+[.)]\s+(.*)', s)
        if m:
            flush(); items = []
            while i < len(lines) and re.match(r'^\s*\d+[.)]\s+', lines[i]): items.append(re.sub(r'^\s*\d+[.)]\s+', '', lines[i]).strip()); i += 1
            out.append(('ol', items)); continue
        para.append(s); i += 1
    flush()
    return out


def _runs(text):
    """[(text, bold)] from **bold** markup."""
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    return [(p[2:-2], True) if p.startswith('**') and p.endswith('**') and len(p) > 4 else (p, False) for p in parts if p]


def plain(text):
    return re.sub(r'\*\*([^*]+)\*\*', r'\1', text)


# ---------------- Word ----------------
def _esc(t):
    return t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')


def _wp(text, style=None, num=None):
    ppr = ''
    if style or num:
        ppr = '<w:pPr>' + (f'<w:pStyle w:val="{style}"/>' if style else '') + \
              (f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num}"/></w:numPr>' if num else '') + '</w:pPr>'
    runs = ''.join(f'<w:r>{"<w:rPr><w:b/></w:rPr>" if b else ""}<w:t xml:space="preserve">{_esc(t)}</w:t></w:r>' for t, b in _runs(text))
    return f'<w:p>{ppr}{runs}</w:p>'


def _wtable(rows):
    cols = max(len(r) for r in rows)
    width = 9300 // cols
    grid = ''.join(f'<w:gridCol w:w="{width}"/>' for _ in range(cols))
    b = ''.join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="B9CBD8"/>' for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'))
    out = [f'<w:tbl><w:tblPr><w:tblW w:w="{width * cols}" w:type="dxa"/><w:tblBorders>{b}</w:tblBorders>'
           '<w:tblCellMar><w:left w:w="80" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tblCellMar></w:tblPr>'
           f'<w:tblGrid>{grid}</w:tblGrid>']
    for n, r in enumerate(rows):
        cells = ''
        for x in list(r) + [''] * (cols - len(r)):
            shade = '<w:shd w:val="clear" w:color="auto" w:fill="E3F1F6"/>' if n == 0 else ''
            txt = f'**{x}**' if n == 0 and x and not x.startswith('**') else x
            cells += f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{shade}</w:tcPr>{_wp(txt)}</w:tc>'
        out.append(f'<w:tr>{"<w:trPr><w:tblHeader/></w:trPr>" if n == 0 else ""}{cells}</w:tr>')
    out.append('</w:tbl><w:p/>')
    return ''.join(out)


def to_docx(title, content):
    import knowledge
    numbering = knowledge.NUMBERING.replace('</w:numbering>',
        '<w:abstractNum w:abstractNumId="1"><w:multiLevelType w:val="hybridMultilevel"/><w:lvl w:ilvl="0"><w:start w:val="1"/>'
        '<w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>'
        '</w:abstractNum><w:num w:numId="2"><w:abstractNumId w:val="1"/></w:num></w:numbering>')
    styles = knowledge.STYLES.replace('</w:styles>',
        '<w:style w:type="paragraph" w:styleId="Heading4"><w:name w:val="heading 4"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/>'
        '<w:pPr><w:keepNext/><w:spacing w:before="120" w:after="40"/><w:outlineLvl w:val="3"/></w:pPr><w:rPr><w:b/><w:i/></w:rPr></w:style></w:styles>')
    parts = [_wp(title, 'Title')]
    for b in blocks(content):
        if b[0] == 'h': parts.append(_wp(b[2], {1: 'Heading2', 2: 'Heading3', 3: 'Heading4'}[b[1]]))
        elif b[0] == 'p': parts.append(_wp(b[1]))
        elif b[0] == 'ul': parts += [_wp(x, num=1) for x in b[1]]
        elif b[0] == 'ol': parts += [_wp(x, num=2) for x in b[1]]
        elif b[0] == 'table': parts.append(_wtable(b[1]))
    ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
    sect = ('<w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
            '<w:pgMar w:top="1300" w:right="1300" w:bottom="1300" w:left="1300" w:header="600" w:footer="600" w:gutter="0"/></w:sectPr>')
    files = {
        '[Content_Types].xml': ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
            '<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>'
            '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/></Types>'),
        '_rels/.rels': ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/></Relationships>'),
        'word/_rels/document.xml.rels': ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/></Relationships>'),
        'docProps/core.xml': ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>' + _esc(title) + '</dc:title><dc:creator>Alice</dc:creator></cp:coreProperties>'),
        'word/document.xml': f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {ns}><w:body>{"".join(parts)}{sect}</w:body></w:document>',
        'word/styles.xml': styles, 'word/numbering.xml': numbering}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items(): z.writestr(name, data)
    return buf.getvalue()


# ---------------- Excel ----------------
def _cell(v):
    if isinstance(v, (int, float)) or v is None: return v
    s = str(v)
    if re.fullmatch(r'-?\d+', s.replace(',', '')) and len(s) < 16 and not s.startswith('0') or s == '0': return int(s.replace(',', ''))
    if re.fullmatch(r'-?\d+\.\d+', s.replace(',', '')): return float(s.replace(',', ''))
    return s[:32000]


def to_xlsx(title, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    used = set()
    for n, sh in enumerate(sheets[:MAX_SHEETS]):
        name = re.sub(r'[\[\]:*?/\\]', '', str(sh.get('name') or f'Sheet{n + 1}'))[:31] or f'Sheet{n + 1}'
        while name.lower() in used: name = (name[:28] + str(n + 1))
        used.add(name.lower())
        ws = wb.create_sheet(name)
        rows = [list(r)[:MAX_COLS] for r in (sh.get('rows') or [])[:MAX_ROWS] if isinstance(r, (list, tuple))]
        widths = {}
        for r in rows:
            ws.append([_cell(v) for v in r])
            for i, v in enumerate(r, 1): widths[i] = max(widths.get(i, 0), min(len(str(v)), 60))
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='075E79')
                c.alignment = Alignment(vertical='center', wrap_text=True)
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
        for i, wdt in widths.items(): ws.column_dimensions[get_column_letter(i)].width = max(10, wdt + 2)
    if not wb.sheetnames: wb.create_sheet('Sheet1')
    wb.properties.title = title; wb.properties.creator = 'Alice'
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ---------------- PDF (no external library) ----------------
W_REG = [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278, 556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556, 1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556, 333, 556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556, 556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584, 761, 556, 500, 222, 556, 333, 1000, 556, 556, 333, 1000, 667, 333, 1000, 500, 611, 500, 500, 222, 222, 333, 333, 350, 556, 1000, 333, 1000, 500, 333, 944, 500, 500, 667, 278, 333, 556, 556, 556, 556, 260, 556, 333, 737, 370, 556, 584, 333, 737, 333, 400, 584, 333, 333, 333, 556, 537, 278, 333, 333, 365, 556, 834, 834, 834, 611, 667, 667, 667, 667, 667, 667, 1000, 722, 667, 667, 667, 667, 278, 278, 278, 278, 722, 722, 778, 778, 778, 778, 778, 584, 778, 722, 722, 722, 722, 667, 667, 611, 556, 556, 556, 556, 556, 556, 889, 500, 556, 556, 556, 556, 278, 278, 278, 278, 556, 556, 556, 556, 556, 556, 556, 584, 611, 556, 556, 556, 556, 500, 556, 500]
W_BOLD = [278, 333, 474, 556, 556, 889, 722, 238, 333, 333, 389, 584, 278, 333, 278, 278, 556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 333, 333, 584, 584, 584, 611, 975, 722, 722, 722, 722, 667, 611, 778, 722, 278, 556, 722, 611, 833, 722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 333, 278, 333, 584, 556, 333, 556, 611, 556, 611, 556, 333, 611, 611, 278, 278, 556, 278, 889, 611, 611, 611, 611, 389, 556, 333, 611, 556, 778, 556, 556, 500, 389, 280, 389, 584, 761, 556, 500, 278, 556, 500, 1000, 556, 556, 333, 1000, 667, 333, 1000, 500, 611, 500, 500, 278, 278, 500, 500, 350, 556, 1000, 333, 1000, 556, 333, 944, 500, 500, 667, 278, 333, 556, 556, 556, 556, 280, 556, 333, 737, 370, 556, 584, 333, 737, 333, 400, 584, 333, 333, 333, 611, 556, 278, 333, 333, 365, 556, 834, 834, 834, 611, 722, 722, 722, 722, 722, 722, 1000, 722, 667, 667, 667, 667, 278, 278, 278, 278, 722, 722, 778, 778, 778, 778, 778, 584, 778, 722, 722, 722, 722, 667, 667, 611, 556, 556, 556, 556, 556, 556, 889, 556, 556, 556, 556, 556, 278, 278, 278, 278, 611, 611, 611, 611, 611, 611, 611, 584, 611, 611, 611, 611, 611, 556, 611, 556]
PAGE_W, PAGE_H, MARGIN = 595.28, 841.89, 56


def _width(text, size, bold=False):
    table = W_BOLD if bold else W_REG
    return sum(table[b - 32] if b >= 32 else 500 for b in text.encode('cp1252', 'replace')) * size / 1000


def _wrap(text, size, width, bold=False):
    lines, cur = [], ''
    for word in text.split():
        trial = (cur + ' ' + word) if cur else word
        if _width(trial, size, bold) <= width: cur = trial; continue
        if cur: lines.append(cur)
        while _width(word, size, bold) > width:              # very long word: break it
            k = len(word)
            while k > 1 and _width(word[:k], size, bold) > width: k -= 1
            lines.append(word[:k]); word = word[k:]
        cur = word
    if cur: lines.append(cur)
    return lines or ['']


def _pdf_str(text):
    raw = text.encode('cp1252', 'replace')
    return '(' + raw.replace(b'\\', b'\\\\').replace(b'(', b'\\(').replace(b')', b'\\)').decode('latin-1') + ')'


class _Pdf:
    def __init__(self, title):
        self.title, self.pages, self.ops, self.y = title, [], [], 0
        self.new_page()

    def new_page(self):
        if self.ops: self.pages.append(self.ops)
        self.ops, self.y = [], PAGE_H - MARGIN

    def need(self, h):
        if self.y - h < MARGIN + 20: self.new_page()

    def text(self, x, y, s, size, bold=False, colour=(0.08, 0.2, 0.29)):
        self.ops.append(f'BT /{"F2" if bold else "F1"} {size} Tf {colour[0]} {colour[1]} {colour[2]} rg {x:.2f} {y:.2f} Td {_pdf_str(s)} Tj ET')

    def rect(self, x, y, w, h, fill=None, stroke=None):
        if fill: self.ops.append(f'{fill[0]} {fill[1]} {fill[2]} rg {x:.2f} {y:.2f} {w:.2f} {h:.2f} re f')
        if stroke: self.ops.append(f'0.5 w {stroke[0]} {stroke[1]} {stroke[2]} RG {x:.2f} {y:.2f} {w:.2f} {h:.2f} re S')

    def para(self, text, size=10.5, bold=False, indent=0, before=0, after=5, colour=(0.08, 0.2, 0.29), bullet=None):
        width = PAGE_W - 2 * MARGIN - indent
        if text.startswith('**') and text.endswith('**') and len(text) > 4: bold, text = True, text[2:-2]
        lines = _wrap(plain(text), size, width, bold)
        lead = size * 1.35
        self.y -= before
        for n, line in enumerate(lines):
            self.need(lead)
            self.y -= lead
            if bullet and n == 0: self.text(MARGIN + indent - 12, self.y, bullet, size, False, colour)
            self.text(MARGIN + indent, self.y, line, size, bold, colour)
        self.y -= after

    def table(self, rows):
        cols = max(len(r) for r in rows)
        width = PAGE_W - 2 * MARGIN
        cw = width / cols
        size, pad = 9, 4
        for n, r in enumerate(rows):
            cells = [_wrap(plain(x), size, cw - 2 * pad, n == 0) for x in list(r) + [''] * (cols - len(r))]
            h = max(len(c) for c in cells) * size * 1.3 + 2 * pad
            if self.y - h < MARGIN + 20:
                self.new_page()
            top = self.y
            for i, lines in enumerate(cells):
                x = MARGIN + i * cw
                self.rect(x, top - h, cw, h, fill=(0.89, 0.945, 0.965) if n == 0 else None, stroke=(0.72, 0.8, 0.85))
                yy = top - pad
                for line in lines:
                    yy -= size * 1.3
                    self.text(x + pad, yy + 2, line, size, n == 0)
            self.y = top - h
        self.y -= 8

    def build(self):
        self.pages.append(self.ops)
        objs = ['<< /Type /Catalog /Pages 2 0 R >>', None,
                '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',
                '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>']
        kids = []
        total = len(self.pages)
        for n, ops in enumerate(self.pages, 1):
            footer = f'BT /F1 8 Tf 0.36 0.45 0.52 rg {MARGIN} 30 Td {_pdf_str(self.title)} Tj ET ' \
                     f'BT /F1 8 Tf 0.36 0.45 0.52 rg {PAGE_W - MARGIN - 50} 30 Td {_pdf_str(f"Page {n} of {total}")} Tj ET'
            stream = ('\n'.join(ops) + '\n' + footer).encode('latin-1', 'replace')
            objs.append(f'<< /Length {len(stream)} >>\nstream\n'.encode('latin-1') + stream + b'\nendstream')
            objs.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {PAGE_W} {PAGE_H}] /Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {len(objs)} 0 R >>')
            kids.append(f'{len(objs)} 0 R')
        objs[1] = f'<< /Type /Pages /Kids [{" ".join(kids)}] /Count {len(kids)} >>'
        objs.append(f'<< /Title {_pdf_str(self.title)} /Producer (Alice) /CreationDate (D:{datetime.now(timezone.utc).strftime("%Y%m%d%H%M%SZ")}) >>')
        out = io.BytesIO(); out.write(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
        offsets = []
        for i, o in enumerate(objs, 1):
            offsets.append(out.tell())
            body = o if isinstance(o, bytes) else o.encode('latin-1', 'replace')
            out.write(f'{i} 0 obj\n'.encode() + body + b'\nendobj\n')
        xref = out.tell()
        out.write(f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode())
        for off in offsets: out.write(f'{off:010d} 00000 n \n'.encode())
        out.write(f'trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info {len(objs)} 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
        return out.getvalue()


def to_pdf(title, content):
    p = _Pdf(title)
    p.para(title, 20, True, after=10, colour=(0.04, 0.23, 0.33))
    for b in blocks(content):
        if b[0] == 'h': p.para(b[2], {1: 15, 2: 12.5, 3: 11}[b[1]], True, before=8, after=4, colour=(0.03, 0.37, 0.47) if b[1] == 1 else (0.08, 0.2, 0.29))
        elif b[0] == 'p': p.para(b[1])
        elif b[0] == 'ul':
            for x in b[1]: p.para(x, indent=16, after=2, bullet='•')
            p.y -= 4
        elif b[0] == 'ol':
            for n, x in enumerate(b[1], 1): p.para(x, indent=18, after=2, bullet=f'{n}.')
            p.y -= 4
        elif b[0] == 'table': p.table(b[1])
    return p.build()


# ---------------- create, store, fetch ----------------
def create(fmt, title, content='', sheets=None, chat_id=''):
    """Build, check and keep a document. Returns {id, name, format, size}."""
    import rules_engine
    if fmt not in FORMATS: raise ValueError('Format must be docx, xlsx or pdf.')
    title = ' '.join((title or '').split())[:150]
    if not title: raise ValueError('Give the document a title.')
    content = content or ''
    if len(content) > MAX_TEXT: raise ValueError('That document is too long (200,000 characters at most).')
    if fmt == 'xlsx':
        if not sheets and content:                       # a pipe table given as text still works
            tables = [b[1] for b in blocks(content) if b[0] == 'table']
            sheets = [{'name': f'Sheet{i + 1}', 'rows': t} for i, t in enumerate(tables)]
        if not sheets or not isinstance(sheets, list): raise ValueError('Give the spreadsheet at least one sheet with rows.')
        text = title + '\n' + '\n'.join('\t'.join(str(v) for v in r) for s in sheets if isinstance(s, dict) for r in (s.get('rows') or [])[:MAX_ROWS] if isinstance(r, (list, tuple)))
    else:
        if not content.strip(): raise ValueError('Give the document some content.')
        text = title + '\n' + content
    name = _slug(title, '.' + fmt)
    rules_engine.check_file(text, name)              # secrets and protective markings are never written to a document
    data = to_docx(title, content) if fmt == 'docx' else to_xlsx(title, sheets) if fmt == 'xlsx' else to_pdf(title, content)
    return keep(fmt, name, data, text, chat_id)


def keep(fmt, name, data, text, chat_id=''):
    """Keep a finished document for download (callers have already run check_file on its text)."""
    did = uuid.uuid4().hex
    with store.db() as c:
        c.execute('INSERT INTO generated_documents(id,chat_id,name,format,size,original,text,created_at) VALUES (?,?,?,?,?,?,?,?)',
                  (did, chat_id or '', name, fmt, len(data), data, text[:MAX_TEXT], store.now()))
        store.audit(c, 'document_created', did, 'human_review', f'{FORMATS[fmt][0]}: {name} ({len(data):,} bytes)')
    store.stamp('document', did)
    return {'id': did, 'name': name, 'format': fmt, 'kind': FORMATS[fmt][0], 'size': len(data)}


def get(did):
    with store.db() as c:
        r = c.execute('SELECT id,name,format,original FROM generated_documents WHERE id=?', (did,)).fetchone()
    if not r: raise LookupError('Document not found.')
    return {'name': r['name'], 'media_type': FORMATS[r['format']][1], 'data': bytes(r['original'])}
