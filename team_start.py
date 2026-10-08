"""The Start a job screen (Stefan, 8 Oct 2026; he approved the mock-up): what the page needs, what each file is before it is added,
the "Ready to start" checklist worked out from the form, and the typical run time and AI cost of this team's past jobs (never
invented: hidden when there are none). Nothing here stores anything; starting the job is teams.start_job, with its own checks."""
import base64
import io
import re
import statistics
from datetime import datetime
from pathlib import Path

import substrate_store as store
import teams

ROLES = {'drawing': 'Drawing', 'spec': 'Specification', 'schedule': 'Schedule', 'template': 'Cost/pricing template', 'brief': 'Other'}
SCALE = re.compile(r'\b1\s*:\s*\d{1,4}\b|\bscale\b|\bscaled\b', re.I)
TYPES = {'.pdf': 'PDF', '.docx': 'Word', '.xlsx': 'Excel', '.xlsm': 'Excel', '.xls': 'Excel (old format)', '.csv': 'CSV', '.txt': 'Text', '.md': 'Text'}


def page(tid):
    """Everything the Start a job screen shows before anything is typed."""
    import organisations, pricing_templates, doc_library, rules_engine
    t = teams.get(tid)
    members = {m['id']: m for m in t['members']}
    jts = []
    for jt in t['job_types']:
        work = {}
        for s in jt['stages']:
            work.setdefault(s['member'], []).append(s.get('task') or s['title'])
        jts.append({'id': jt['id'], 'name': jt['name'], 'description': jt.get('description', ''), 'client_facing': teams.facing(jt),
                    'members': [{'id': mid, 'role': (members.get(mid) or {}).get('role', ''), 'initials': teams._initials((members.get(mid) or {}).get('role', '')),
                                 'lead': mid == teams.lead_id(t), 'does': tasks} for mid, tasks in work.items()],
                    'typical': typical(tid, jt['id'])})
    rs = rules_engine.rate_sources()
    try: orgs = [{'name': o['name'], 'client': bool(o.get('is_client'))} for o in organisations.listing()['organisations']]
    except Exception: orgs = []
    return {'team': {'id': t['id'], 'name': t['name'], **teams.identity(t), 'href': teams.team_url(t['id'])}, 'job_types': jts,
            'autonomy': t['autonomy'], 'autonomy_options': teams.AUTONOMY, 'roles': ROLES,
            'rate_sources': {'order': rs['order'], 'allowed': rs['allowed'], 'names': dict(rules_engine.RATE_SOURCES), 'on': rs['on'],
                             'href': '/admin/rules?rule=rate_sources#rules'},
            'templates': pricing_templates.team_page(tid), 'organisations': orgs, 'org_defaults': _org_defaults(orgs),
            'library': any(s['id'] for s in doc_library.sources()), 'icons': {k: v[1] for k, v in teams.ICONS.items()}, 'nav': teams._nav()}


def _org_defaults(orgs):
    import pricing_templates
    out = {}
    for o in orgs:
        p = pricing_templates.org_default(o['name'])
        if p: out[o['name']] = pricing_templates.describe(p)
    return out


def typical(tid, job_type):
    """This team's finished jobs of this type: the typical (median) time to reach your sign-off and AI cost. None when there are none."""
    import team_costs
    with store.db() as c:
        jobs = [dict(r) for r in c.execute("SELECT id, created_at, ai_cost FROM team_jobs WHERE team_id=? AND job_type=? AND status='done'", (tid, job_type))]
        firsts = {r[0]: r[1] for r in c.execute("SELECT job_id, min(created_at) FROM team_steps WHERE kind='signoff' AND job_id IN "
                                                "(SELECT id FROM team_jobs WHERE team_id=? AND job_type=? AND status='done') GROUP BY job_id", (tid, job_type))}
    secs, costs = [], []
    for j in jobs:
        if j['id'] in firsts:
            try: secs.append((datetime.fromisoformat(firsts[j['id']]) - datetime.fromisoformat(j['created_at'])).total_seconds())
            except ValueError: pass
        costs.append(float(j['ai_cost'] or 0))
    if not jobs: return None
    out = {'jobs': len(jobs)}
    if secs:
        m = statistics.median(secs)
        n, unit = ((max(1, round(m / 60)), 'minute') if m < 5400 else (round(m / 3600, 1), 'hour') if m < 172800 else (round(m / 86400), 'day'))
        out['run_time'] = f'about {n:g} {unit}{"" if n == 1 else "s"}'
    if costs: out['ai_cost'] = team_costs.money(statistics.median(costs))
    return out


def _bytes(data):
    try: return base64.b64decode(data or '', validate=True)
    except ValueError: raise ValueError('The file could not be read.') from None


def inspect(tid, name='', data=None, path=''):
    """What a file is, before it is added: its type, pages (or sheets or lines), its Purview label, Alice's guess at its role, whether a
    drawing gives a scale, and anything the rules would refuse. Nothing is stored and no block is logged (a preview)."""
    import doc_library, purview_labels, pricing_templates, rules_engine
    teams.get(tid)
    src = 'upload'
    if path:
        p = doc_library.resolve(path) or pricing_templates.resolve(path)
        if not p: raise ValueError('Not found in the document sources.')
        name, raw, src = p.name, p.read_bytes(), 'library'
    else:
        name = ' '.join(str(name or '').split())[:120]
        if not name: raise ValueError('Each document needs a name.')
        raw = _bytes(data)
        if len(raw) > 15 * 1024 * 1024: raise ValueError(f'{name} is larger than 15 MB.')
    ext = Path(name).suffix.lower()
    out = {'name': name, 'source': src, 'path': path or '', 'type': TYPES.get(ext, ext.lstrip('.').upper() or 'File'), 'pages': None, 'sheets': None,
           'lines': None, 'label': '', 'readable': True, 'problem': '', 'scale': None, 'template': None}
    lbl = purview_labels.read_label(name, raw)
    if lbl: out['label'] = lbl.get('name') or 'A label without a name'
    text = ''
    if ext == '.xls':
        out.update(readable=False, problem='Old Excel format (.xls): Alice cannot read it. Save it as .xlsx.')
    else:
        try:
            text = teams._doc_text(name, raw)
            if ext == '.pdf': out['pages'] = text.count('[Page ')
            elif ext in ('.xlsx', '.xlsm'):
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
                out['sheets'] = len(wb.sheetnames)
            elif ext in ('.csv', '.txt', '.md'): out['lines'] = sum(1 for x in text.splitlines() if x.strip())
        except Exception:
            out.update(readable=False, problem='Alice could not read this file. Use Word, PDF, Excel, CSV or text.')
    if text:
        tok = rules_engine.PREVIEW.set(True)                     # a preview: the same checks, no block lines
        try: rules_engine.check_file(text, name)
        except rules_engine.RuleViolation as e: out['problem'] = str(e)
        finally: rules_engine.PREVIEW.reset(tok)
    if ext in pricing_templates.EXT_OK and out['readable']:
        try:
            m = pricing_templates.detect_code(raw, name)
            if m: out['template'] = {'looks_like': _looks_like_template(raw, name, m), 'mode': m['mode'], 'sheets': len(m['sheets'])}
        except Exception: pass
    out['guess'] = guess(name, text, out)
    if out['guess'] == 'drawing': out['scale'] = bool(SCALE.search(text or ''))
    return out


def _looks_like_template(raw, name, m):
    """A spreadsheet whose header has a rate column that is mostly empty below it is a pricing template, not a priced schedule."""
    from openpyxl.utils import column_index_from_string as ci
    import pricing_templates
    wb = pricing_templates._load(raw, name)
    filled = empty = 0
    for s in m['sheets']:
        ws = wb[s['sheet']]
        rc, dc = ci(s['columns']['rate']), ci(s['columns']['description'])
        for r in range(s['header_row'] + 1, min(ws.max_row, s['header_row'] + 200) + 1):
            if ws.cell(r, dc).value in (None, ''): continue
            if ws.cell(r, rc).value in (None, '') or pricing_templates._is_formula(ws.cell(r, rc).value): empty += 1
            else: filled += 1
    return filled <= empty


def guess(name, text, info=None):
    """Alice's guess at a file's role, from its name and contents; Stefan can change it."""
    n = (name or '').lower()
    if info and (info.get('template') or {}).get('looks_like'): return 'template'
    if re.search(r'template|pricing|cost plan|boq|bill of quant', n) and Path(n).suffix in ('.xlsx', '.xlsm', '.csv', '.xls'): return 'template'
    if re.search(r'draw|\bplan\b|plans|elevation|section|\.dwg|layout|\bga\b', n): return 'drawing'
    if re.search(r'spec', n): return 'spec'
    if re.search(r'schedule|\.csv$|\.xlsx$|\.xlsm$', n): return 'schedule'
    if re.search(r'\bscale\b|\b1\s*:\s*\d{2,4}\b', text or '', re.I) and re.search(r'drawing|elevation|section', text or '', re.I): return 'drawing'
    return 'brief'


def check(tid, form):
    """The "Ready to start" checklist, worked out from the form: [{key, level (required | warning | info), ok, text}] and ready (every
    required item met). Warnings are things that may lead to estimates or approximate quantities; they do not stop the job."""
    import pricing_templates, clients
    t = teams.get(tid)
    f = form or {}
    jt = next((x for x in t['job_types'] if x['id'] == f.get('job_type')), None)
    docs = [d for d in f.get('documents') or [] if isinstance(d, dict)]
    work = [d for d in docs if d.get('role') != 'template']
    items = []

    def add(key, level, ok, text): items.append({'key': key, 'level': level, 'ok': bool(ok), 'text': text})
    title = ' '.join(str(f.get('title') or '').split())
    add('title', 'required', title, 'The job has a name' if title else 'Give the job a name')
    add('job_type', 'required', jt, f'What\'s needed: {jt["name"]}' if jt else 'Choose what\'s needed')
    words = len(str(f.get('brief') or '').split())
    add('brief', 'required', words >= 5, 'The brief says what is wanted' if words >= 5 else 'Write the brief: what is wanted, in a few sentences')
    add('documents', 'required', work, f'{len(work)} document{"s" if len(work) != 1 else ""} to work from' if work else 'Add the specification, schedules or drawings')
    bad = [d['name'] for d in docs if d.get('problem')]
    if bad: add('refused', 'required', False, 'Remove what Alice\'s rules refuse: ' + ', '.join(bad[:4]))
    tpl = f.get('template') or {}
    tpath = tpl.get('path') or ''
    if tpath:
        d = pricing_templates.describe(tpath)
        import rules_engine
        tok = rules_engine.PREVIEW.set(True)                         # checked as you type: nothing logged until the job starts
        try:
            pricing_templates.allowed_for(tpath, f.get('client_name') or '', teams.facing(jt) if jt else True) if d.get('client') else None
            allowed = True
        except ValueError as e:
            allowed, why = False, str(e)
        finally: rules_engine.PREVIEW.reset(tok)
        if not allowed: add('template', 'required', False, why)
        elif d['mapping'] == 'confirmed': add('template', 'info', True, f'Pricing template: {d["name"]} ({"shared" if d["shared"] else "for " + d["client"]}, mapping confirmed)')
        else: add('template', 'warning', False, f'Pricing template: {d["name"]}: {d["mapping_label"].lower()}. Confirm it before pricing finishes, or the template is not filled.')
    else:
        add('template', 'warning', False, 'No pricing template: Alice\'s own layout will be used')
    loc = ' '.join(str(f.get('location') or '').split())
    add('location', 'warning', loc, f'Location: {loc}' if loc else 'No location: Market Trends and the regional factor will use national figures')
    for d in work:
        if d.get('role') == 'drawing' and d.get('scale') is False:
            add('scale:' + d['name'][:60], 'warning', False, f'{d["name"]}: no scale found, so quantities taken from it are approximate and may lead to estimates')
    if work and all(d.get('role') == 'drawing' for d in work):
        add('only_drawings', 'warning', False, 'Only drawings: every quantity will be approximate. A specification or schedule gives firmer quantities.')
    cl = f.get('client_name') or ''
    if cl and cl not in clients.names(): add('client', 'info', True, f'{cl} is not marked Client: the job uses General material only for client-facing work')
    return {'items': items, 'ready': all(i['ok'] for i in items if i['level'] == 'required')}
